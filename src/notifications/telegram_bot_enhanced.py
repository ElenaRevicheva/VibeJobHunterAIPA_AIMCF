"""
Enhanced Telegram Bot for VibeJobHunter
Interactive commands for job hunting control and status

ADDED: Voice + free-form questions (additive, non-overwriting)
- Voice messages → transcribed with Whisper → answered from real job data
- Text questions (e.g. "what companies did I apply to today?") → honest answers

16 Oct 2026 (Elena, bilingual): answers are grounded in her shared memory (telegram_qa_helper.answer_grounded) and
come back in the language she used; screenshots are read; a voice note in gets a text AND a voice reply, and so does
a typed message that asks for voice. Only her TELEGRAM_CHAT_ID reaches the bot. Model and user text is sent plain.
PDF and Word .docx documents are read from their text layer and answered the same way (handle_document).
18 Oct 2026: a request for a PDF or Word file (EN/ES, typed or spoken) gets the answer AND the file, drawn by cto-aipa's
shared renderer; "letter for <company>" sends that deal's ✅ letter note word for word (_reply_with_document).
19 Oct 2026: the document desk — an explicit request rule (_doc_request), HTML as a third format, a fit document as the
deal's own kit notes word for word, a client quote from her catalog word for word, and every model-written document
held to her sources in code (telegram_qa_helper.gate_document).
"""

import os
import re
import asyncio
import tempfile
import shutil
import subprocess
import sys
from dotenv import load_dotenv
load_dotenv()
from datetime import datetime, timedelta
from typing import Optional, List, Dict
import json
import logging

try:
    from telegram import Update, InlineKeyboardButton, InlineKeyboardMarkup
    from telegram.ext import (
        Application,
        CommandHandler,
        CallbackQueryHandler,
        ContextTypes,
        MessageHandler,
        filters
    )
    TELEGRAM_AVAILABLE = True
except ImportError:
    TELEGRAM_AVAILABLE = False
    print(" python-telegram-bot not installed. Install with: pip install python-telegram-bot")

logger = logging.getLogger(__name__)

# ═══════════════════════════════════════════════════════════════════════════
# 16 Oct 2026 (Elena): the screens say only what really happens today. Checked on Oracle against the code and the
# journal; one copy each, shared by the /command and its button. Telegram Markdown v1: no "_" outside backticks.
# Buttons whose screens showed April data or a fake "0 jobs" (Today, Jobs, Stats, Outreach, Pause, Resume, Status,
# the "coming soon" CMO ones) are off the menu; their handlers are untouched.
# ═══════════════════════════════════════════════════════════════════════════
MENU_TEXT = """📋 *VIBEJOBHUNTER MENU*

💼 *YOUR JOB FLOW — nothing to remember*
1️⃣ New job → the kit fills the HubSpot deal by itself (company note, CV, letter, Comet prompt) → open 🔥 I act today → apply. The daily apply queue and kit alerts arrive here.
2️⃣ Interview booked (LinkedIn, WhatsApp, anywhere) → in the *CTO AIPA bot* send `/research_employer <company>` → if the company has an open job deal, deep research on the company + its products is saved on it and the bot confirms when OpenClaw has it (no deal yet: give Claude Code the posting link first)
3️⃣ In OpenClaw: `/new` → `prep me for <company>` → a mock interview on that company's likely questions, steering clear of what the research says not to claim

Choose an option below:"""

START_TEXT = """🤖 *VibeJobHunter*

I find and screen remote AI jobs every hour and put the good ones in HubSpot 🔥 I act today, where the kit writes your materials. *You apply* — I do not apply or email employers for you.

/menu — your job flow
/workflow — what runs by itself
/manual — what is yours
/help — commands that work"""

WORKFLOW_TEXT = """🔄 *HOW VIBEJOBHUNTER WORKS*

It finds and screens jobs, and the kit writes your materials. *You apply.* It does not apply, fill in forms or email employers for you.

*1️⃣ Finds jobs*: one cycle about every hour
• Job boards of 160 companies (Greenhouse, Ashby, Lever, Workable and more), Himalayas, Get on Board, YC OSS, Hacker News, Dice, RemoteOK, WeWorkRemotely, Remotive, Puente, AI-Native-Builder
• A Google Jobs feed, twice a day. Jobs that pass your rules and the judge go straight into 🔥 I act today with no alert here (the 08:15 queue lists them); borderline ones arrive here as BORDERLINE.
• Off on purpose: Wellfound, AI-Jobs, YC Work at a Startup, LinkedIn (Bright Data). Torre is failing right now.

*2️⃣ Screens them*
• Re-reads the live posting and drops it if closed
• The job must be fully remote, open to LATAM or worldwide, and AI-augmented work, with no posted pay under $3,000/month
• Scores each job: 60+ goes to your learned rules, then an AI judge · 55–59 is shown as a lead · under 55 is dropped. 55–59 leads skip your learned rules and the judge.

*3️⃣ Learns from you*: every hour
It re-reads your decided HubSpot job deals and the reason you gave for each No-fit, including screenshots on the deal. Those reasons become rules and examples for the judge.

*4️⃣ Alerts you here and creates the deal*
📋 Apply yourself (⚠️ UNVERIFIED when the posting could not be read) → a job deal in HubSpot 🔥 I act today

*5️⃣ The kit fills the deal*: every 10 min, from CTO AIPA
🔎 company brief · CV shaped to the posting, or your tailored CV with the reason · ✅ cover letter · 🎯 role defense · 🛡️ technical defense · 📋 Comet prompt

*6️⃣ Mornings*
• 07:55 Panama: closed postings are moved out of 🔥 I act today to Lost
• 08:15 Panama: the apply queue, sent here only when there are new jobs
• 08:45 Panama: a kit alert, sent here only if a deal is missing a piece

*7️⃣ Replies*
Your Zoho and Gmail inboxes are checked every hourly cycle, looking back 24 hours. A reply the AI reads as positive or a question (sometimes an intro or a newsletter, not only employers) goes to HubSpot and is alerted here once. An interview invite with a booking link should do the same; that was added 15 Oct and none has arrived yet.

*8️⃣ Evenings*
• 20:00 Panama: a daily public post on LinkedIn and Instagram, written by the AI (a template only if every model fails), posted without a preview

*👤 Your part*
Apply from 🔥 I act today → move the deal to ⏳ Sent, or to Lost with a reason. Interview booked anywhere → CTO AIPA bot `/research_employer <company>` → OpenClaw `/new`, then `prep me for <company>`.

*Switched off:* auto-apply, ATS form filling and founder emails. The old follow-up emailer still checks but has nothing to send (last one 24 May)."""

MANUAL_TEXT = """👤 *WHAT YOU DO BY HAND*

The bot finds and screens jobs, and the kit writes your materials on the HubSpot deal. *Applying, replying and booking calls are yours.* Nothing is sent for you.

*Every morning*
1️⃣ Open HubSpot → 🔥 I act today. When there are new jobs, the 08:15 Panama apply queue here lists them.
2️⃣ Open a deal and read the 🔎 company brief, 🎯 role defense and 🛡️ technical defense
3️⃣ Apply on the posting with the CV and ✅ cover letter on the deal. The 📋 Comet prompt is there if you use Comet for the form.
4️⃣ ⚠️ UNVERIFIED deal: open the link first and check the job is still open
5️⃣ Move the deal: ⏳ Sent once you have applied, or Lost with one line on why it is not a fit (a screenshot on the deal works too). Within the hour, that reason trains the judge.

*A job you found yourself*
Give the link to Claude Code and it becomes a deal with the same kit. This bot cannot do that.

*When an employer replies*
• The alert comes here. You answer and book the call yourself.
• Whose-turn reminders come at 07:00 Panama in the *CTO AIPA bot*, with a HubSpot task

*Interview booked* (LinkedIn, WhatsApp, email, anywhere)
1️⃣ *CTO AIPA bot:* `/research_employer <company>` → if the company has an open job deal, the research is saved on it and the bot replies "✅ OpenClaw has it now" (no deal yet: give Claude Code the posting link first)
2️⃣ *OpenClaw:* `/new` → `prep me for <company>` → a mock interview on that company's likely questions, steering clear of what the research says not to claim. Add `en español` for a Spanish interview.
3️⃣ After the call, move the deal on in HubSpot"""

HELP_TEXT = """📖 *COMMANDS THAT WORK*

/menu: your 3-step job flow, with buttons
/workflow: what runs by itself, and what is yours
/manual: your steps every morning, when a reply comes, and when an interview is booked
/priority: the YC priority company list. It does not change which jobs reach you. Skip `/priority refresh` and the 🔄 Refresh button on that screen: either freezes the whole bot and the job hunter for up to 2 minutes.
/help: this list
⚠️ The 📘 menu button posts publicly with no preview: LinkedIn and Instagram, plus a Hashnode and dev.to article whenever a CTO update is waiting.

*Arrives here by itself*
• 📋 Apply yourself / ⚠️ UNVERIFIED: a new job, already a HubSpot deal
• 08:15 Panama: the apply queue, only when there are new jobs
• 08:45 Panama: a kit alert, only if a deal is missing a piece
• Replies the AI reads as positive or a question (sometimes an intro or a newsletter, not only employers), once each
• BORDERLINE jobs from the Google Jobs feed, worth your eyes
• After each hourly cycle with new jobs: a cycle summary. Its applications and founder outreach lines read 0 because those features are switched off.
• After 12 empty cycles in a row: an alert that there are 0 new jobs
• On every restart: "AUTONOMOUS ENGINE STARTED". Its promises are old; this screen says what really runs.
• Daily digest around 15:00 Panama. Ignore its Auto-applied, Outreach, Follow-up, Warm Intro and Applied Today lines: those features are off, so they read 0, stay empty or show old data.

*Ask about your jobs* by text, voice note, screenshot or document (PDF or Word .docx), in English or Spanish. Answers come from your live HubSpot job deals, your verified facts and your corrections, and the bot says so when it does not know. A voice note gets a text and a voice reply; when you type, ask for voice ("reply with voice", "respóndeme con voz"). A screenshot of a posting gets fit and gaps, an interviewer's question gets an answer from your verified examples, a recruiter's message gets a reply draft. A PDF or Word document is read the same way, and a contract, NDA or offer gets its key terms, dates, money and the clauses to check (a reading aid, not legal advice). A scan with no text layer, a password-protected file or an old .doc is not guessed at: the bot says so and what to send instead. Russian voice notes are answered in English.

*For company research and interview prep:* CTO AIPA bot `/research_employer <company>` · OpenClaw `prep me for <company>` (English or Spanish)"""


def _menu_keyboard():
    """Only buttons whose screens are true today (16 Oct 2026)."""
    return InlineKeyboardMarkup([
        [InlineKeyboardButton("🔄 How it works", callback_data="workflow"),
         InlineKeyboardButton("👤 What is mine", callback_data="manual")],
        [InlineKeyboardButton("🎯 Priority Companies", callback_data="priority")],
        [InlineKeyboardButton("📘 Post publicly NOW (LinkedIn + IG)", callback_data="post_linkedin")],
    ])


# ═══════════════════════════════════════════════════════════════════════════
# THE SAME VOICE RULE in all three bots (Elena, 16 Oct 2026): a voice note in → the answer as text AND a voice note;
# a typed message that asks for voice (voice / voz / audio / "out loud" / "respóndeme con voz", any language) → text
# AND voice; otherwise text only. A product term ("voice agent", "agente de voz") is not a request.
# ═══════════════════════════════════════════════════════════════════════════
_VOICE_WORDS = re.compile(r"\b(voice|voz|voix|audio|aloud|out\s+loud|голос\w*|вслух|аудио)\b", re.IGNORECASE)
_VOICE_PRODUCT_TERMS = re.compile(
    r"\bvoice[\s-]*(agents?|assistants?|ai|bots?|apps?|latency|models?|interfaces?|cloning|actors?|over)\b|"
    r"\b(agentes?|asistentes?|bots?|modelos?|latencia|interfaz)\s+de\s+voz\b", re.IGNORECASE)


def _asks_for_voice(text: str) -> bool:
    """True when a typed message (or a screenshot caption) asks for the answer by voice."""
    if not text:
        return False
    return bool(_VOICE_WORDS.search(_VOICE_PRODUCT_TERMS.sub(" ", text)))


def _for_speech(text: str) -> str:
    """Port of cto-aipa src/telegram-bot.ts forSpeech: no links, code, Markdown, emoji or /commands; at most 1,500
    characters, cut on a sentence end found after character 500."""
    if not text:
        return ""
    s = re.sub(r"```[\s\S]*?```", " ", text)
    s = re.sub(r"https?://\S+", "", s)
    s = re.sub(r"(?m)^\s*/\w+.*$", "", s)
    s = s.replace("_", " ")
    s = re.sub(r"[*`\[\]#>]", "", s)
    try:
        import regex  # stdlib re has no \p{Extended_Pictographic}
        s = regex.sub(r"\p{Extended_Pictographic}|\uFE0F", "", s)
    except ImportError:
        s = re.sub("[\U0001F000-\U0001FAFF\u2600-\u27BF\uFE0F]", "", s)
    s = "\n".join(l for l in (re.sub(r"[ \t]+", " ", x).strip() for x in s.split("\n")) if l)
    if len(s) <= 1500:
        return s
    cut = s[:1500]
    end = max(cut.rfind(". "), cut.rfind(".\n"), cut.rfind("! "), cut.rfind("? "), cut.rfind("\n"))
    return (cut[:end + 1] if end > 500 else cut[:cut.rfind(" ")]).strip()


# ═══════════════════════════════════════════════════════════════════════════
# A DOCUMENT ON REQUEST (18 Oct 2026, Elena: "make bots recognise and create pdfs and word out of the task I send by
# voice or text in eng or esp"). THE RECOGNISER — 19 Oct 2026: the SAME rule as CTO AIPA's live desk, ported line by
# line from cto-aipa src/telegram-bot.ts (documentTask / documentRequestIn, 6582fe58), so both bots answer one message
# the same way. VJH's own earlier rule (four rounds of blocklists) missed her own "hazme en Word la carta para Pitz" and
# fired on "Haz clic en el PDF", "Crea una carpeta para los PDF", "Escribe 'urgente' en el asunto del correo con el PDF".
# The rule, in CTO's words:
#   · the request sits in ONE line — the whole message, or its first or last line (greeting / thanks lines aside); the
#     last line counts only when the first is not her own question, it is no posting line ("To apply…") and nothing
#     before it introduces someone else's text ("Esto dice el reclutador:", "Hi Elena,");
#   · a request VERB OPENS that line — before it only fillers ("Hola, …", "Ok, ahora …", "please", "VJH,"), "Para
#     <cliente>:" and an asking form ("¿me puedes…", "can you…", "necesito que me…", "I need you to…");
#   · and it GOVERNS a document kind or a format: between them only the words of its noun phrase (articles, the format,
#     a size, a tone, "please", a day) — so "hazme en Word la carta", "dame en PDF la evaluación" are requests, and
#     "haz CLIC en el PDF", "crea una CARPETA para los PDF" are not; after the kind only its format, a name, a
#     preposition or a word describing it ("la cotización BUTTON", "el informe RESPONSIVE", "una carta DEL BANCO" are not);
#   · a weak verb ("dame / mándame / necesito / quiero / I need") asks for a file only with its format said, and never
#     for "la / the / mi" one that exists; an infinitive only after "¿puedes / can you"; a bare plural ("create reports")
#     is a duty, never a request;
#   · not a request: a question (ending in "?" without "can you / ¿me puedes…"), a reminder or task, a note to herself,
#     "never mind / nah / olvídalo", later or not now ("next week", "cuando tenga el precio", "solo aquí en el chat"), a
#     document that exists ("…ya está enviada", "…que hiciste ayer"), a change to one ("…más corta", "…de nuevo" =
#     again), talk about the feature ("…every Monday", "can we…"), code talk ("Make the quote button work").
# Lower recall is bought for precision, as in CTO ("evaluación en PDF para Pitz" with no verb is chat).
# VJH keeps, beside CTO's rule: a message that is ONLY a format ("PDF", "en Word por favor", "pásamelo a Word") files
# the answer she got last (VJH's own follow-up), and three precision guards — a message greeting Elena by name is
# someone else's, an English "quote" with no client, service or format is a citation, and résumé / summary / brief
# count as a document only with a format said ("hazme un resumen de la semana" stays a chat answer).
# ═══════════════════════════════════════════════════════════════════════════
_DOC_FOLD = str.maketrans("áéíóúü", "aeiouu")
_DOC_FMT_TOKENS = {"pdf": "pdf", "word": "docx", "docx": "docx", "html": "html"}
# a message made only of these words and a format is "that answer as a file" ("PDF", "en Word por favor", "pásamelo").
# 19 Oct review: "por" is no bare word ("Gracias por el PDF!" re-filed the last answer): "por favor" is folded to
# "porfa" first (_doc_bare_tokens), and "gracias por / thanks for …" is never a request.
_DOC_BARE_WORDS = frozenset("""please pls thanks thank you gracias in as into to a an the en como un una el la
lo al de del of it this that these those esto eso ese este also too tambien y and or o file archivo format formato
version copy copia doc document documento your tu last ultima ultimo answer respuesta reply previous anterior mensaje
message now ahora can could puedes podrias me mi i want need quiero necesito quisiera would like get have same mismo
misma one ms microsoft send give make create generate export convert turn save put print haz hazme hazlo hazla
hazmelo hazmela crea creame genera generame exporta convierte conviertelo conviertela pasa pasalo pasala pasamelo
pasamela pon ponlo ponla ponmelo ponmela dame damelo damela envia enviame enviamelo manda mandame mandamelo version
otra vez again ok okay sí si yes ya pasas mandas envias das pones haces conviertes porfa urgente urgent asap over
back quick rapido pero but mejor better instead rather page pagina""".split())
_DOC_LABEL_RE = re.compile(r"^\s*(?:format|formato|file|archivo|type|tipo|cv|resume)\s*:")
_DOC_BARE_OWNED = re.compile(r"\b(?:of\s+(?:the|this|that|my|your)|del|de\s+(?:la|el|ese|esa|este|esta|mi|tu))\s+"
                             r"(?:ms\s+)?(?:pdf|word|docx|html)\b")
# a message written TO Elena by someone else (pasted): it opens by greeting her ("Hi Elena, … send me a PDF…")
_DOC_TO_ELENA = re.compile(
    r"^[\W_]*(?:(?:hi|hello|hey|dear|hola|buen[oa]s(?:\s+(?:dias|tardes|noches))?|estimad[oa]|querid[oa]|"
    r"good\s+(?:morning|afternoon|evening))[\s,!.:;-]+)?(?:ms\.?\s+|sra\.?\s+|senora\s+)?elena\b")
# her own services by name (an English "quote" of one of them is a client quote, not a citation)
_DOC_HER_SERVICE = re.compile(
    r"\b(?:diagnostic\w*|diagnostico|blueprint|analisis\s+profundo|full\s+analysis|auditoria\s+(?:web|tecnica)|"
    r"web\s+audit|informe\s+tecnico\s+preliminar|preliminary(?:\s+technical)?\s+web\s+audit|preliminar|"
    r"ai\s+growth\s+operator|plan\s+de\s+implementacion|implementation\s+plan)\b")
# The document kinds, in EN / ES (folded) — VJH's downstream vocabulary (_doc_kind_asked, _doc_two_kinds).
_DOC_KIND_NOUNS = (
    ("quote", r"cotizacion(?:es)?|presupuestos?|quotes?|quotations?|estimates?"),
    ("proposal", r"propuestas?(?:\s+comercial(?:es)?)?|proposals?"),
    ("evaluation", r"evaluacion(?:es)?|evaluations?|assessments?|scorecards?|defensas?|defen[cs]es?|"
                   r"requisito\s+por\s+requisito|requirement\s+by\s+requirement"),
    ("letter", r"cover\s+letters?|(?<!offer )letters?|cartas?(?!\s+(?:de\s+)?oferta)(?:\s+de\s+(?:presentacion|motivacion))?"),
    # 19 Oct review: "interview questions / preguntas de entrevista / mock interview" are chat practice (the grounded
    # Q&A reads the mock-interview sheet), never a prep file — "Hazme preguntas de entrevista para Pitz" asks to be asked
    ("prep", r"interview\s+prep(?:aration)?(?:\s+(?:sheets?|docs?|documents?|notes))?|prep\s+sheets?|cheat\s*sheets?|"
             r"preparacion(?:es)?(?:\s+(?:de|para)\s+(?:la\s+|una\s+)?entrevistas?)?"),
    ("presentation", r"presentacion(?:es)?|presentations?|pitch\s+decks?|decks?|slides|diapositivas"),
    ("promo", r"promocion(?:es)?|promos?|promotions?|flyers?|folletos?|brochures?|landing(?:\s+pages?)?|"
              r"paginas?\s+(?:web|html|de\s+aterrizaje)|web\s*pages?"),
    ("analysis", r"analisis|analys[ie]s|estudios?\s+de\s+mercado|market\s+(?:analysis|research|study)"),
    ("onepager", r"one[\s-]?pagers?|resumen(?:es)?\s+ejecutivos?|executive\s+summar(?:y|ies)|fichas?"),
    ("report", r"informes?|reportes?|reports?|resumen(?:es)?|summar(?:y|ies)|briefs?"),
    ("answer", r"documento|document|doc"),
)

# ── the port (cto-aipa src/telegram-bot.ts 9652-10167). \p{L} = [^\W\d_]; \p{L}\p{N}_ = \w. ──────────────────────────
_DL = r"[^\W\d_]"                                    # a letter, any script
_DNB = r"(?<!\w)"
_DNA = r"(?!\w)"
_DUP = "A-ZÀ-ÖØ-ÞА-ЯЁ"                               # \p{Lu}
_DLOW = "a-zß-öø-ÿа-яё"                              # \p{Ll}
_DOC_PDF = r"(?:pdf|p\.\s?d\.\s?f\.?|pe\s+de\s+efe|пдф)"
_DOC_WORD = r"(?:docx|ms[\s-]?word|word|ворд" + _DL + r"*)"
_DOC_HTML = r"(?:html?|хтмл)"
_DOC_FMT = f"(?:{_DOC_PDF}|{_DOC_WORD}|{_DOC_HTML})"
# The kinds, CTO's order (the earliest one in the line decides; on a tie, the first here). Beside CTO's words: VJH's
# kit names (scorecard, the role / technical defence), prep and cheat sheets, a market study, a web page — and
# résumé / summary / brief as "report_fmt", a document only when a format is said as the output.
_DOC_KINDS = [
    ("prep", r"(?:preparaci[oó]n\s+(?:para|de)\s+(?:la\s+|una\s+)?entrevista|gu[ií]a\s+(?:para|de)\s+(?:la\s+)?entrevista|"
             r"interview\s+prep(?:aration)?(?:\s+(?:sheets?|docs?|notes))?|prep\s+sheets?|cheat\s*sheets?|подготовк"
             + _DL + r"*\s+к\s+(?:интервью|собеседовани" + _DL + r"*))"),
    ("answer", r"(?:respuestas?\s+(?:a|de|para)\s+(?:la\s+|el\s+|las\s+)?(?:solicitud|aplicaci[oó]n|postulaci[oó]n|formulario|"
               r"preguntas)|(?:application|form)\s+answers?|answers?\s+(?:to|for)\s+(?:the\s+)?(?:application|form|questions)|"
               r"ответ" + _DL + r"*\s+на\s+вопрос" + _DL + r"*)"),
    ("letter", r"(?:carta(?!\s+(?:de\s+)?oferta)(?:\s+de\s+(?:presentaci[oó]n|motivaci[oó]n))?|cover\s+letter|(?<!offer )letter|"
               r"сопроводительн" + _DL + r"*\s+письм" + _DL + r"*|письм" + _DL + r"*)"),
    ("quote", r"(?:cotizaci[oó]n(?:es)?|presupuestos?|quotes?|quotations?|estimates?|смет" + _DL + r"*)"),
    ("proposal", r"(?:propuestas?(?:\s+comercial(?:es)?)?|proposals?|коммерческ" + _DL + r"*\s+предложени" + _DL
                 + r"*|предложени" + _DL + r"*)"),
    ("evaluation", r"(?:evaluaci[oó]n(?:es)?|valoraci[oó]n(?:es)?|evaluations?|assessments?|scorecards?|"
                   r"defensas?\s+(?:del\s+rol|t[eé]cnicas?)|(?:role|technical)\s+defen[cs]es?|оценк" + _DL + r"*)"),
    ("report", r"(?:informes?|reportes?|reports?|отч[её]т" + _DL + r"*|доклад" + _DL + r"*)"),
    ("presentation", r"(?:presentaci[oó]n(?:es)?|presentations?|pitch\s+decks?|decks?|slides|diapositivas|презентаци"
                     + _DL + r"*)"),
    ("analysis", r"(?:an[aá]lisis|analysis|analyses|estudios?\s+de\s+mercado|market\s+(?:research|study)|анализ" + _DL + r"*)"),
    ("promo", r"(?:promoci[oó]n(?:es)?|promos?|promotions?|landing(?:\s+page)?s?|p[aá]gina\s+de\s+aterrizaje|"
              r"p[aá]ginas?\s+(?:web|html)|web\s*pages?|flyers?|folletos?|brochures?|промо" + _DL + r"*|лендинг" + _DL + r"*)"),
    ("other", r"(?:one[\s-]?pagers?|resumen\s+ejecutivo|executive\s+summary|ficha(?:\s+t[eé]cnica)?|documento|document|"
              r"документ)"),
    ("report_fmt", r"(?:res[uú]men(?:es)?|summar(?:y|ies)|briefs?)"),
    ("cv", r"(?:cv|curr[ií]cul(?:um|o)(?:\s+vitae)?|r[eé]sum[eé]|resume|hoja\s+de\s+vida)"),
]
# The deal's kit documents — the cover letter, the evaluation (role / technical defence), interview prep, the CV — EXIST
# on her job deal, and VJH files them word for word. So "mándame LA carta para Pitz en Word", "send me THE evaluation for
# the Rwazi role in Word" ask for a file here, where CTO's weak-verb rule ("la / the" = one that exists → chat) would
# answer in chat. The one deliberate difference from CTO's rule; the format must still be said.
_DOC_KIT_KINDS = frozenset(("letter", "evaluation", "prep", "cv"))
_DOC_KIND_ANY = "(?:" + "|".join(p for _k, p in _DOC_KINDS) + ")"
_DOC_HTML_KINDS = frozenset(("presentation", "promo"))
# The format said AS THE OUTPUT: "en PDF", "as a Word doc", "un PDF", "PDF file", "cotización PDF", "PDF y Word", or
# alone at the end ("…, PDF por favor"). "the PDF reader" or "why the PDF export failed" ask for no file. (CTO's sixth
# alternative, a lookbehind of variable width, is _DOC_FMT_CHAIN below: "en PDF y Word".)
_DOC_AS_OUTPUT = re.compile("|".join([
    _DNB + r"(?:en|in|as|como|to|into|formato|format|в|во|на)\s+(?:(?:ms|formato|format|формате)\s+)?" + _DOC_FMT + _DNA,
    r"(?<![\w./\\-])" + _DOC_FMT + r"\s+(?:file|doc|document|version|format|formato|archivo|documento|versi[oó]n|файл"
    + _DL + r"*|документ" + _DL + r"*|" + _DOC_KIND_ANY + ")" + _DNA,
    r"\(\s*" + _DOC_FMT + r"(?:\s*(?:y|and|or|o|\+|/|,)\s*" + _DOC_FMT + r")?\s*\)",
    _DNB + _DOC_KIND_ANY + r"\s+(?:(?:en|in|в)\s+)?" + _DOC_FMT + _DNA,
    r"(?<![\w./\\-])" + _DOC_FMT + r"\s*(?:y|and|or|o|и|или|\+|/|&|,)\s*" + _DOC_FMT + _DNA,
    r"(?:^|[,;:.!—–-]\s*)" + _DOC_FMT + r"(?:\s+(?:por favor|please|pls|пожалуйста))?\s*[.!]*\s*$",
    r"^\s*" + _DOC_FMT + r"\s*[:—–-]",
    _DNB + r"(?:un|una|a|an)\s+" + _DOC_FMT + _DNA,
]), re.I)
_DOC_FMT_CHAIN = re.compile(r"\s*(?:y|and|or|o|и|или|\+|/|&|,)\s*" + _DOC_FMT + _DNA, re.I)
_DOC_FMT_END = re.compile(_DOC_FMT + r"$", re.I)
_DOC_VERB = re.compile(
    _DNB + r"(?:haz(?:me|melo|nos)?|h[aá]zmel[oa]|hacer(?:me|lo|melo)?|h[aá]game|haga|hagas|haces|har[ií]as|"
    r"crea(?:r|me|rme|melo|s)?|cr[eé]ame(?:lo)?|prepar(?:a|ar|ame|arme|amelo|as|es)|prep[aá]r[ae]me(?:lo)?|"
    r"genera(?:r|me|rme)?|gen[eé]r[ae]me(?:lo)?|generes|escrib(?:e|ir|eme|irme|as)|escr[ií]b[ae]me(?:lo)?|"
    r"redact(?:a|ar|ame|arme|as|es)|red[aá]ct[ae]me(?:lo)?|elabor(?:a|ar|ame|arme|as|es)|elab[oó]r[ae]me(?:lo)?|"
    r"arm(?:a|ar|ame|arme|es|as)|escribes|generas|[aá]rmame(?:lo)?|m[aá]ndame(?:lo)?|mandar(?:me)?|mandes|mandas|"
    r"env[ií]ame(?:lo)?|enviar(?:me)?|env[ií]es|env[ií]as|p[aá]same(?:l[oa])?|dame|darme|crees|necesito|quiero|quisiera|"
    r"make(?!\s+sure)|"
    r"create|write|draft|prepare|generate|build|produce|compose|put\s+together|turn\s+(?:it|this|that)\s+into|give\s+me|"
    r"send\s+me|i\s+need|i\s+want|i'd\s+like|сделай" + _DL + r"*|сделать|создай" + _DL + r"*|создать|напиши" + _DL
    + r"*|написать|подготовь" + _DL + r"*|подготовить|составь" + _DL + r"*|составить|сгенерируй" + _DL + r"*|сформируй"
    + _DL + r"*|оформи" + _DL + r"*|пришли|скинь|нуж(?:ен|на|но)|хочу)" + _DNA, re.I)
# A question ABOUT documents ("how do I make a PDF report?") is answered, not built; a Trello task or reminder that
# mentions one ("crea una tarea: enviar la cotización en PDF a Pitz") stays a task.
_DOC_QUESTION = re.compile(
    r"^\s*¿?\s*(?:how|why|what(?:'s|\s+is|\s+are)|where|did|does|do\s+you\s+have|have\s+you|has|is|are|was|were|c[oó]mo|"
    r"por\s*qu[eé]|qu[eé]\s+(?:es|son)|d[oó]nde|hiciste|has\s+hecho|ya\s+(?:hiciste|tienes)|как|почему|зачем|где|"
    r"что\s+такое|ты\s+сделал" + _DL + r"*)" + _DNA, re.I)
_DOC_NOT_A_FILE = re.compile(
    _DNB + r"(?:trello|tareas?|tasks?|recu[eé]rdame|remind\s+me|reminders?|recordatorios?|задач" + _DL + r"*|напомни"
    + _DL + r"*|карточк" + _DL + r"*)" + _DNA, re.I)
_DOC_CARD = re.compile(_DNB + r"(?:cards?|tarjetas?)(?!\s+de\s+(?:cr[eé]dito|d[eé]bito))" + _DNA, re.I)
_DOC_CARD_PAY = re.compile(r"(?:by|con|with|por|via|v[ií]a|credit|debit|en|de)\s+$", re.I)   # "pay by card"
_DOC_NOTE_TO_SELF = re.compile(
    r"^\s*(?:recuerda|recordar|acu[eé]rdate|no\s+olvides|remember|don'?t\s+forget|запомни|не\s+забудь)(?!\w)", re.I)
# Before the verb only: greetings and fillers (and the bot's own name, as CTO's has "cto / aipa"; voice fillers "eh /
# pues / um"), "Para <cliente>:" and an asking form.
_DOC_LEAD = re.compile(
    r"^[\s¿¡\"“«(]*(?:(?:hola|hi|hello|hey|oye|ok|okay|okey|vale|bueno|mira|buenas|buenos\s+d[ií]as|buenas\s+(?:tardes|"
    r"noches)|good\s+(?:morning|afternoon|evening)|porfa|por\s+favor|please|pls|cto|aipa|vjh|vibejobhunter|bot|eh|ehm|"
    r"em|pues|um|uh|ahora|now|tambi[eé]n|also|entonces|so|perfecto|perfect|genial|great|gracias|thanks|thank\s+you|"
    r"listo|este(?=\s*,)|привет|пожалуйста|слушай|теперь)[\s,!.:;—–-]+)*", re.I)    # held-out p8d: "este, hazme…"
_DOC_FOR_PREFIX = re.compile(r"^(?:para|for|для)\s+((?:[^\s,:;.!?]+\s+){0,4}[^\s,:;.!?]+)\s*[,:;—–-]\s*", re.I)
_DOC_ASK = re.compile(
    r"^(?:¿\s*)?(?:(?:(?:me|nos)\s+)?(?:puedes|podr[ií]as|pod[eé]s|puede|podr[ií]a)\s+(?:(?:por\s+favor|porfa)\s+)?|"
    r"si\s+(?:puedes|puede)\s+|(?:can|could|would|will)\s+you\s+(?:please\s+)?|if\s+you\s+can\s+(?:please\s+)?|"
    r"(?:necesito|quiero|quisiera)\s+que\s+|i\s+(?:need|want|would\s+like|'d\s+like)\s+you\s+to\s+|help\s+me\s+(?:to\s+)?|"
    r"ay[uú]dame\s+a\s+|let'?s\s+|(?:ты\s+)?(?:можешь|можете|сможешь|сможете)\s+(?:мне\s+)?|не\s+мог" + _DL
    + r"*\s+бы\s+(?:ты\s+|вы\s+)?(?:мне\s+)?)?(?:(?:me|nos)\s+)?$", re.I)
_DOC_POSTING_LINE = re.compile(
    r"(?<!" + _DL + r")(?:to\s+apply|apply\s+(?:now|here|at|via|by|through)|para\s+(?:postular|aplicar|postularte|"
    r"aplicar\s+al?)|post[uú]late|send\s+(?:us\s+)?your\s+(?:cv|resume|r[eé]sum[eé])|env[ií]a(?:nos)?\s+tu\s+(?:cv|"
    r"hoja\s+de\s+vida|curr[ií]culum))(?!" + _DL + r")|[\w.+-]+@[\w-]+\.[\w.]+", re.I)
_DOC_GREETING_LINE = re.compile(
    r"^[\s¡¿]*(?:(?:hola|hi|hello|hey|oye|buenas|buenos\s+d[ií]as|buenas\s+(?:tardes|noches)|good\s+(?:morning|"
    r"afternoon|evening)|gracias|muchas\s+gracias|thanks|thank\s+you|thx|saludos|un\s+abrazo|bye|chao|ok|okay|привет|"
    r"спасибо)[\s,!.¡?]*)+$", re.I)
# Between a request verb and the kind it asks for there is only that noun phrase.
_DOC_NP_WORDS = frozenset((
    "un una unos unas el la los las lo a an the some another otra otro otras otros my mi mis tu tus your su sus our "
    "nuestra nuestro this that these those esta este esa ese estas estos esas esos me te le les nos se you us it de del of "
    "con with en in as como para for por y and o or pdf word docx html htm ms formato format file archivo documento doc "
    "version "
    "versión please pls porfa favor ahora now ya quick quickly rápido rápida new nueva nuevo nuevas nuevos short corta "
    "corto breve brief simple sencilla sencillo formal full completa completo detailed detallada detallado professional "
    "profesional pequeña pequeño small little one-page página page hoja commercial comercial técnica técnico technical "
    "clara claro clear bonita bonito nice good buena bueno buen final updated actualizada actualizado мне нам новую новый "
    "новое короткую краткую короткий краткий urgente urgentemente pronto hoy mañana today tomorrow tonight asap lunes "
    "martes miércoles jueves viernes sábado domingo monday tuesday wednesday thursday friday saturday sunday by up dos "
    "tres cuatro cinco two three four five сегодня завтра срочно").split())
# English noun modifiers, allowed between the verb and the kind only ("a WEEKLY report", "an INVESTOR deck")
_DOC_PRE_MODS = frozenset((
    "weekly monthly quarterly annual yearly daily market sales marketing investor investors client customer business "
    "project status progress competitor competitors competition company product products service services bilingual "
    "executive pitch strategy growth seo geo aeo audit ai visibility website site web one-page two-page follow-up "
    "partnership collaboration pricing price cost budget").split())
# Code talk: with no format said, "THE report" plus a word of code names a part of a system, not a document to write.
_DOC_CODE_TALK = re.compile(
    r"(?<!" + _DL + r")(?:headers?|footers?|navbars?|sidebars?|hero\s+section|cabeceras?|encabezados?|pie\s+de\s+p[aá]gina|"
    r"cron(?:job)?s?|endpoints?|apis?|scripts?|functions?|funci[oó]n|funciones|modules?|m[oó]dulos?|pipelines?|"
    r"webhooks?|deploy" + _DL + r"*|despliegue|desplegar|repos?|repositor" + _DL + r"*|database|base\s+de\s+datos|db|"
    r"schema|esquema|servers?|servidor(?:es)?|pm2|components?|componentes?|commits?|branch|rama|merge|bugs?|refactor"
    + _DL + r"*|typescript|javascript|python|node|npm|css|frontend|backend|buttons?|bot[oó]n|botones|templates?|"
    r"plantillas?|generators?|generador(?:es)?|layers?|capa|flows?|workflows?|scenarios?|escenarios?|lambda|c[oó]digo|"
    r"code|tests?|unit\s+tests?|logs?|ui|ux|queue|cola|cach[eé]|config" + _DL + r"*|env|variables?)(?!" + _DL + r")",
    re.I)
_DOC_INDEFINITE = frozenset("un una unos unas a an another otra otro otras otros some new nueva nuevo nuevas nuevos".split())
_DOC_NUMERAL = re.compile(r"^(?:\d+(?:-" + _DL + r"+)?|dos|tres|cuatro|cinco|two|three|four|five)$")
_DOC_DEFINITE = frozenset("el la los las lo the my mi mis tu tus your su sus our nuestra nuestro this that these those "
                          "esta este esa ese estas estos esas esos".split())
_DOC_WEAK_VERB = re.compile(
    r"^(?:dame|darme|dar|m[aá]ndame(?:lo)?|mandar(?:me)?|mandes|mandas|env[ií]ame(?:lo)?|enviar(?:me)?|env[ií]es|p[aá]same(?:l[oa])?|"
    r"env[ií]as|necesito|quiero|quisiera|give\s+me|send\s+me|i\s+need|i\s+want|i'd\s+like|пришли|скинь|нуж(?:ен|на|но)|"
    r"хочу)$", re.I)
_DOC_INFINITIVE = re.compile(
    r"^(?:(?:hacer|crear|preparar|generar|escribir|redactar|elaborar|armar|mandar|enviar|dar)(?:me|lo|la|nos|melo|mela)?|"
    r"сделать|создать|написать|подготовить|составить|прислать)$", re.I)
_DOC_HELPER_BEFORE = re.compile(
    r"(?:^|[\s¿,])(?:puedes|podr[ií]as|pod[eé]s|podr[ií]a|puede|ayudas\s+a|ay[uú]dame\s+a|can\s+you|could\s+you|"
    r"would\s+you|will\s+you|help\s+me(?:\s+to)?|you\s+to|можешь|можете|сможешь|сможете|помоги)\s+"
    r"(?:(?:me|te|nos|please|por\s+favor)\s+)?$", re.I)
_DOC_EDIT_AFTER = re.compile(
    r"^\s*,?\s*(?:m[aá]s|menos|more|less|mejor|peor|better|worse|shorter|longer|bigger|smaller|clearer|simpler|nicer|"
    r"cleaner|again|de\s+nuevo|otra\s+vez|nuevamente|igual|the\s+same|look|looks|sound|seem|be|fit|se\s+vea|quede|"
    r"luzca|короче|длиннее|лучше|снова|ещё\s+раз)(?!" + _DL + r")", re.I)
_DOC_AFTER_OK = frozenset((
    "para por for sobre about acerca regarding re on de del of con with en in as como que that which who y and o or u e "
    "ni nor to al a per so plus versus vs desde me te le nos les you it о об обо про для по на в во с со и или к у "
    "mensual semanal trimestral anual diario diaria monthly weekly quarterly yearly annual daily personalizada "
    "personalizado personalized personalised ejecutiva ejecutivo executive larga largo long extensa extenso informal "
    "bilingüe bilingual atractiva atractivo persuasiva persuasivo convincente visual creativa creativo impactante moderna "
    "moderno elegante inicial preliminar oficial general comparativa comparativo competitiva competitivo estratégica "
    "estratégico financiera financiero objetiva objetivo honesta honesto honest interna interno externa externo dirigida "
    "dirigido addressed tailored adaptada adaptado aimed enfocada enfocado centrada centrado focused basada basado based "
    "lista listo ready explaining explicando showing mostrando summarizing summarising resumiendo highlighting destacando "
    "resaltando covering cubriendo comparing comparando outlining describing describiendo detailing detallando listing "
    "including incluyendo presenting presentando introducing thanking agradeciendo offering ofreciendo proposing "
    "proponiendo asking pidiendo telling diciendo contando recommending recomendando analyzing analysing analizando "
    "evaluating evaluando targeting dirigida real").split())
# VJH: how the kit's evaluation is asked for, right after the kind ("evaluación requisito por requisito", "step by step")
_DOC_AFTER_PHRASE = re.compile(r"\s*(?:requisito\s+por\s+requisito|requirement\s+by\s+requirement|punto\s+por\s+punto|"
                               r"point\s+by\s+point|paso\s+a\s+paso|step\s+by\s+step)(?!\w)", re.I)
# VJH: "Can you make one-pagers, or only PDFs?" asks what the bot can do (an alternative question), not for a file
_DOC_OR_ONLY = re.compile(r"(?<!\w)(?:or|o)\s+(?:only|just|simply|solo|s[oó]lo|solamente|nada\s+m[aá]s)(?!\w)[^?]*\?\s*$",
                          re.I)
_DOC_CONTAINER_OF = re.compile(
    r"(?:^|\s)(?:p[aá]gina|page|hoja|formato|format|plantilla|template)\s+(?:de|del|of|para|for)"
    r"(?:\s+(?:el|la|los|las|the|un|una|a|an|mi|my))?\s*$", re.I)
_DOC_AFTER_REFUSE = re.compile(
    r"^\s*(?:(?:de|of)\s+(?:recomendaci[oó]n|referencia|recommendation|reference)\s+)?(?:from(?!" + _DL + r")|от(?!" + _DL
    + r")|(?:de|del)\s+(?:(?:mi|mis|el|la|los|las|un|una|su|sus|tu|tus|nuestr[oa]s?)\s+)?(?:banco|plomer[oa]|pintor[a]?|"
    r"m[eé]dic[oa]|doctor[a]?|dentista|abogad[oa]|notari[oa]|contador[a]?|caser[oa]|arrendador[a]?|due[nñ][oa]|"
    r"propietari[oa]|exjef[ea]|ex\s+jef[ea]|jef[ea]|escuela|colegio|universidad|embajada|consulado|gobierno|ministerio|"
    r"migraci[oó]n|seguro|aseguradora|hospital|dise[nñ]ador[a]?|proveedor[a]?|arquitect[oa]|electricista|mec[aá]nic[oa]|"
    r"carpinter[oa]|vecin[oa]|profesor[a]?|maestr[oa]|empleador[a]?|landlord|laboratorio|sangre|orina|heces)(?!" + _DL
    + r"))", re.I)
# She says not as a file, or not now, or takes it back ("nah", "never mind", "olvídalo"), or puts it off.
_DOC_NOT_NOW = re.compile(
    _DNB + r"(?:(?:just\s+)?tell\s+me\s+(?:here|in\s+(?:the\s+)?chat)|d[ií]melo\s+(?:aqu[ií]|ac[aá]|por\s+aqu[ií])|"
    r"in\s+your\s+head|en\s+tu\s+cabeza|de\s+cabeza|mentally|mentalmente|nah|never\s*mind|forget\s+it|olv[ií]dalo|"
    # held-out p8b: "…, bueno no, déjalo" — "déjalo" ending the message or right after "no", "scratch that"
    r"(?:no,?\s+|mejor\s+)d[eé]jalo|d[eé]jalo(?:\s+as[ií])?(?=\s*[.!]*\s*$)|scratch\s+that|forget\s+about\s+it|"
    r"don'?t\s+bother|skip\s+it|"
    r"mejor\s+(?:no|nada)|next\s+(?:monday|tuesday|wednesday|thursday|friday|saturday|sunday)|el\s+pr[oó]ximo\s+(?:lunes|"
    r"martes|mi[eé]rcoles|jueves|viernes|s[aá]bado|domingo)|el\s+(?:lunes|martes|mi[eé]rcoles|jueves|viernes|s[aá]bado|"
    r"domingo)\s+que\s+viene|(?:once|when|as\s+soon\s+as|after)\s+i\s+(?:have|get|know|hear|talk|speak|receive)|"
    r"(?:cuando|en\s+cuanto|apenas|una\s+vez\s+que)\s+(?:tenga|sepa|reciba|hable|me\s+(?:digan|confirmen|respondan|"
    r"manden|pasen))|despu[eé]s\s+de\s+que|after\s+(?:the|our|my)\s+(?:call|meeting)|despu[eé]s\s+de\s+(?:la|nuestra|mi)\s+"
    r"(?:llamada|reuni[oó]n)|solo\s+aqu[ií]|aqu[ií]\s+(?:mismo\s+)?en\s+el\s+chat|en\s+el\s+chat|por\s+(?:el\s+)?chat|"
    r"in\s+(?:the\s+)?chat|here\s+in\s+(?:the\s+)?chat|just\s+here|no\s+file|without\s+a\s+file|sin\s+(?:archivo|"
    r"documento)|not\s+now|ahora\s+no|todav[ií]a\s+no|a[uú]n\s+no|not\s+yet|next\s+(?:week|month|year)|"
    r"la\s+pr[oó]xima\s+semana|la\s+semana\s+(?:que\s+viene|pr[oó]xima)|el\s+pr[oó]ximo\s+mes|el\s+mes\s+que\s+viene|"
    r"alg[uú]n\s+d[ií]a|someday|some\s+day|one\s+day|later|m[aá]s\s+tarde|не\s+сейчас|потом|позже)" + _DNA, re.I)
_DOC_COMPLAINT = re.compile(
    r"(?<!" + _DL + r")(?:otra\s+vez|cada\s+vez|siempre|nunca|always|never|every\s+time|again|una\s+y\s+otra)(?!" + _DL
    + r")", re.I)
_DOC_YOU_DO = re.compile(r"^(?:mandas|env[ií]as|haces|das|escribes|generas|creas|preparas)$", re.I)
# The first line introduces someone else's text, so the last line is theirs, not her request.
_DOC_INTRO_OTHER = re.compile(
    r":\s*$|(?<!" + _DL + r")(?:escribi[oó]|escribe|dice|dijo|dicen|wrote|writes|says|said|puso)(?=\s*(?:[:\"“«]|$))|"
    r"(?<!" + _DL + r")(?:(?:este|ese|esta|esa|su|el|this|that|his|her|their|the)\s+(?:mensaje|message|correo|email|"
    r"e-mail|mail|whatsapp|texto|text)|qu[eé]\s+(?:opinas|te\s+parece|piensas)|what\s+do\s+you\s+think|opinas|reenv"
    + _DL + r"*|forward" + _DL + r"*|resp[oó]nde" + _DL + r"*|contest[aá]" + _DL + r"*|reply|respond)(?!" + _DL + r")|"
    r"^[\s¡¿]*(?:hola|hi|hello|hey|dear|estimad[ao]|querid[ao]|buenas|buenos\s+d[ií]as|buenas\s+(?:tardes|noches))"
    r"[\s,]+elena(?!" + _DL + r")", re.I)
_DOC_POLITE = re.compile(
    r"^\s*¿?\s*(?:(?:can|could|would|will)\s+you\s+(?:please\s+)?|(?:me\s+|nos\s+)?(?:puedes|podr[ií]as|pod[eé]s)\s+|"
    r"me\s+(?:haces|har[ií]as|preparas|preparar[ií]as|creas|escribes|redactas|armas|generas|elaboras|mandas|env[ií]as|"
    r"das)" + _DNA + r"|(?:ты\s+)?(?:можешь|можете|сможешь|сможете)\s+|не\s+мог" + _DL + r"*\s+бы\s+)", re.I)
_DOC_NOT_FMT = re.compile(
    _DNB + r"(?:no|not|ni|nor|не|без)\s+(?:(?:en|in|as|como|a|an|un|una|el|la|в|во|на)\s+){0,2}"
    r"(?:(?:ms|formato|format|формате)\s+)?" + _DOC_FMT + _DNA, re.I)
_DOC_FMT_LEAD = re.compile(
    r"^\s*,?\s*(?:(?:en|in|as|como|to|into|formato|format|в|во|на|un|una|a|an|el|la|the)\s+){0,3}"
    r"(?:(?:ms|formato|format|формате)\s+)?" + _DOC_FMT + r"(?:\s*(?:y|and|or|o|и|или|\+|/|&|,)\s*(?:(?:en|in|as|a|an|un|"
    r"una)\s+)?" + _DOC_FMT + r")*(?:\s+(?:file|doc|document|version|archivo|documento|versi[oó]n|файл" + _DL
    + r"*|документ" + _DL + r"*))?" + _DNA, re.I)
_DOC_FMT_AFTER_KIND = re.compile(r"\s+(?:en|in|as|como|в|во)\s+(?:(?:ms|formato|format|формате)\s+)?" + _DOC_FMT + _DNA,
                                 re.I)
# 19 Oct review (fix 8): a format with "el / the / ese / mi" right in front of it names a file that already EXISTS.
# After a place word ("EN EL Word", "IN THE PDF", "AL PDF", "DEL PDF") it is a place: what she asks is done inside that
# file ("Escribe en el Word mi teléfono"), and a kind written as its complement ("en el PDF DE LA CARTA") is part of
# that place. With no place word ("haz el PDF de la carta más bonito", "make the PDF version shorter") it is the thing
# being changed. Either way it is never the verb's object, and never an output phrase ("the PDF version"), except in
# two shapes where it is the FORMAT of a document the line asks for: next to that document's kind ("hazme la carta para
# Pitz en el Word", "hazme en el Word la carta para Pitz"), or written "el PDF de la carta" with no place word ("Hazme
# el PDF de la carta para Pitz") — the kind then decides, and the checks after it (an edit, an existing file, a word
# that does not describe it) still apply.
_DOC_FMT_DEF_BEFORE = re.compile(
    r"(?:^|[\s,(¿¡\"“«])(?:el|la|lo|los|las|the|this|that|these|those|ese|esa|este|esta|esos|esas|estos|estas|mi|mis|"
    r"my|tu|tus|your|su|sus|our|nuestro|nuestra)\s+$", re.I)
_DOC_FMT_PLACE_BEFORE = re.compile(
    r"(?:^|[\s,(¿¡\"“«])(?:(?:en|in|into|inside|on|onto|to|a|de|of|from|desde|dentro\s+de|within|at)\s+(?:el|la|lo|los|"
    r"las|the|this|that|ese|esa|este|esta|mi|my|tu|your|su|our|nuestro|nuestra)|al|del)\s+$", re.I)
_DOC_FMT_OF_KIND = re.compile(
    r"\s+(?:(?:file|doc|document|version|archivo|documento|versi[oó]n)\s+)?(?:de|del|of)\s+(?:(?:la|el|las|los|the|this|"
    r"esta|este|esa|ese|mi|my|tu|your)\s+)?(?=" + _DOC_KIND_ANY + _DNA + ")", re.I)
# What follows the kind says the document already exists: "…para Pitz no llegó", "…que hiciste ayer", "…you made".
_DOC_EXISTS_AFTER = re.compile(
    r"^\s*(?:(?:no|not|не|nunca|never|still|todav[ií]a|a[uú]n|already|уже)\s+)?(?:ya|est[aá]n?|estaba|fue|fueron|era|"
    r"lleg[oó]|llegaron|sali[oó]|qued[oó]|tiene|ten[ií]a|sigue|parece|enviad[oa]s?|mandad[oa]s?|rechazad[oa]s?|"
    r"aprobad[oa]s?|was|were|is|are|got|has|had|failed|fails|rejected|looks|seems|came|arrived|went|did|does|doesn't|"
    r"didn't|isn't|wasn't|был" + _DL + r"*|пришл" + _DL + r"*|отправлен" + _DL + r"*|отклон[её]н" + _DL + r"*|выглядит|"
    r"(?:que|that|which)\s+(?:me\s+|te\s+|ya\s+)?(?:hiciste|hicimos|hice|preparaste|creaste|escribiste|mandaste|enviaste|"
    r"redactaste|generaste|mand[eé]|envi[eé])|(?:(?:that|which)\s+)?(?:you|we|i)\s+(?:made|wrote|created|prepared|sent|"
    r"did|generated|drafted)|(?:который|которую|которое|что)\s+(?:ты|вы|я)\s+(?:сделал|написал|прислал|подготовил|"
    r"отправил)" + _DL + r"*)" + _DNA, re.I)
# Talk about the feature, in the request's own clause: "…in PDF automatically every Monday", "can we / podemos …".
_DOC_META = re.compile(
    _DNB + r"(?:automatically|autom[aá]ticamente|автоматически|every\s+(?:day|week|month|morning|evening|monday|tuesday|"
    r"wednesday|thursday|friday|saturday|sunday)|cada\s+(?:d[ií]a|semana|mes|ma[nñ]ana|tarde|lunes|martes|mi[eé]rcoles|"
    r"jueves|viernes|s[aá]bado|domingo)|todos\s+los\s+(?:d[ií]as|lunes|martes|mi[eé]rcoles|jueves|viernes|s[aá]bados|"
    r"domingos)|кажд" + _DL + r"*|еженедельн" + _DL + r"*|ежедневн" + _DL + r"*|can\s+we|could\s+we|podemos|"
    r"podr[ií]amos|можем|можно\s+ли)" + _DNA, re.I)
_DOC_FILLER = frozenset((
    "una uno unos unas los las del por para con que favor porfa hola gracias puedes podrias podrías me mi the for and "
    "with please thanks hello this that can you could would мне для или пожалуйста привет спасибо можешь").split())
# A company name as she says it: capitalised words, Spanish joiners allowed ("Clínica de la Mujer").
_DOC_NAME = (f"[{_DUP}0-9][\\w&'’.\\-]*(?:\\s+(?:(?:de|del|la|las|los|&)\\s+)?[{_DUP}0-9][\\w&'’.\\-]*){{0,4}}")
_DOC_FOR_LEAD = re.compile(
    r"^\s*,?\s*(?:para|for|для|de|del|of|at|to|al|a)\s+(?:(?:the|la|el|los|las|mi|my|a|an|un|una|nuestr[oa]|our)\s+)*"
    r"(?:(?:empresa|compa[nñ][ií]a|company|cliente|client|cl[ií]nica|clinic|team|equipo|startup|puesto|role|компании|"
    r"клиента)\s+(?:(?:de|of|at|en)\s+)?)?" + _DOC_NAME + r"(?:\s+(?:team|equipo))?")
_DOC_NAME_STOP = frozenset((
    "pdf word docx doc html htm ms elena revicheva aideazz cto aipa vjh openclaw espaluz vibejobhunter atuona hubspot "
    "telegram trello cursor claude chatgpt gemini whatsapp linkedin google i ok ai ia seo geo aeo cv hola hello hi hey "
    "please gracias thanks english spanish russian espanol ingles ruso enero febrero marzo abril mayo junio julio agosto "
    "septiembre setiembre octubre noviembre diciembre january february march april may june july august september "
    "october november december lunes martes miercoles jueves viernes sabado domingo monday tuesday wednesday thursday "
    "friday saturday sunday api apis url sql json crm erp ux ui llm llms rag mcp sdk").split())
# AIdeazz's own offer and product names are never a company (CTO also reads the catalog's titles: written out here, so
# the recogniser reads no file while she waits).
_DOC_OWN_NAMES = [re.compile(re.escape(n).replace(r"\ ", r"\s+"), re.I) for n in sorted({
    "Quick AI Growth Operator Diagnostic", "Diagnóstico rápido AI Growth Operator", "Preliminary technical web audit",
    "Informe técnico preliminar (web)", "Implementation blueprint (full analysis)",
    "Análisis profundo + plan de implementación", "AI Growth Operators", "AI Growth Operator", "Growth Operators",
    "Growth Operator", "AI Visibility Audit API", "AI Visibility Audit", "Identity Audit", "AI Ops Wiki", "Audit API",
    "AIdeazz AI Lab", "AIdeazz", "EspaLuz", "VibeJobHunter", "VibeJob Hunter", "CTO AIPA", "OpenClaw",
    "ATUONA AI Film Studio", "ATUONA", "WHITESPACE", "ALGOM Alpha"}, key=len, reverse=True)]
_DOC_GEO = frozenset(("panama", "panama city", "ciudad de panama", "republica de panama", "colombia", "costa rica",
                      "mexico", "latam", "latinoamerica", "america latina", "latin america", "usa", "eeuu",
                      "estados unidos", "united states", "espana", "spain", "europa", "europe", "miami", "bogota",
                      "medellin", "chiriqui", "boquete", "coronado", "venezuela", "peru", "chile", "argentina",
                      "ecuador", "rusia", "russia"))
_DOC_SUBJ = re.compile(
    _DNB + r"(?:[Pp]ara|[Ff]or|[Дд]ля|al|[Ss]obre|[Aa]bout|[Aa]cerca\s+de|[Rr]egarding|[Tt]o|a|[Dd]irigid[oa]s?\s+a|"
    r"[Aa]ddressed\s+to|(?:[Pp]uesto|[Rr]ol|[Vv]acante|[Pp]osici[oó]n|[Cc]argo)\s+en|(?:[Rr]ole|[Jj]ob|[Pp]osition|"
    r"[Oo]pening)\s+(?:at|with|in))\s+(?:(?:[Tt]he|la|el|los|las|a|an|una|un|mi|my|nuestr[oa]|our|su|sus|their)\s+)*"
    r"(?:(?:empresa|compa[nñ][ií]a|company|cliente|client|cl[ií]nica|clinic|team|equipo|startup|negocio|business|marca|"
    r"brand|puesto|role|vacante|job|компании|клиента)\s+(?:(?:de|of|at|en)\s+)?)?(" + _DOC_NAME + ")")
_DOC_SUBJ_WEAK = re.compile(
    _DNB + r"(?:the|el|la|al|my|mi)\s+(" + _DOC_NAME + f")(?:\\s+[{_DLOW}](?:{_DL}|-)*){{0,2}}"
    r"\s+(?:job|role|position|posting|opening|vacancy|interview|offer)(?!" + _DL + r")|" + _DNB
    + r"(?:en|at|with|con)\s+(?:(?:la|el|the)\s+)?(" + _DOC_NAME + ")")
_DOC_SUBJ_OF = re.compile(
    r"^\s*,?\s*(?:de|del|of)\s+(?:(?:the|la|el|los|las|mi|my|su)\s+)*(?:(?:empresa|compa[nñ][ií]a|company|cliente|"
    r"client|cl[ií]nica|clinic|negocio|business|marca|brand)\s+(?:(?:de|of)\s+)?)?(" + _DOC_NAME + ")")
_DOC_REF_BEFORE = re.compile(
    r"(?<!" + _DL + r")(?:como|igual\s+(?:que|a)|parecid[oa]s?\s+a|similar(?:es)?\s+(?:a|to)|like|same\s+as|based\s+on|"
    r"basad[oa]s?\s+en|tomando\s+como|que\s+(?:hicimos|hice|hiciste|mandamos|mand[eé]|mandaste|enviamos|envi[eé]|"
    r"enviaste|preparamos|preparaste)|(?:we|i|you)\s+(?:did|made|sent|wrote|prepared)|anterior|previous|"
    r"last\s+(?:one|time)|the\s+other|del\s+otro|de\s+la\s+otra)(?!" + _DL + r")", re.I)
_DOC_ROLE_HEADS = None    # filled below (folded)
_DOC_ROLE_WORDS = None


def _doc_fold_name(s: str) -> str:
    """CTO docFold: accents off (NFD), lower case."""
    import unicodedata
    return re.sub(r"[̀-ͯ]", "", unicodedata.normalize("NFD", s or "")).lower()


_DOC_ROLE_HEADS = frozenset(_doc_fold_name(w) for w in (
    "engineer engineers ingeniero ingeniera ingenieros lead leads manager managers gerente director directora head "
    "developer developers desarrollador desarrolladora specialist especialista analyst analista architect arquitecto "
    "arquitecta consultant consultor consultora designer diseñador diseñadora coordinator coordinador coordinadora "
    "associate assistant asistente officer intern writer redactor redactora creator creador creadora trainer tutor "
    "expert experto experta strategist estratega generalist builder scientist científico científica researcher "
    "investigador investigadora sdet").split())
_DOC_ROLE_WORDS = frozenset(_doc_fold_name(w) for w in (
    "operations ops operaciones automation automatización growth product producto data datos ml ai ia software senior sr "
    "junior jr remote remoto remota staff principal founding fullstack full-stack backend frontend marketing sales ventas "
    "customer success solutions soluciones support soporte prompt llm machine learning research technical técnico "
    "técnica tech it qa devops cloud security seguridad mobile web video content contenido community social media "
    "project proyecto program programa mid level ii iii chief executive ejecutivo ejecutiva vp").split())
_DOC_KIND_ONLY = re.compile("^" + _DOC_KIND_ANY + "$", re.I)


def _doc_clean_name(raw: str) -> str:
    """CTO docCleanName: "Pitz's PDF" → "Pitz"; '' when nothing name-like is left."""
    def bad(w):
        return (re.sub(r"[^\w&]+|_", "", _doc_fold_name(w)) in _DOC_NAME_STOP
                or re.fullmatch(r"(?:de|del|la|las|los|&)", w, re.I) is not None or bool(_DOC_KIND_ONLY.match(w)))
    toks = re.sub(r"[.'’-]+$", "", re.sub(r"['’]s\b", "", raw or "", flags=re.I)).split()
    while toks and bad(toks[0]):
        toks.pop(0)
    while toks and bad(toks[-1]):
        toks.pop()
    name = re.sub(r"[.'’-]+$", "", " ".join(toks))
    return name if re.search(_DL + "{2,}", name) else ""


def _doc_company_name(raw: str) -> str:
    """CTO docCompanyName: a role title dropped from the end ("Clara AI Operations Lead" → "Clara"); never a place."""
    toks = _doc_clean_name(raw).split()

    def key(w):
        return re.sub(r"[^\w-]+|_", "", _doc_fold_name(w))
    if toks and key(toks[-1]) in _DOC_ROLE_HEADS:
        toks.pop()
        while toks and (key(toks[-1]) in _DOC_ROLE_WORDS or key(toks[-1]) in _DOC_ROLE_HEADS):
            toks.pop()
    n = _doc_clean_name(" ".join(toks))
    return n if n and re.sub(r"[^\w]+|_", " ", _doc_fold_name(n)).strip() not in _DOC_GEO else ""


def _doc_strip_lead(t: str) -> str:
    """CTO: t.replace(DOC_LEAD, '').replace(DOC_FOR_PREFIX, '').replace(DOC_LEAD, '')."""
    return _DOC_LEAD.sub("", _DOC_FOR_PREFIX.sub("", _DOC_LEAD.sub("", t, 1), 1), 1)


def _doc_subject(t: str, kind_end: int) -> str:
    """CTO docSubject: the ONE company the document is for ("Para <X>:" first, else a name after "para / for / sobre /
    about / de" past the kind, never one given as a reference — "como la que hicimos para Mundo Dental"); ''."""
    s = t
    for rx in _DOC_OWN_NAMES:
        s = rx.sub(lambda m: "·" * len(m.group(0)), s)             # same length: kind_end stays valid
    pre = _DOC_FOR_PREFIX.match(_DOC_LEAD.sub("", s, 1))
    if pre:
        m = re.search(_DOC_NAME, pre.group(1))
        n = _doc_company_name(m.group(0)) if m else ""
        if n:
            return n
    tail = s[kind_end:]
    stop = re.search(r"[.!?;\n]\s", tail)
    if stop:
        tail = tail[:stop.start()]
    lead = tail
    for _i in range(3):
        lead = _DOC_FMT_LEAD.sub("", lead, 1)
    of = _DOC_SUBJ_OF.match(lead)
    if of:
        n = _doc_company_name(of.group(1))
        if n:
            return n
    cands = sorted([(m.start(), m.group(1)) for m in re.finditer(r"[«\"“]([^»\"”\n]{2,60})[»\"”]", tail)]
                   + [(m.start(), m.group(1)) for m in _DOC_SUBJ.finditer(tail)])

    def weak():
        return sorted((m.start(), m.group(1) or m.group(2)) for m in _DOC_SUBJ_WEAK.finditer(tail))
    for lst in ([cands, weak()] if cands else [weak()]):
        for at, name in lst:
            rest = tail[tail.find(name, at) + len(name):]
            if re.match(r"\s+(?:job|role|position|opening|vacancy|puesto|vacante|cargo|posici[oó]n)\s+(?:at|with|in|en|"
                        f"con)\\s+[{_DUP}]", rest):
                continue
            if re.match(r"\s+(?:case|cases|caso|casos|example|ejemplo|study|estudio|model|modelo|template|plantilla)(?!"
                        + _DL + ")", rest, re.I):
                continue
            before = tail[:at]
            frm = max([before.rfind(","), before.rfind(";"), before.rfind(":")]
                      + [m.end() - 1 for m in re.finditer(r"\s(?:pero|but|sino|y\s+no)\s", before, re.I)])
            if _DOC_REF_BEFORE.search(before[frm + 1:]):
                continue
            n = _doc_company_name(name)
            if n:
                return n
    return ""


def _doc_outputs(t: str) -> List[tuple]:
    """CTO DOC_AS_OUTPUT_G matches as (start, end, text): the formats said as the output, and a format chained after
    one ("en PDF y Word")."""
    outs = []
    for m in _DOC_AS_OUTPUT.finditer(t):
        outs.append((m.start(), m.end(), m.group(0)))
        end = m.end()
        while _DOC_FMT_END.search(t[:end]):
            c = _DOC_FMT_CHAIN.match(t, end)
            if not c:
                break
            outs.append((c.start(), c.end(), c.group(0)))
            end = c.end()
    return outs


def _doc_without_outputs(t: str) -> str:
    """t with every output phrase blanked (CTO: scope.replace(DOC_AS_OUTPUT_G, ' '))."""
    out, pos = [], 0
    for a, b, _x in sorted(_doc_outputs(t)):
        if a >= pos:
            out += [t[pos:a], " "]
            pos = b
    return "".join(out) + t[pos:]


def _doc_ref_formats(t: str, t_pos: str, outs: List[tuple], kinds: List[tuple]) -> tuple:
    """Fix 8 (see _DOC_FMT_DEF_BEFORE): the formats written with "el / the …" in front. → (outs without them, plus
    the ones that are a document's format; the starts of the format words that are never the verb's object; the starts
    of the kind words that are part of a place, "en el PDF DE LA CARTA")."""
    keep, skip, part_of = [], {}, set()
    for m in re.finditer(_DNB + _DOC_FMT + _DNA, t_pos, re.I):
        pre = t_pos[:m.start()]
        place = bool(_DOC_FMT_PLACE_BEFORE.search(pre))
        if not place and not _DOC_FMT_DEF_BEFORE.search(pre):
            continue
        skip[m.start()] = m.end()
        of = _DOC_FMT_OF_KIND.match(t_pos, m.end())
        if of:
            if place:
                part_of.update(h[0] for h in kinds if h[0] == of.end())
            else:
                keep.append((m.start(), m.end(), m.group(0)))
            continue
        before = any(h[3] <= m.start() and not re.search(r"[.;:!?]", t[h[3]:m.start()]) for h in kinds)
        nxt = next((h for h in kinds if h[0] >= m.end()), None)
        gap = [w for w in re.split(r"[\s,]+", t[m.end():nxt[0]].lower()) if w] if nxt else None
        if before or (gap is not None and not re.search(r"[.;:!?]", t[m.end():nxt[0]]) and len(gap) <= 4
                      and all(w in _DOC_NP_WORDS and w not in ("de", "del", "of") for w in gap)):
            keep.append((m.start(), m.end(), m.group(0)))
    if skip:
        outs = [o for o in outs if not any(o[0] < b and o[1] > a for a, b in skip.items())] + keep
    return outs, set(skip), part_of


def _doc_not_a_file(t: str) -> bool:
    """A Trello task, a reminder, a card (never "pay by card")."""
    return bool(_DOC_NOT_A_FILE.search(t)) or any(not _DOC_CARD_PAY.search(t[:m.start()]) for m in _DOC_CARD.finditer(t))


def _doc_request_in(t: str, whole: str) -> Optional[dict]:
    """CTO documentRequestIn: one line → {kind (CTO's), formats, fmt_said, word, kword, subject} or None. word = the
    kind or format word the verb governs; kword = the kind's own word ('' when a format stands alone: "un PDF")."""
    if not t or _DOC_QUESTION.match(t) or _doc_not_a_file(t) or _DOC_NOTE_TO_SELF.match(t) or _DOC_NOT_NOW.search(t):
        return None
    if re.search(r"\?\s*$", t) and not _DOC_POLITE.match(_doc_strip_lead(t)):
        return None
    t_pos = _DOC_NOT_FMT.sub(lambda m: " " * len(m.group(0)), t)      # same length: indices stay valid
    outs = _doc_outputs(t_pos)
    hits = []                                     # (at, order, kind, end, word, the kind's own word)
    for order, (k, rx) in enumerate(_DOC_KINDS):
        for m in re.finditer(_DNB + rx + _DNA, t, re.I):
            hits.append((m.start(), order, k, m.end(), m.group(0), m.group(0)))
    hits.sort(key=lambda h: (h[0], h[1]))
    outs, ref_fmts, in_place = _doc_ref_formats(t, t_pos, outs, hits)      # fix 8: "en el PDF", "the PDF version"
    # "un resumen" is a file only with its format said
    hits = [(h[0], h[1], "report" if h[2] == "report_fmt" else h[2]) + h[3:] for h in hits
            if h[0] not in in_place and (h[2] != "report_fmt" or outs)]
    for m in re.finditer(_DNB + _DOC_FMT + _DNA, t_pos, re.I):
        at = m.start()
        if at in ref_fmts:
            continue                                                 # an existing file, never the verb's object
        nxt = next((h for h in hits if h[0] > at and h[1] < 99 and not re.search(r"[.;:!?]", t[at:h[0]])), None)
        hits.append((at, 99, nxt[2] if nxt else "other", m.end(), m.group(0), nxt[4] if nxt else ""))
    if not hits:
        return None
    hits.sort(key=lambda h: (h[0], h[1]))

    def verb_ok(v) -> bool:
        at = v.start()
        if not _DOC_ASK.match(_doc_strip_lead(t[:at])):
            return False
        if _DOC_YOU_DO.match(v.group(0)) and _DOC_COMPLAINT.search(t):
            return False                     # "Me mandas una cotización mal hecha OTRA VEZ": a complaint
        return not (_DOC_INFINITIVE.match(v.group(0)) and not _DOC_HELPER_BEFORE.search(t[max(0, at - 60):at]))
    verbs = [v for v in _DOC_VERB.finditer(t) if verb_ok(v)]

    def governed(h) -> bool:
        for v in verbs:
            if v.end() > h[0]:
                continue
            between = t[v.end():h[0]]
            if re.search(r"[.;:!?¿¡]", between) or _DOC_CONTAINER_OF.search(between):
                continue
            words = [w for w in (re.sub(r"[»\"”')]+$", "", re.sub(r"^[«\"“'(]+", "", x))
                                 for x in re.split(r"[\s,]+", between.lower())) if w]
            if len(words) > 5 or not all(w in _DOC_NP_WORDS or w in _DOC_PRE_MODS or re.fullmatch(r"\d+(?:-" + _DL + r"+)?", w)
                                         for w in words):
                continue
            if any(w == "favor" and (i == 0 or words[i - 1] != "por") for i, w in enumerate(words)):
                continue                     # "Hazme UN FAVOR, la propuesta…"
            if any(w == "nuevo" and i > 0 and words[i - 1] == "de" for i, w in enumerate(words)):
                continue                     # "Mándame DE NUEVO la propuesta": again, not new
            indefinite = any(w in _DOC_INDEFINITE for w in words)
            definite = any(w in _DOC_DEFINITE for w in words)
            if (_DOC_WEAK_VERB.match(re.sub(r"\s+", " ", v.group(0))) and ((not indefinite and definite) or not outs)
                    and not (outs and h[1] < 99 and h[2] in _DOC_KIT_KINDS)):
                continue                     # a weak verb: a NEW document with its format said (or a kit document)
            if (h[2] not in ("answer", "prep", "other") and re.search(r"s$", h[4], re.I) and not indefinite
                    and not definite and not any(_DOC_NUMERAL.match(w) for w in words)):
                continue                     # a bare plural: a duty ("create reports"), never a request
            return True
        return False
    kind_head = re.compile(r"\s+" + _DOC_KIND_ANY + _DNA, re.I)

    def headed(h) -> bool:
        rest = t[h[3]:]
        for _i in range(3):
            rest = _DOC_FMT_LEAD.sub("", rest, 1)
            k = kind_head.match(rest)
            if k:
                rest = rest[k.end():]
        if _DOC_AFTER_REFUSE.match(rest) or _DOC_AFTER_REFUSE.match(_DOC_FMT_AFTER_KIND.sub("", rest)):
            return False
        if _DOC_AFTER_PHRASE.match(rest):
            return True                      # VJH: "una evaluación REQUISITO POR REQUISITO de la vacante…"
        wm = re.match(r"\s*(\S+)", rest)
        w = wm.group(1) if wm else ""
        if not w or (re.match(r"[^\w]|_", w) and not re.match(r"[«\"“'(]" + _DL, w)):
            return True                      # the end, punctuation, a sign
        word = re.sub(r"(?:[^\w'’-]|_)+$", "", re.sub(r"^[«\"“'(]+", "", w))
        if word[:1].isupper():
            return True                      # a name
        low = word.lower()
        return low in _DOC_AFTER_OK or (low in _DOC_NP_WORDS and not re.fullmatch(
            r"p[aá]gina|page|hoja|formato|format|file|archivo|documento|doc", low))
    hit = next((h for h in hits if governed(h) and headed(h)), None)
    if not hit:
        return None
    at, order, kind, end, word, kword = hit
    tail = t[end:]
    after = tail
    for _i in range(3):
        after = _DOC_FOR_LEAD.sub("", _DOC_FMT_LEAD.sub("", after, 1), 1)
    if _DOC_EXISTS_AFTER.match(after) or _DOC_EDIT_AFTER.match(after):
        return None
    # "Write a letter? Nah, a WhatsApp is enough": a "?" right after the document asks, unless an asking form opens it
    if re.match(r"\s*\?", after) and not _DOC_POLITE.match(_doc_strip_lead(t)):
        return None
    cut = re.search(r"[:.;!?]|\s(?:about|sobre|con|with|including|incluyendo|о|об|про)\s", tail, re.I)
    if _DOC_META.search(t[:end] + (tail[:cut.start()] if cut else tail)):
        return None
    said = " ".join(x for _a, _b, x in outs)
    formats = []
    if re.search(_DNB + _DOC_PDF, said, re.I):
        formats.append("pdf")
    if re.search(_DNB + _DOC_WORD + _DNA, said, re.I):
        formats.append("docx")
    if re.search(_DNB + _DOC_HTML + _DNA, said, re.I):
        formats.append("html")
    fmt_said = bool(formats)
    if not fmt_said:
        formats.append("html" if kind in _DOC_HTML_KINDS else "pdf")
    scope = whole if len(whole) <= 400 else t
    subject = _doc_subject(t, end)
    rest = _DOC_VERB.sub(" ", re.sub(_DNB + _DOC_KIND_ANY + _DNA, " ", _doc_without_outputs(scope), flags=re.I))
    substance = len([w for w in re.findall(_DL + "{3,}", rest) if w.lower() not in _DOC_FILLER])
    own_named = any(rx.search(scope) for rx in _DOC_OWN_NAMES)
    # "Make the report" and nothing else — no format, no company, no content — points at a thing that exists
    before = t[:at].split()
    article = (before[-1] if before else "").lower()
    if not fmt_said and article in _DOC_DEFINITE and substance < 2 and not subject and not own_named:
        return None
    np_ = [w.lower() for w in before[-4:]]
    # "Make the (sales) report for the API work", "Haz la cotización del bot": code talk, not a document to write
    if ((not fmt_said or order == 99) and any(w in _DOC_DEFINITE for w in np_)
            and not any(w in _DOC_INDEFINITE for w in np_) and _DOC_CODE_TALK.search(t)):
        return None
    return {"kind": kind, "formats": tuple(formats), "fmt_said": fmt_said, "word": word, "kword": kword,
            "subject": subject}


def _doc_task(text: str) -> Optional[dict]:
    """CTO documentTask: the request line — the whole message, or its first or last line (greeting / thanks lines
    aside; the last only when the first is not her question, it is no posting line and nothing before it introduces
    someone else's text) → {…, "line"} or None."""
    raw = (text or "").strip()
    if not raw or len(raw) > 4000 or raw.startswith("/"):
        return None
    lines = [ln.strip() for ln in re.split(r"\r?\n", raw) if ln.strip() and not _DOC_GREETING_LINE.match(ln.strip())]
    if not lines:
        return None
    first, last = lines[0], lines[-1]
    pasted = any(_DOC_INTRO_OTHER.search(ln) for ln in lines[:-1]) or bool(re.match(r"[\"“«'‘]", last))
    areas = [first] + ([last] if len(lines) > 1 and not re.search(r"\?\s*$", first) and not _DOC_POSTING_LINE.search(last)
                       and not pasted else [])
    whole = re.sub(r"\s+", " ", raw)
    for area in areas:
        line = re.sub(r"\s+", " ", area)
        r = _doc_request_in(line, whole)
        if r:
            return {**r, "line": line}
    return None


def _doc_bare_tokens(folded: str) -> List[str]:
    """The words of a folded message for the bare test ("es urgente", "esta respuesta" carry no task)."""
    s = re.sub(r"\b(?:es|it's|its|it\s+is)\s+(?:urgente|urgent)\b", " ", folded)
    s = re.sub(r"\bpor\s+favor\b", "porfa", s)
    s = re.sub(r"\b(?:esta|esa)\s+(respuesta)\b", r"\1", s)
    return re.findall(r"[a-z0-9ñ']+", s)


def _doc_quote_client(raw: str) -> bool:
    """English "quote" is a client quote only with a client after "for": a name ("for Café Luna") or a business ("for
    the gym") — never "a quote from Steve Jobs", "a quote about resilience" or "a quote for my LinkedIn post"."""
    m = re.search(r"\b(?:quotes?|quotations?)\b(?P<t>.*)$", raw or "", re.I | re.S)
    if not m:
        return False
    return bool(re.match(
        r"[^.?!]{0,40}?\bfor\s+(?:the\s+|a\s+|an\s+|my\s+client\s+|our\s+client\s+)?(?:[A-ZÁÉÍÓÚÑ0-9]|"
        r"(?:dental\s+|vet\s+|beauty\s+|yoga\s+|law\s+)?(?:clinics?|gym|restaurant|hotel|caf[eé]|bakery|studio|spa|"
        r"shop|store|school|academy|salon|office|firm|agency|client|customer|dentist|doctor|practice|business|startup|"
        r"pharmacy|hostel|bar|boutique|realtor|real\s+estate)\b)", m.group("t")))


_DOC_ONEPAGER_WORD = re.compile(r"one[\s-]?pager|resumen\s+ejecutivo|executive\s+summary|ficha", re.I)


def _doc_request(text: str) -> Optional[dict]:
    """{formats, kind, sentence, raw, bare} when her message asks for a file (CTO's rule, above); else None. kind is
    VJH's word (quote, proposal, evaluation, letter, prep, presentation, promo, analysis, onepager, report, answer), or
    None when the object is only a format ("hazme un PDF")."""
    if not text or not text.strip():
        return None
    folded = text.lower().translate(_DOC_FOLD)
    if _DOC_TO_ELENA.match(folded):
        return None                           # someone else's message to her ("Hi Elena, … send me a PDF…")
    toks = _doc_bare_tokens(folded)
    # fix 8 (held-out p8): "Make a copy OF THE Word file" — a format owned by "of the / del / de la" is a file that
    # exists, never "make it a file" ("ponlo en el Word", "pásalo al PDF" still are)
    if (toks and len(toks) <= 12 and any(t in _DOC_FMT_TOKENS for t in toks) and not _DOC_LABEL_RE.match(folded)
            and all(t in _DOC_BARE_WORDS or t in _DOC_FMT_TOKENS for t in toks)
            and not _DOC_BARE_OWNED.search(folded)):
        fmts = tuple(dict.fromkeys(_DOC_FMT_TOKENS[t] for t in toks if t in _DOC_FMT_TOKENS))
        return {"formats": fmts, "kind": None, "sentence": text.strip().lower(), "raw": text.strip(), "bare": True}
    task = _doc_task(text)
    if not task:
        return None
    line = task["line"]
    kind = task["kind"]
    if _DOC_OR_ONLY.search(line):
        return None                           # "…, or only plain text?": what the bot can do, not a file
    if kind == "quote" and re.fullmatch(r"quotes?|quotations?", task["kword"], re.I) and not task["fmt_said"] and not (
            _DOC_HER_SERVICE.search(line.lower().translate(_DOC_FOLD)) or _doc_quote_client(line)):
        return None                           # "Write a quote for my LinkedIn post": a citation
    if kind == "cv":
        kind = "answer"                       # the CV refusal (_DOC_CV_ASK_RE): the kit's CV lives on the deal
    if kind == "other":                       # "hazme un PDF con…": the format alone is the object → None
        kind = None if not task["kword"] else "onepager" if _DOC_ONEPAGER_WORD.search(task["kword"]) else "answer"
    return {"formats": task["formats"], "kind": kind, "sentence": line.lower(), "raw": line, "bare": False}


def _document_formats(text: str) -> tuple:
    """('pdf',) / ('docx',) / ('html',) / ('pdf', 'docx') … when a typed or spoken message asks for a file; () when
    not."""
    req = _doc_request(text)
    return req["formats"] if req else ()


def _asks_for_document(text: str) -> bool:
    return bool(_doc_request(text))


def _doc_fold(text: str) -> str:
    return re.sub(r"\s+", " ", (text or "").lower().translate(_DOC_FOLD)).strip()


_DOC_COVER_RE = re.compile(r"\b(?:cover|motivation)\s+letter|\bcarta\s+de\s+(?:presentacion|motivacion)")
# 18 Oct review: a CV is the apply kit's posting-shaped one on the deal, never model-written here. Only when the CV is
# the thing she asks to have made ("hazme mi CV en Word", "write my CV for Pitz in PDF", "CV for Pitz as PDF"); a
# request that only mentions one ("a document of the jobs where my CV was sent") is not. "resumen" is a summary.
_DOC_CV_ASK_RE = re.compile(
    r"(?:\b(?:make|write|draft|create|build|prepare|update|tailor|send|give|get|haz|escribe|redacta|crea|prepara|arma|"
    r"actualiza|adapta|envia|manda|pasa|da)(?:me|melo|mela)?\s+(?:me\s+)?|^\s*)"
    r"(?:(?:a|an|my|mi|un|el|the|new|nuevo|tailored|updated|actualizado)\s+)*"
    r"(?:cv|curriculum(?:\s+vitae)?|curriculo|resume|hoja\s+de\s+vida)\b")
# 19 Oct (rule 2): a document about HER FIT for a job is never written here — it is the deal's own kit notes. Decided
# by STRUCTURE, whatever the document's noun (review blocker: "how well I match Pitz", "mis fortalezas frente a lo que
# pide Pitz", "a presentation pitching me to Pitz", "carta … por qué encajo en el puesto" all reached the model):
# a live job-deal company named anywhere (_doc_companies), or her fit words next to a job or a company (_doc_fit).
_DOC_FIT_HER = re.compile(
    r"\b(?:(?:mi|my)\s+(?:encaje|fit|candidatura|candidacy|application|postulacion)|encaj\w*|"
    r"(?:i|i'm|i\s+am|me|my\s+profile|mi\s+perfil|soy)\s+(?:\w+\s+){0,3}?(?:match\w*|right\s+for|(?:a\s+)?(?:good\s+|"
    r"great\s+|strong\s+|perfect\s+)?fit)\b|am\s+i\s+a\s+(?:\w+\s+)?fit|is\s+(?:this|it)\s+a\s+(?:\w+\s+)?fit|"
    r"cumplo|do\s+i\s+meet|i\s+meet|(?:am\s+i|i'm|i\s+am)\s+(?:\w+\s+)?qualified|qualify|candidat\w*|hire\s+me|"
    r"hiring\s+me|pitch(?:ing)?\s+me|sell(?:ing)?\s+me|vendeme|venderme|"
    r"por\s+que\s+(?:yo|a\s+mi|me\s+contrat\w*|contratarme|deberian)|why\s+(?:me|i\s+should|i'm|i\s+am|"
    r"should\s+\w+\s+hire)|(?:mis|my)\s+(?:fortalezas|strengths|debilidades|weaknesses|gaps|brechas))\b")
# what a job asks ("los requisitos", "lo que pide Pitz"): a fit document next to a job, or next to her and a company
_DOC_FIT_REQ = re.compile(
    r"\b(?:requisitos?|requirements?|requerimientos|criterios\s+de\s+la\s+oferta|(?:lo\s+que|what)\s+(?:\w+\s+){0,2}?"
    r"(?:pide|piden|busca|buscan|necesita|necesitan|asks?|wants?|needs?|(?:is|are)\s+looking))\b")
# fit words that speak of a hiring decision by themselves ("¿por qué encajo?", "hire me"): a fit document even with no
# job or company named — then she is asked which job
_DOC_FIT_SELF = re.compile(
    r"\b(?:encaj\w*|cumplo|do\s+i\s+meet|i\s+meet|am\s+i\s+a\s+fit|hire\s+me|hiring\s+me|pitch(?:ing)?\s+me|"
    r"sell(?:ing)?\s+me|vendeme|venderme|(?:mi|my)\s+(?:encaje|fit|candidatura|candidacy)|why\s+me|por\s+que\s+yo)\b")
# her in the first person, without a fit word ("sobre mí", "my experience"): a fit document only next to a job
_DOC_FIT_ME = re.compile(
    r"\b(?:sobre\s+mi|about\s+me|(?:mi|my|mis)\s+(?:perfil|profile|experiencia|experience|cv|background|trayectoria|"
    r"habilidades|skills|logros|achievements)|fortalezas|strengths)\b")
_DOC_FIT_JOB = re.compile(
    r"\b(?:(?:el|la|un|una|the|a|this|that|este|esta|ese|esa)\s+(?:puesto|vacante|cargo|role|position|job|opening|"
    r"vacancy|oferta|posting|trabajo|empleo)|(?:puesto|vacante|cargo|rol|role|position|job|opening)\s+"
    r"(?:en|de|at|with|in|of)|reclutador\w*|recruiters?|hiring\s+managers?|entrevistador\w*|interviewers?)\b")
# a fit document about an interview ("…the Pitz interview questions and my answers"): the prep kit, not the evaluation
_DOC_FIT_INTERVIEW = re.compile(r"\b(?:entrevistas?|interviews?|mock|preguntas|questions)\b")


def _doc_kind_asked(text: str) -> str:
    """The kind of the document noun she used first (quote, proposal, evaluation, letter, prep, presentation, promo,
    analysis, onepager, report), else answer."""
    low, best = _doc_fold(text), ("answer", 10 ** 9)
    for kind, pat in _DOC_KIND_NOUNS:
        if kind == "answer":
            continue
        m = re.search(r"\b(?:" + pat + r")(?![\w-])", low)
        if m and m.start() < best[1]:
            best = (kind, m.start())
    return best[0]


def _doc_two_kinds(sentence: str) -> bool:
    """Two documents asked in one message ("la carta para Pitz en PDF y la cotización para El Faro en Word")."""
    low = _doc_fold(sentence)
    hits = [(m.start(), m.end(), k) for k, p in _DOC_KIND_NOUNS if k not in ("answer", "report")
            for m in re.finditer(r"\b(?:" + p + r")(?![\w-])", low)]
    hits.sort()
    for (a1, e1, k1), (a2, _e2, k2) in zip(hits, hits[1:]):
        between = low[e1:a2]
        # "…la cotización con el diagnóstico y el análisis profundo": the second noun is what the first one holds
        if re.search(r"\b(?:con|with|del|de\s+(?:la|las|los|el)|of\s+the|sobre|about|incluyendo|including|que\s+incluya|"
                     r"that\s+includes)\b", between):
            continue
        if k1 != k2 and re.search(r"\b(?:y|and|e|plus|tambien|also)\s+(?:el|la|los|las|un|una|the|a|an|my|mi)\s+\S*$",
                                  between + " "):
            return True
    return False


def _doc_lang_asked(text: str) -> Optional[str]:
    """'en' / 'es' only when she names the document's language ("in English", "en español")."""
    m = re.search(r"\b(?:in|en)\s+(english|ingles|spanish|espa[nñ]ol|castellano)\b", _doc_fold(text))
    return None if not m else ("en" if m.group(1) in ("english", "ingles") else "es")


# Whisper verbose_json names the language in English words.
_WHISPER_LANG = {"spanish": "es", "english": "en", "russian": "ru", "es": "es", "en": "en", "ru": "ru"}

# Screenshot reader (VJH-3): the vision functions of scripts/judge_feedback_sync.py, given this prompt.
_SCREENSHOT_PROMPT = (
    "On the FIRST line write exactly one of these labels: TYPE: job posting | TYPE: interviewer question | "
    "TYPE: recruiter or employer message | TYPE: other "
    "(interviewer question = a question she must answer: an interview, screening or application-form question). "
    "Then transcribe ALL the text visible in this screenshot, exactly as written and in its original language, "
    "keeping the line breaks. Do not summarise, translate or add anything. On the very last line write exactly: "
    "END OF TEXT")
_SCREENSHOT_KINDS = ("job posting", "interviewer question", "recruiter or employer message", "other")
_SCREENSHOT_MAX_BYTES = 6_000_000     # the same cap as judge_feedback_sync._file_bytes
_JUDGE_SYNC = None


def _judge_sync_module():
    """scripts/judge_feedback_sync.py, loaded once (its main() is guarded; the top level is constants only)."""
    global _JUDGE_SYNC
    if _JUDGE_SYNC is None:
        import importlib.util
        from pathlib import Path
        path = Path(__file__).resolve().parents[2] / "scripts" / "judge_feedback_sync.py"
        spec = importlib.util.spec_from_file_location("vjh_judge_feedback_sync", path)
        mod = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(mod)
        _JUDGE_SYNC = mod
    return _JUDGE_SYNC


# ═══════════════════════════════════════════════════════════════════════════
# One-tap job-alert buttons (18 Oct 2026, Elena: "⏳ I applied" / "❌ Not a fit").
# telegram_notifier.add_alert_buttons puts them under a "📋 Apply yourself" / "⚠️ UNVERIFIED" alert, carrying the
# HubSpot deal id. HubSpot is reached the way this bot already reaches it for screenshots: judge_feedback_sync's
# _hubspot_key + _hs, the key and call the hourly learner uses. No new endpoint, service or dependency.
# ⏳ writes "Applied (via VJH alert button, <date>)" (the learner reads "Applied" on a ⏳ Sent deal as HER
# application) and moves the deal to ⏳ Sent. ❌ asks for one line and writes HER words, nothing else, as the note,
# so judge_feedback_sync._rejection_reason reads it as "her reason"; then the deal goes to ❌ No fit.
# A stage moves only out of an open column, read from HubSpot at tap time: a second tap changes nothing.
# ═══════════════════════════════════════════════════════════════════════════
_ALERT_SENT, _ALERT_LOST = "decisionmakerboughtin", "closedlost"
_ALERT_OPEN = ("appointmentscheduled", "qualifiedtobuy", "presentationscheduled")
_ALERT_LABEL = {"appointmentscheduled": "🤖 AI working", "qualifiedtobuy": "🔥 I act today",
                "presentationscheduled": "⚡ I act this week", "decisionmakerboughtin": "⏳ Sent",
                "contractsent": "💬 They replied", "closedwon": "✅ Won", "closedlost": "❌ No fit"}
_NOFIT_PLAIN_WINDOW_S = 15 * 60       # her next plain message is the reason for 15 minutes
_NOFIT_REPLY_WINDOW_S = 24 * 3600     # a reply to the question itself counts for a day
_NOFIT_CANCEL = ("cancel", "cancelar")
_NOFIT_MAX = 1500


def _alert_hs(path: str, body=None) -> dict:
    jfs = _judge_sync_module()
    key = jfs._hubspot_key()
    if not key:
        raise RuntimeError("no HubSpot key")
    return jfs._hs(key, path, body, timeout=20)


_LEARNER_NOTE_CAP = 8                 # judge_feedback_sync._fetch_notes_batch reads `to[:8]` notes per deal


def _alert_deal(deal: str) -> tuple:
    """(stage, deal name, notes on the deal) read from HubSpot NOW, never from the alert, which may be days old.
    The note count rides on the same GET: the ❌ question must not promise "VJH learns from it" on a deal whose
    notes already fill the learner's per-deal cap."""
    res = _alert_hs(f"/crm/v3/objects/deals/{deal}?properties=dealstage,dealname&associations=notes")
    props = res.get("properties") or {}
    notes = {str(r.get("id")) for r in (((res.get("associations") or {}).get("notes") or {}).get("results") or [])
             if isinstance(r, dict) and r.get("id")}
    return props.get("dealstage") or "", re.sub(r"^\[[^\]]*\]\s*", "", props.get("dealname") or ""), len(notes)


def _alert_write(deal: str, note: str, stage: Optional[str]) -> str:
    """The note first, then the stage, so a deal in ❌ No fit always carries her reason.
    'ok' | 'note_only' (HubSpot refused the stage move); raises when the note itself failed (nothing written)."""
    from datetime import timezone as _tz
    _alert_hs("/crm/v3/objects/notes", {
        "properties": {"hs_note_body": note,
                       "hs_timestamp": datetime.now(_tz.utc).strftime("%Y-%m-%dT%H:%M:%S.000Z")},
        "associations": [{"to": {"id": deal},
                          "types": [{"associationCategory": "HUBSPOT_DEFINED", "associationTypeId": 214}]}]})
    if not stage:
        return "ok"
    try:
        res = _alert_hs("/crm/v3/objects/deals/batch/update",
                        {"inputs": [{"id": deal, "properties": {"dealstage": stage}}]})
    except Exception as e:
        logger.warning(f"🔘 alert button: note saved, stage move refused ({str(e)[:120]})")
        return "note_only"
    if isinstance(res, dict) and res.get("errors"):          # a 207 multi-status is not an HTTP error
        logger.warning(f"🔘 alert button: note saved, stage move refused ({str(res['errors'])[:120]})")
        return "note_only"
    return "ok"


def _panama_today() -> str:
    from datetime import timezone as _tz
    return datetime.now(_tz(timedelta(hours=-5))).strftime("%Y-%m-%d")       # Panama has no DST


def _hs_problem(e: Exception, nothing: str = "Nothing changed.", retry: str = "Try again in a minute.") -> str:
    if getattr(e, "code", None) == 404:
        return f"⚠️ This deal is no longer in HubSpot. {nothing}"
    return f"⚠️ HubSpot could not be reached or refused the change. {nothing} {retry}"


def _image_mime(raw: bytes) -> str:
    """The image signature check of judge_feedback_sync._file_bytes, with the mime taken from the magic bytes."""
    head = bytes(raw[:12]) if raw else b""
    if head.startswith(b"\xff\xd8"):
        return "image/jpeg"
    if head.startswith(b"\x89PNG"):
        return "image/png"
    if head.startswith(b"RIFF") and head[8:12] == b"WEBP":
        return "image/webp"
    if head.startswith(b"GIF8"):
        return "image/gif"
    return ""


def _screenshot_kind(label: str) -> str:
    """The vision model's TYPE label → one of _SCREENSHOT_KINDS, by stem ('TYPE: interview question', 'TYPE: posting',
    'TYPE: employer message' all count). '' when the label names no kind at all (not a TYPE label)."""
    label = (label or "").lower()
    for stems, kind in ((("interview",), "interviewer question"),
                        (("recruit", "employer"), "recruiter or employer message"),
                        (("posting", "job", "vacan", "oferta"), "job posting"),
                        (("question", "screening", "pregunta"), "interviewer question"),
                        (("message", "mensaje"), "recruiter or employer message"),
                        (("other", "otro"), "other")):
        if any(s in label for s in stems):
            return kind
    return ""


def _split_screenshot_read(txt: str):
    """(kind, text, cut) from the vision reply. The TYPE line is asked for FIRST, so a long posting cut off at the
    token limit keeps its type; a TYPE line at the end is accepted too. cut = the 'END OF TEXT' line is missing and
    no TYPE line closes the read."""
    lines = (txt or "").strip().splitlines()
    ended = bool(lines) and re.match(r"^\W*END OF TEXT\W*$", lines[-1].strip(), re.IGNORECASE) is not None
    if ended:
        lines = lines[:-1]
    n = len(lines)
    for i in list(range(min(3, n))) + list(range(n - 1, max(n - 4, 2), -1)):
        m = re.match(r"^\W*TYPE\s*:\s*(.+?)\W*$", lines[i], re.IGNORECASE)
        kind = _screenshot_kind(m.group(1)) if m else ""
        if kind:
            body = "\n".join(lines[:i] + lines[i + 1:]).strip()
            return kind, body, (not ended and i < n - 1)
    return "other", "\n".join(lines).strip(), not ended


# ═══════════════════════════════════════════════════════════════════════════
# Documents (16 Oct 2026, Elena: "like Chalán in Pitz"): a PDF or Word .docx is read from its TEXT LAYER with the
# libraries already in the venv (PyPDF2 3.0.1, python-docx 1.2.0), in a thread, and answered like a screenshot. A scan
# with no text layer, a password, a file over 20 MB or another type gets an honest line, never a guess.
# ═══════════════════════════════════════════════════════════════════════════
_DOC_MAX_BYTES = 20_000_000       # Telegram's getFile limit for bots
_DOC_MAX_CHARS = 15_000           # text kept for the answer; the rest is cut and the reply says so
_DOC_MAX_PAGES = 40               # PDF pages read at most
_DOC_MIN_TEXT = 30                # fewer readable characters than this = no text layer (a scan, photos)
_DOCX_MAX_UNZIPPED = 200_000_000  # a .docx's parts together, uncompressed: python-docx loads every part into memory
_PDF_MIME = "application/pdf"
_DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
_MC_FALLBACK = "{http://schemas.openxmlformats.org/markup-compatibility/2006}Fallback"


def _document_gate(filename: str, mime: str) -> str:
    """Before downloading: 'read' (a PDF or .docx by name or mime, or a file with no extension: its bytes decide),
    'doc' (an old Word .doc) or 'other' (another type by name, e.g. .xlsx)."""
    ext = os.path.splitext(filename or "")[1].lower()
    mime = (mime or "").lower()
    if ext in (".pdf", ".docx") or mime in (_PDF_MIME, _DOCX_MIME):
        return "read"
    if ext == ".doc" or mime == "application/msword":
        return "doc"
    return "other" if ext else "read"


def _document_format(raw: bytes) -> str:
    """'pdf' | 'docx' | 'ole' (an old Word .doc, or any password-protected Office file) | '' — from the bytes, never
    from the name or the mime type the phone attached."""
    head = bytes(raw[:1024]) if raw else b""
    if b"%PDF-" in head:
        return "pdf"
    if head.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"):
        return "ole"
    if head.startswith(b"PK\x03\x04"):
        import io
        import zipfile
        try:
            with zipfile.ZipFile(io.BytesIO(raw)) as z:
                return "docx" if "word/document.xml" in z.namelist() else ""
        except Exception:
            return ""
    return ""


def _pdf_text(raw: bytes) -> dict:
    """PyPDF2: the text of the first _DOC_MAX_PAGES pages, stopping once _DOC_MAX_CHARS are read. A PDF that opens
    with an empty password (copy-restricted only) is read; any other encryption is 'locked'. AES needs PyCryptodome,
    which Oracle's venv does not have: PyPDF2 then raises DependencyError — when opening an AES-256 file, or on every
    page of an AES-128 one — and that is 'locked' too, never 'damaged' or 'a scan'. A page whose extraction raised is
    counted apart from a page with no text ('failed_pages'), so a reader error is never reported as a scan."""
    from io import BytesIO
    from PyPDF2 import PdfReader
    try:
        from PyPDF2.errors import DependencyError
    except ImportError:                   # an older PyPDF2 without it: nothing of that type is ever raised
        class DependencyError(Exception):
            pass
    try:
        reader = PdfReader(BytesIO(raw))
    except DependencyError:               # AES-256 and no crypto library in the venv
        return {"status": "locked"}
    if reader.is_encrypted:
        try:
            opened = reader.decrypt("")
        except Exception:                 # e.g. AES with no crypto library in the venv
            opened = 0
        if not opened:
            return {"status": "locked"}
    total = len(reader.pages)
    parts, chars, read, empty, failed = [], 0, 0, 0, 0
    for i in range(min(total, _DOC_MAX_PAGES)):
        read += 1
        try:
            t = (reader.pages[i].extract_text() or "").strip()
        except DependencyError:           # AES-128 opened with the empty password, but its text needs PyCryptodome
            return {"status": "locked"}
        except Exception as e:
            logger.info(f"📄 PDF page {i + 1} text could not be extracted: {str(e)[:120]}")
            failed += 1
            continue
        if len(re.sub(r"\s+", "", t)) < 3:
            empty += 1
            continue
        parts.append(t)
        chars += len(t)
        if chars > _DOC_MAX_CHARS:
            break
    return {"status": "ok", "text": "\n\n".join(parts), "pages": total, "pages_read": read, "empty_pages": empty,
            "failed_pages": failed}


def _docx_text(raw: bytes) -> dict:
    """python-docx: every paragraph in reading order — body, tables, content controls, text boxes and tracked
    insertions (deleted text and a text box's duplicate fallback copy are skipped). Headers and footers are left out."""
    import io
    import zipfile
    with zipfile.ZipFile(io.BytesIO(raw)) as z:
        # python-docx loads EVERY part into memory, so the whole package is capped, not only the text part
        if (sum(i.file_size for i in z.infolist()) > _DOCX_MAX_UNZIPPED
                or z.getinfo("word/document.xml").file_size > 50_000_000):
            raise ValueError("the .docx is too large uncompressed (a zip bomb, or huge embedded media)")
    import docx
    body = docx.Document(io.BytesIO(raw)).element.body
    lines, chars = [], 0
    for p in body.iter(_W + "p"):
        if next(p.iterancestors(_MC_FALLBACK), None) is not None:
            continue
        bits = []
        for el in p.iter(_W + "t", _W + "tab", _W + "br", _W + "cr"):
            # only this paragraph's own runs: not a tab stop in its properties, not a text box's paragraph inside it
            if el.getparent().tag != _W + "r" or next(el.iterancestors(_W + "p"), None) is not p:
                continue
            bits.append((el.text or "") if el.tag == _W + "t" else ("\t" if el.tag == _W + "tab" else "\n"))
        line = "".join(bits).strip()
        if line:
            lines.append(line)
            chars += len(line) + 1
            if chars > _DOC_MAX_CHARS:
                break
    return {"status": "ok", "text": "\n".join(lines)}


def _read_document(raw: bytes) -> dict:
    """{status, fmt, text, cut, cut_at, pages, pages_read, empty_pages, failed_pages}. status: ok | no_text | locked |
    unreadable | no_reader (PyPDF2 / python-docx missing here) | doc (an old .doc or a password-protected Word file) |
    other (not a PDF or Word file). Blocking: run it in a thread."""
    fmt = _document_format(raw)
    res = {"status": "ok", "fmt": fmt, "text": "", "cut": False, "cut_at": "", "pages": 0, "pages_read": 0,
           "empty_pages": 0, "failed_pages": 0}
    if fmt not in ("pdf", "docx"):
        res["status"] = "doc" if fmt == "ole" else "other"
        return res
    try:
        res.update(_pdf_text(raw) if fmt == "pdf" else _docx_text(raw))
    except ImportError as e:                       # the reader library is missing here: never blame the file
        logger.warning(f"📄 No {fmt} reader installed: {e}")
        res["status"] = "no_reader"
    except Exception as e:
        logger.warning(f"📄 {fmt} could not be opened: {str(e)[:160]}")
        res["status"] = "unreadable"
    if res["status"] != "ok":
        return res
    text = (res["text"] or "").strip()
    if len(re.sub(r"\s+", "", text)) < _DOC_MIN_TEXT:
        # pages the reader failed on are not "a scan": with no text read, that is an unreadable file
        res.update(status="unreadable" if res.get("failed_pages") else "no_text", text="")
        return res
    if len(text) > _DOC_MAX_CHARS:
        text = text[:_DOC_MAX_CHARS]
        space = text.rfind(" ")
        text, res["cut_at"] = (text[:space] if space > _DOC_MAX_CHARS - 200 else text), "chars"
    elif res["pages_read"] < res["pages"]:
        res["cut_at"] = "pages"
    res.update(text=text, cut=bool(res["cut_at"]))
    return res


def _document_problem_reply(status: str, label: str, label_es: str, ext: str = "", pages: int = 0,
                            pages_read: int = 0) -> str:
    """The honest line for a document the bot did not read: what happened and what to send instead. It never says
    anything about the content."""
    if status == "no_text":
        if pages and pages > pages_read:
            where = f"the first {pages_read} of its {pages} pages have"
        elif pages:
            where = f"its {pages} page has" if pages == 1 else f"its {pages} pages have"
        else:
            where = "it has"
        return (f"📄 I opened {label}, but {where} no text I can read: it looks like a scan or photos saved as a "
                "document. I will not guess what it says. Send screenshots of the pages as photos and I will read "
                f"them. / Abrí {label_es}, pero no tiene texto que pueda leer: parece un escaneo o fotos guardadas "
                "como documento. No voy a adivinar su contenido. Envía capturas de las páginas como fotos y las leo.")
    # a file NAMED .pdf/.docx whose bytes are not one must not be called "a .pdf file"
    kind = f"a {ext} file" if ext and ext not in (".pdf", ".docx") else "not a real PDF or Word file"
    # any old Office binary (or a password-protected Office file) looks the same: say Word only when named Word
    word = ext in (".doc", ".docx")
    ole_en = ("an old Word .doc file or a password-protected Word file" if word else
              "an old Office file (.doc, .xls, .ppt) or a password-protected Office file")
    ole_es = "un Word .doc antiguo o un Word con contraseña" if word else         "un archivo de Office antiguo (.doc, .xls, .ppt) o un archivo de Office con contraseña"
    return {
        "too_big": (f"📄 {label} is over 20 MB. Telegram does not let a bot download files that big, so I did not "
                    "read it. Send a smaller file or screenshots of the pages that matter. / "
                    f"{label_es} pesa más de 20 MB: Telegram no deja que un bot descargue archivos tan grandes, así "
                    "que no lo leí. Envía uno más pequeño o capturas de las páginas importantes."),
        "doc": (f"📄 {label} is {ole_en}. I read only PDF and Word "
                ".docx without a password, so I did not read it and I will not guess what it says. Save it as .docx "
                f"or PDF and send it again. / {label_es} es {ole_es}. Solo leo "
                "PDF y Word .docx sin contraseña, así que no lo leí y no voy a adivinar su contenido. Guárdalo como "
                ".docx o PDF y envíalo otra vez."),
        "other": (f"📄 I read PDF and Word .docx documents, screenshots and photos. {label} is {kind}, so I "
                  "did not read it and I will not guess what it says. Send it as a PDF or .docx, or as a screenshot. "
                  f"/ Leo documentos PDF y Word .docx, capturas y fotos. {label_es} es otro tipo de archivo, así que "
                  "no lo leí y no voy a adivinar su contenido. Envíalo como PDF o .docx, o como captura."),
        "locked": (f"📄 {label} is password-protected or encrypted, so I could not open it and I will not guess what "
                   "it says. Send a copy without the password, or screenshots of the pages. / "
                   f"{label_es} tiene contraseña o está cifrado: no pude abrirlo y no voy a adivinar su contenido. "
                   "Envía una copia sin contraseña o capturas de las páginas."),
        "timeout": (f"📄 Reading {label} took too long, so I stopped. I will not guess what it says. Send a smaller "
                    "file or screenshots of the pages that matter. / "
                    f"Leer {label_es} tardó demasiado y paré. No voy a adivinar su contenido. Envía un archivo más "
                    "pequeño o capturas de las páginas importantes."),
        "download": (f"📄 I could not download {label}. Please send it again. / "
                     f"No pude descargar {label_es}. Envíalo otra vez."),
        "no_reader": (f"📄 I cannot read {label} right now: the reader for this file type is missing on the server. "
                      "The file is not the problem, and I will not guess what it says. Send it as a PDF or as "
                      f"screenshots. / No puedo leer {label_es} ahora: falta el lector de este tipo de archivo en el "
                      "servidor. El archivo no tiene la culpa y no voy a adivinar su contenido. Envíalo como PDF o "
                      "como capturas."),
    }.get(status, (f"📄 I could not read {label}: the file looks damaged, or it is not a real PDF or Word file. I "
                   "will not guess what it says. Export it again as PDF, or send screenshots of the pages. / "
                   f"No pude leer {label_es}: el archivo parece dañado o no es un PDF o Word real. No voy a adivinar "
                   "su contenido. Expórtalo otra vez como PDF o envía capturas de las páginas."))


def _document_notes(res: dict, label: str) -> str:
    """Notes the bot adds to the answer itself, so a cut or a scanned page is never left to the model to mention."""
    notes = []
    if res.get("cut"):
        if res.get("cut_at") == "pages":
            part = f"the first {res['pages_read']} of its {res['pages']} pages"
        elif res.get("pages", 0) > res.get("pages_read", 0):
            part = f"the first {_DOC_MAX_CHARS:,} characters, from {res['pages_read']} of its {res['pages']} pages"
        else:
            part = f"the first {_DOC_MAX_CHARS:,} characters"
        notes.append(f"(📄 I read only the first part of {label}: {part}. The rest is not in this answer. / "
                     "Leí solo la primera parte del documento; el resto no está en esta respuesta.)")
    if res.get("empty_pages"):
        verb = "has" if res["empty_pages"] == 1 else "have"
        notes.append(f"(📄 {res['empty_pages']} of the {res['pages_read']} pages I read {verb} no readable text (blank, "
                     "or scans or images), so what they show is not in this answer. Send them as screenshots if they "
                     f"matter. / {res['empty_pages']} de las páginas leídas no tienen texto legible (en blanco, o "
                     "escaneos); envíalas como capturas si importan.)")
    if res.get("failed_pages"):
        verb = "was" if res["failed_pages"] == 1 else "were"
        notes.append(f"(📄 {res['failed_pages']} of the {res['pages_read']} pages I read {verb} not readable: the PDF "
                     "reader failed on them, so what they show is not in this answer. Send them as screenshots if "
                     f"they matter. / {res['failed_pages']} de las páginas leídas no se pudieron leer (falló el lector "
                     "de PDF); envíalas como capturas si importan.)")
    return ("\n\n" + "\n\n".join(notes)) if notes else ""


# ═══════════════════════════════════════════════════════════════════════════
# THE FILE ITSELF (18 Oct 2026). Drawn by cto-aipa's shared renderer, run in place (~/cto-aipa/scripts/
# build-lane-cv.cjs --text=<spec> --format=pdf|docx): one renderer for every bot, with its checks — a character the
# font cannot draw stops it, words in == words drawn (PDF) / written and read back by python-docx (DOCX), at most 20
# pages, bracketed placeholders listed. Word goes through the renderer's own docx mode too (python-docx in this venv,
# run by the renderer with -I), so a .docx gets exactly the same parsing and checks as the PDF instead of a second
# parser here. Blocking: callers run these in a thread. A temp folder per file, removed before returning.
# "letter for <company>" reads the deal's own ✅ letter note with hs-deal-prep.cjs (read-only, never rewritten).
# Nothing is uploaded: VJH has no file-upload path to HubSpot (crm_hub → /api/crm-event carries no file). cto-aipa
# has one (scripts/hs-files.cjs uploadOutreachFile + addNoteAttachments, `files` scope; hs-attach-deal-files.cjs
# --deal=<id>), but calling it from VJH would be a new HubSpot write path: only with Elena's go-ahead.
# ═══════════════════════════════════════════════════════════════════════════
_DOC_RENDER_TIMEOUT_S = 90
_DOC_PREP_TIMEOUT_S = 60
_DOC_FILE_MAX = 45_000_000            # under Telegram's 50 MB for a bot's document
_DOC_LAST_WINDOW_S = 6 * 3600         # a bare "PDF" turns the answer of the last 6 hours into a file
_DOC_READY_MARK = "✅ READY TO SEND — cover letter"      # cto-aipa hs-fill-apply-kit.cjs READY_MARK
_DOC_ACTING = ("qualifiedtobuy", "presentationscheduled")
# Marks that carry a verdict. The renderer refuses only some of them and drops the rest as emoji with ok:true
# (★★★★☆ printed as nothing, 18 Oct review): VJH refuses the whole set before rendering.
_DOC_VERDICT_MARKS = re.compile(
    "[☐-☒✅✓-✘❌❎⛔⭕★☆⭐➕➖⬆⬇⬛⬜"
    "\U0001F44D\U0001F44E\U0001F534\U0001F535\U0001F53A\U0001F53B\U0001F4C8\U0001F4C9\U0001F5F4-\U0001F5F8\U0001F6AB"
    "\U0001F7E0-\U0001F7EB]")
_DOC_PREP_HEAD = re.compile(
    "^===== (DEAL: .*|📌 JOB POSTING|🔎 COMPANY BRIEF|🎯 ROLE DEFENSE.*|🛡️? TECHNICAL DEFENSE|✉️? LETTER|"
    "📝 NOTE|🎤 THIS COMPANY'S LIKELY QUESTIONS.*|✅ VERIFIED EXAMPLES.*|✅ EJEMPLOS VERIFICADOS.*|NO DEEP-DIVE.*|"
    "READ-PROOF.*)$")
_DOC_NOTE_LABELS = ("📌 JOB POSTING", "🔎 COMPANY BRIEF", "✉", "🎯 ROLE DEFENSE", "🛡", "📝 NOTE")
_DOC_QUESTIONS_LABEL = "🎤 THIS COMPANY'S LIKELY QUESTIONS"
_DOC_ROLE_STOP = _DOC_BARE_WORDS | frozenset("""cover letter carta presentacion para for pdf word docx html deal job role
puesto trabajo empresa company compania hiring vjh lead la los las the de del en with con about sobre at al hazme haz make
create write redacta redactame prepara preparame genera generame dame send me enviame mandame quiero necesito want need
english ingles spanish espanol castellano por favor please evaluacion evaluation assessment scorecard defensa defense
defence prep preparacion entrevista interview sheet cheat questions preguntas mock requisito requisitos requirement
requirements encaje fit vacante vacancy opening posting oferta analisis analysis informe report resumen summary
one pager propuesta proposal cotizacion quote presentation promo cumplo meet qualify candidata candidate por each
que encajo encajaria why should hire hiring how match matches ejecutivo executive perfil profile promocion promotion
presentacion pager fortalezas strengths frente pide piden busca buscan comparando comparing experiencia experience
right pitching pitch mis sobre about well tan bien explicando explaining reclutador recruiter mi my me""".split())
# the renderer's kinds (build-lane-cv.cjs TEXT_KINDS); the bot's own kinds map onto them
_DOC_RENDER_KIND = {"analysis": "report", "presentation": "other", "promo": "other", "onepager": "other"}
_DOC_FMT_NAME = {"pdf": "PDF", "docx": "Word", "html": "HTML"}
_DOC_FMT_LABEL = {"pdf": "PDF", "docx": "Word (.docx)", "html": "HTML page (opens in any browser)"}
# the example in "which role?" names what she asked for
_DOC_EXAMPLE_WORD = {"letter": ("letter", "carta"), "evaluation": ("evaluation", "evaluación"),
                     "prep": ("interview prep", "preparación de entrevista"), "quote": ("quote", "cotización"),
                     "proposal": ("proposal", "propuesta"), "analysis": ("analysis", "análisis"),
                     "presentation": ("presentation", "presentación"), "promo": ("promotion", "promoción"),
                     "onepager": ("one-pager", "one-pager"), "report": ("report", "informe"),
                     "answer": ("document", "documento")}
# Rule 2 (19 Oct 2026): a document about HER FIT is the deal's own reviewed kit notes, word for word, never written here.
# Each kind names the notes it files, in this order; the 🎤 likely questions (from the 🔎 deep-dive) close a prep sheet.
_DOC_KIT_PARTS = {"evaluation": ("🎯 ROLE DEFENSE", "🛡"), "prep": ("🔎 COMPANY BRIEF", "🎯 ROLE DEFENSE", "🛡")}
_DOC_KIT_NAMES = {"🎯 ROLE DEFENSE": ("🎯 role defense", "🎯 defensa del rol"),
                  "🛡": ("🛡️ technical defense", "🛡️ defensa técnica"),
                  "🔎 COMPANY BRIEF": ("🔎 company brief", "🔎 resumen de la empresa")}


def _cto_aipa_dir() -> str:
    try:
        from src.notifications.telegram_qa_helper import _cto_aipa_dir as qa_dir
        return str(qa_dir())
    except Exception:
        return os.getenv("VJH_CTO_AIPA_DIR") or os.path.join(os.path.expanduser("~"), "cto-aipa")


def _doc_env() -> dict:
    """A small environment for the node scripts: no API keys of this bot leave it. The renderer finds its libraries and
    fonts by the kit's own folders; hs-deal-prep reads cto-aipa's own .env for its HubSpot key (hs-env.cjs)."""
    home = os.path.expanduser("~")
    env = {"PATH": os.environ.get("PATH", "/usr/bin:/bin"), "HOME": home, "LANG": "C.UTF-8"}
    for k in ("SYSTEMROOT", "SystemRoot", "TEMP", "TMP"):        # Windows needs these to start node at all
        if os.environ.get(k):
            env[k] = os.environ[k]
    deps, fonts = os.path.join(home, "cv-build-deps", "node_modules"), os.path.join(home, "cv-build-fonts")
    if os.path.isdir(deps):
        env["NODE_PATH"] = deps
    if os.path.isdir(fonts):
        env["CV_FONT_DIR"] = fonts
    if sys.executable:
        env["VJH_PYTHON"] = sys.executable                         # this venv has python-docx (the renderer's docx)
    return env


def _node_bin() -> str:
    return shutil.which("node") or "/usr/bin/node"


def _renderer_script() -> str:
    return os.path.join(_cto_aipa_dir(), "scripts", "build-lane-cv.cjs")


def _renderer_ready(script: Optional[str] = None, fmt: Optional[str] = None) -> bool:
    """True only when the renderer on disk HAS text mode. An older build-lane-cv.cjs ignores --text and rebuilds every
    lane CV in docs/applications/cv-by-lane instead (Oracle's copy on 18 Oct is that one): never run it then.
    fmt="html": the text mode must also have its HTML branch (19 Oct; a text-mode renderer without it refuses html)."""
    script = script or _renderer_script()
    try:
        with open(script, "r", encoding="utf-8", errors="replace") as f:
            src = f.read(400_000)
    except OSError:
        return False
    ready = "const TEXT_MODE" in src and "async function buildText(" in src
    return ready and ("function textHtml(" in src if fmt == "html" else True)


def _doc_clean_error(err: str) -> str:
    """The renderer's own refusal line ("text: …"), else the first error line, without file paths."""
    lines = [ln.strip() for ln in (err or "").strip().splitlines() if ln.strip()]
    lines = ([ln for ln in lines if ln.startswith("text:")] or [ln for ln in lines if "rror" in ln] or lines)[:1]
    return re.sub(r"(?:(?<=[\s(\"'])|^)(?:[A-Za-z]:)?[\\/][^\s:'\"]*[\\/][^\s:'\")]*", "<file>",
                  lines[-1] if lines else "")[:160]


def _render_document(spec: dict, fmt: str, base: str) -> dict:
    """{ok, bytes, name, placeholders, emoji_dropped, pages} or {ok: False, problem, detail}. problem: no_renderer |
    no_html | no_libs | refused | timeout | no_file. The file is read back and checked (%PDF- / a zip / <!doctype
    html>) before it counts as made."""
    script = _renderer_script()
    if fmt not in ("pdf", "docx", "html") or not _renderer_ready(script):
        return {"ok": False, "problem": "no_renderer", "detail": ""}
    if fmt == "html" and not _renderer_ready(script, "html"):
        return {"ok": False, "problem": "no_html", "detail": ""}
    tmp = tempfile.mkdtemp(prefix="vjh-doc-")
    try:
        spec_path = os.path.join(tmp, "spec.json")
        with open(spec_path, "w", encoding="utf-8") as f:
            json.dump({k: v for k, v in spec.items() if k != "out"}, f, ensure_ascii=False)
        out = os.path.join(tmp, f"{base}.{fmt}")
        try:
            r = subprocess.run([_node_bin(), script, f"--text={spec_path}", f"--format={fmt}", f"--out={out}"],
                               cwd=_cto_aipa_dir(), env=_doc_env(), capture_output=True, timeout=_DOC_RENDER_TIMEOUT_S)
        except subprocess.TimeoutExpired:
            return {"ok": False, "problem": "timeout", "detail": ""}
        except OSError as e:                                        # no node on this machine
            return {"ok": False, "problem": "no_libs", "detail": str(e)[:120]}
        stdout = (r.stdout or b"").decode("utf-8", "replace") if isinstance(r.stdout, bytes) else (r.stdout or "")
        stderr = (r.stderr or b"").decode("utf-8", "replace") if isinstance(r.stderr, bytes) else (r.stderr or "")
        try:
            res = json.loads(stdout.strip().splitlines()[-1]) if stdout.strip() else {}
        except ValueError:
            res = {}
        if r.returncode != 0 or not res.get("ok"):
            detail = stderr or stdout
            problem = "no_libs" if re.search(r"Cannot find module|ENOENT", detail) else "refused"
            return {"ok": False, "problem": problem, "detail": detail}
        try:
            with open(out, "rb") as f:
                data = f.read(_DOC_FILE_MAX + 1)
        except OSError:
            return {"ok": False, "problem": "no_file", "detail": ""}
        magic = {"pdf": b"%PDF-", "docx": b"PK\x03\x04", "html": b"<!doctype html>"}[fmt]
        if not data[:len(magic)].lower().startswith(magic.lower()) or len(data) > _DOC_FILE_MAX:
            return {"ok": False, "problem": "no_file", "detail": ""}
        return {"ok": True, "bytes": data, "name": f"{base}.{fmt}", "pages": res.get("pages"),
                "placeholders": [p for p in (res.get("placeholders") or []) if isinstance(p, str)],
                "emoji_dropped": int(res.get("emoji_dropped") or 0)}
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def _placeholders(text: str) -> List[str]:
    """The renderer's rule, for the text when no file was made: [precio] / [date] are listed; [ ] / [x] (checkboxes)
    and [1] / [2,3] (citations) are not."""
    seen = []
    for p in re.findall(r"\[[^\[\]\n]{1,300}\]", text or ""):
        if re.fullmatch(r"\[\s*[xX]?\s*\]", p) or re.fullmatch(r"\[\d+([,–-]\d+)*\]", p) or p in seen:
            continue
        seen.append(p)
    return seen


def _doc_file_base(*parts: str) -> str:
    """ASCII file name from the parts: 'Cover_Letter_Elena_Revicheva_Developer_Mid_Pitz'."""
    import unicodedata
    s = "_".join(p for p in parts if p)
    s = unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode()
    s = re.sub(r"[^A-Za-z0-9-]+", "_", s).strip("_-")
    return (s[:90].rstrip("_") or "Documento")


def _run_deal_prep(company: str, role: str = "") -> str:
    """cto-aipa scripts/hs-deal-prep.cjs --company (READ-ONLY: it only reads HubSpot). stdout, or '' when it cannot run."""
    script = os.path.join(_cto_aipa_dir(), "scripts", "hs-deal-prep.cjs")
    if not os.path.isfile(script):
        return ""
    args = [_node_bin(), script, f"--company={company}"] + ([f"--role={role}"] if role else [])
    try:
        r = subprocess.run(args, cwd=_cto_aipa_dir(), env=_doc_env(), capture_output=True,
                           timeout=_DOC_PREP_TIMEOUT_S)
    except (subprocess.TimeoutExpired, OSError) as e:
        logger.warning(f"📄 hs-deal-prep did not run: {str(e)[:120]}")
        return ""
    out = (r.stdout or b"").decode("utf-8", "replace") if isinstance(r.stdout, bytes) else (r.stdout or "")
    if r.returncode != 0:
        err = (r.stderr or b"").decode("utf-8", "replace") if isinstance(r.stderr, bytes) else (r.stderr or "")
        logger.warning(f"📄 hs-deal-prep failed: {_doc_clean_error(err)}")
        return ""
    return out


def _parse_deal_prep(out: str) -> dict:
    """hs-deal-prep's printout → {status: ok | none | not_job | unavailable, listed: [(name, stage, id)], name, stage,
    id, sections: [(label, text)], questions: [str]}. Only the deal's notes are sections: the mock-interview blocks after
    them (verified examples, READ-PROOF / NO DEEP-DIVE orders) are left out. questions = the "🎤 THIS COMPANY'S LIKELY
    QUESTIONS" lines hs-deal-prep copies from the deal's 🔎 deep-dive note (for an interview prep sheet)."""
    res = {"status": "unavailable", "listed": [], "name": "", "stage": "", "id": "", "sections": [], "questions": []}
    if not (out or "").strip():
        return res
    if re.search(r"^No \[HIRING-", out, re.M):
        res["status"] = "none"
        return res
    if re.search(r"^Deal \S+ is not a job deal", out, re.M):
        res["status"] = "not_job"
        return res
    label, body = None, []

    def close():
        if label and label.startswith(_DOC_NOTE_LABELS):
            res["sections"].append((label, "\n".join(body).strip()))
        elif label and label.startswith(_DOC_QUESTIONS_LABEL):
            res["questions"] = [ln.strip() for ln in body if re.match(r"^\s*\d{1,2}\.\s+\S", ln)]
    for ln in out.splitlines():
        m = re.match(r"^\s+· (.+) \[([a-z]+)\] id=(\d+)\s*$", ln)
        if m and label is None:
            res["listed"].append((m.group(1).strip(), m.group(2), m.group(3)))
            continue
        h = _DOC_PREP_HEAD.match(ln)
        if h:
            close()
            label, body = h.group(1), []
            if label.startswith("DEAL: "):
                res["name"] = label[6:].strip()
            continue
        if label and label.startswith("DEAL: ") and not res["id"]:
            s = re.match(r"^stage: (\S+) · created: .* · id: (\d+)", ln)
            if s:
                res["stage"], res["id"] = s.group(1), s.group(2)
                continue
        body.append(ln)
    close()
    if res["name"]:
        res["status"] = "ok"
    return res


def _deal_letter(prep: dict) -> dict:
    """The deal's ✅ READY TO SEND letter, word for word, without its heading, the Apply line and the CV line.
    {status: ok | none | several | cut | short, text}. Never repaired, never rewritten."""
    found = [t for _l, t in prep.get("sections") or [] if t.lstrip().startswith(_DOC_READY_MARK)]
    if not found:
        return {"status": "none", "text": ""}
    if len(found) > 1:
        return {"status": "several", "text": ""}
    body = found[0].strip()
    if body.endswith("…[cut]"):
        return {"status": "cut", "text": ""}
    lines = [ln.strip() for ln in body.split("\n")][1:]           # the heading (with "Apply: …" glued to it)
    lines = [ln for ln in lines if ln and not ln.startswith("Apply: ")]
    while lines and re.search(r"attached to this note|superseded", lines[-1], re.I):
        lines.pop()
    text = "\n".join(lines)
    if len(text.split()) < 40:
        return {"status": "short", "text": ""}
    return {"status": "ok", "text": text}


def _deal_kit(prep: dict, kind: str) -> dict:
    """The deal's reviewed kit notes for a fit document (rule 2, 19 Oct 2026), word for word: evaluation → 🎯 ROLE
    DEFENSE + 🛡️ TECHNICAL DEFENSE; prep → 🔎 COMPANY BRIEF + both, then the 🎤 likely questions from the deep-dive.
    {status: ok | none | several | cut, parts: [(label, text)], missing: [label], several: label}. A note's own title
    line ("🎯 ROLE DEFENSE — <deal> 📱 👉 APPLY HERE · letter: the ✅ READY TO SEND note …") is navigation, not content:
    it is left out like the letter's heading (it also carries a ✅ the PDF font cannot draw). Nothing else is touched."""
    parts, missing = [], []
    for want in _DOC_KIT_PARTS[kind]:
        found = [t for label, t in prep.get("sections") or [] if label.startswith(want)]
        if not found:
            missing.append(want)
            continue
        if len(found) > 1:
            return {"status": "several", "parts": [], "missing": [], "several": want}
        body = found[0].strip()
        if body.endswith("…[cut]"):
            return {"status": "cut", "parts": [], "missing": [], "several": ""}
        lines = body.split("\n")
        title = re.sub(r"^[^\w]+", "", lines[0]).upper()
        if want != "🔎 COMPANY BRIEF" and title.startswith(("ROLE DEFENSE", "TECHNICAL DEFENSE")):
            lines = lines[1:]
        text = "\n".join(lines).strip()
        if text:
            parts.append((want, text))
    if not any(w in ("🎯 ROLE DEFENSE", "🛡") for w, _t in parts):
        return {"status": "none", "parts": [], "missing": missing, "several": ""}
    if kind == "prep" and prep.get("questions"):
        parts.append((_DOC_QUESTIONS_LABEL, "\n".join(prep["questions"])))
    return {"status": "ok", "parts": parts, "missing": missing, "several": ""}


def _kit_markdown(parts: List[tuple], lang: str) -> str:
    """The kit notes for the renderer: one heading per note (the bot's label), then the note's own lines unchanged."""
    out = []
    for want, text in parts:
        if want == _DOC_QUESTIONS_LABEL:
            head = ("🎤 Likely questions (from the 🔎 deep-dive)" if lang != "es" else
                    "🎤 Preguntas probables (del 🔎 deep-dive)")
        else:
            head = re.sub(r"[^\W\d_]", lambda c: c.group(0).upper(), _DOC_KIT_NAMES[want][1 if lang == "es" else 0],
                          count=1)                                   # "🎯 defensa del rol" → "🎯 Defensa del rol"
        out += [f"## {head}", "", text, ""]
    return "\n".join(out).strip()


def _letter_markdown(text: str) -> str:
    """One letter line = one paragraph; a short sign-off line ending with "," keeps her name right under it."""
    lines = [ln for ln in text.split("\n") if ln.strip()]
    out = []
    for i, ln in enumerate(lines):
        out.append(ln)
        if i < len(lines) - 1:
            keep = ln.rstrip().endswith(",") and len(ln) <= 40 and i >= len(lines) - 4
            out.append("" if not keep else None)
    return "\n".join(x for x in out if x is not None)


def _doc_markdown(body: str, kind: str) -> str:
    """The model's plain text for the renderer, words unchanged: a short line of its own ending with ":" (a section
    title) becomes a heading. A letter is left as it is ("Estimado cliente:" is a salutation), and so are a deal's kit
    notes (evaluation, prep: already laid out by _kit_markdown, never re-shaped)."""
    if kind in ("letter", "evaluation", "prep"):
        return body
    lines = body.split("\n")
    out = []
    for i, ln in enumerate(lines):
        s = ln.strip()
        nxt = next((x.strip() for x in lines[i + 1:] if x.strip()), "")
        if (s.endswith(":") and 3 <= len(s) <= 70 and nxt and not re.match(r"^(?:[-*•>#]|\d{1,3}[.)])", s)
                and "[" not in s):
            out.append("### " + s)
        else:
            out.append(ln)
    return "\n".join(out)


def _deal_title(name: str) -> str:
    """'[HIRING-VJH-LEAD] Developer Mid @ Pitz' → 'Developer Mid @ Pitz'."""
    return re.sub(r"^\s*\[[^\]]*\]\s*", "", name or "").strip()


def _snapshot_companies() -> List[str]:
    """Companies in her live-deals snapshot (OpenClaw deals/TOOLS.md, '### Company — Title (stage)' rows)."""
    try:
        from src.notifications import telegram_qa_helper as q
        raw = q._read_text(q._OPENCLAW_WS / "deals" / "TOOLS.md")
    except Exception:
        return []
    names = []
    for ln in raw.splitlines():
        if ln.startswith("### "):
            # fix 8 (held-out p8): "Addi @ Addi", "WON @ Won.Ai" — each side of the "@" is the company's name
            for co in ln[4:].split(" — ")[0].split(" @ "):
                co = co.strip()
                if len(co) >= 3 and co not in names:
                    names.append(co)
    return names


_DOC_CO_PREP = (r"(?:for|para|about|sobre|de|del|at|en|with|con|a|al|to)\s+(?:the\s+company\s+|(?:la\s+)?empresa\s+|"
                r"(?:el\s+)?deal\s+(?:de\s+)?)?")


# 19 Oct review (fix 8, wrong deal): "una propuesta para Clara Dental", "para la doctora Clara Pérez" went to the Clara
# JOB deal. A snapshot company is the one she means only when it stands alone: never right after a person's title or a
# kind of business ("la doctora Clara", "la clínica Clara"), never the start of a longer capitalised name ("Clara
# Dental", "Clara Pérez"). A role title, a format or a company suffix after it is fine ("Pitz Data Engineer", "Puente
# Talent Partners RevOps", "Pitz Inc").
_DOC_CO_TITLE_BEFORE = re.compile(
    r"(?:^|[\s,(])(?:doctor[a]?|dra?\.?|se[nñ]or[a]?|sra?\.?|srta\.?|don|do[nñ]a|lic\.?|licenciad[oa]|ing\.?|"
    r"ingenier[oa]|prof\.?|profesor[a]?|mr\.?|mrs\.?|ms\.?|miss|dear|estimad[oa]|querid[oa]|client[ea]|clienta|"
    r"cl[ií]nica|clinic|consultorio|[oó]ptica|farmacia|restaurante|restaurant|hotel|spa|sal[oó]n|gimnasio|gym|"
    r"veterinaria|panader[ií]a|tienda|store|caf[eé]|academia|escuela|colegio|school|bufete|despacho|marca|brand|"
    r"amig[oa]|herman[oa]|hija|mam[aá]|madre|vecin[oa])\s+$", re.I)
_DOC_CO_SUFFIX = frozenset("inc llc ltd sa s.a sas srl corp co gmbh plc ag bv".split())


def _doc_co_alone(text: str, start: int, end: int) -> bool:
    """Fix 8: the company at text[start:end] stands alone (see _DOC_CO_TITLE_BEFORE)."""
    if _DOC_CO_TITLE_BEFORE.search(text[:start]):
        return False
    nxt = re.match(r"[ \t]+([A-ZÁÉÍÓÚÑ][\w&'’.-]*)", text[end:])
    if not nxt:
        return True
    w = _doc_fold_name(nxt.group(1)).strip(".'’-")
    return (w in _DOC_NOT_COMPANY or w in _DOC_ROLE_HEADS or w in _DOC_ROLE_WORDS or w in _DOC_CO_SUFFIX
            or w in _DOC_FMT_TOKENS)


def _doc_companies(text: str, prep: bool = False) -> List[str]:
    """The live job-deal companies (her snapshot) her message names: after a preposition ("para Pitz", "de pitz" —
    any case, a voice note is lower-case), or anywhere when written with a capital and not as the message's first word
    ("…con lo que busca Pitz"). prep=True: only after a preposition. A name inside another one ("Pitz" in "Pitz Labs")
    counts once, as the longer. A common word that is also a company ("Clara" / "una propuesta clara", "WITHIN" /
    "within a week") counts only in those two forms. Fix 8: and only standing alone (_doc_co_alone)."""
    text = text or ""
    low = text.lower().translate(_DOC_FOLD)            # the same length as text: positions line up
    lead = len(text) - len(text.lstrip(" ¿¡\"'("))
    hits = []
    for c in _snapshot_companies():
        fc = re.escape(_doc_fold(c)).replace(r"\ ", r"\s+")
        found = False
        for m in re.finditer(r"(?<![\w])" + fc + r"(?![\w])", low):
            if not _doc_co_alone(text, m.start(), m.end()):
                continue
            if re.search(r"\b" + _DOC_CO_PREP + r"$", low[:m.start()]) or (
                    not prep and m.start() > lead and text[m.start()].isupper()):
                found = True
                break
        if found and c not in hits:
            hits.append(c)
    return [c for c in hits if not any(c != o and _doc_fold(c) in _doc_fold(o) for o in hits)]


def _named_company(text: str) -> Optional[str]:
    """The ONE live-deal company she names right after a preposition ("for Pitz", "para Pitz", "de Pitz"); None when
    none or two."""
    hits = _doc_companies(text, prep=True)
    return hits[0] if len(hits) == 1 else None


# capitalised words after a preposition that name no company: her own products, a role, a place, a format, a month
_DOC_NOT_COMPANY = frozenset(_doc_fold("""pdf word html docx ms microsoft ia ai cv elena revicheva aideazz espaluz
vibejobhunter vjh cto aipa openclaw atuona whitespace algom atlas audit api growth operator chief head senior junior
lead director manager engineer developer program product project business remote sdet qa revops sales marketing data
software fullstack full frontend backend operations ops vp ceo coo cfo cmo staff principal specialist analyst
consultant account customer technical solutions automation forward mid panamá panama bogotá bogota colombia latam
latinoamérica latinoamerica méxico mexico usa eeuu lima perú peru chile argentina brasil brazil costa rica europa europe
españa espana spain miami linkedin instagram telegram whatsapp hubspot google zapier make n8n python javascript
typescript node react enero febrero marzo abril mayo junio julio agosto septiembre octubre noviembre diciembre january
february march april may june july august september october november december lunes martes miércoles jueves viernes
sábado domingo monday tuesday wednesday thursday friday saturday sunday english inglés spanish español""").split())


def _doc_caps(text: str) -> List[str]:
    """Capitalised names right after a preposition ("…para Acme Labs", "…at Clara", "…de Acme Labs") that are not
    her own products, a role, a place, a format or a month: the companies of a request whose company is not in her
    snapshot. In the order written."""
    out = []
    for c in re.finditer(r"\b(?:for|para|at|en|con|with|de|del|to|al|a|vs\.?|versus)\s+(?:the\s+company\s+|"
                         r"(?:la\s+)?empresa\s+)?([A-ZÁÉÍÓÚÑ0-9][\w&'+-]*(?:\.\w+)*(?:\s+(?:de\s+|del\s+|&\s+)?"
                         r"[A-ZÁÉÍÓÚÑ0-9][\w&'+-]*(?:\.\w+)*){0,3})", text or ""):
        name = c.group(1).strip(" .'-")
        if _doc_fold(name.split()[0]) in _DOC_NOT_COMPANY or name in out:
            continue
        out.append(name[:60])
    return out


def _doc_fit(text: str, kind: str, cos: List[str], caps: List[str]) -> bool:
    """Rule 2, by structure: is this a document about HER FIT for a job? A live job-deal company named anywhere; an
    interview prep; an evaluation of her or of a job; her fit words ("encajo", "match", "right for", "hire me",
    "mis fortalezas frente a lo que pide…") next to a job or a company; a hiring fit word on its own ("¿por qué
    encajo?"); or "sobre mí / my experience" next to a job. Whatever the document's noun."""
    if cos or kind == "prep":
        return True
    low = _doc_fold(text)
    her, req, me, job = (_DOC_FIT_HER.search(low), _DOC_FIT_REQ.search(low), _DOC_FIT_ME.search(low),
                         _DOC_FIT_JOB.search(low))
    if kind == "evaluation":
        return bool(her or req or me or job)
    return bool((her and (job or caps)) or _DOC_FIT_SELF.search(low) or (me and (job or (req and caps)))
                or (req and job))


# fix 8 (wrong deal): a CLIENT kind names a live job-deal company — it is about that job only with a job or fit word
_DOC_CLIENT_KINDS = frozenset(("proposal", "promo", "onepager", "presentation"))
_DOC_JOB_WORD = re.compile(
    r"\b(?:puesto|vacante|cargo|rol|role|position|job|empleo|trabajo|opening|vacancy|posting|aplicacion|application|"
    r"postulacion|candidatura|candidacy|entrevista\w*|interview\w*|recruiter\w*|reclutador\w*|hiring|contratacion|"
    r"contratarme|hire)\b|\b(?:de|sobre|about)\s+(?:mi|me|myself)\b(?!\s+(?:parte|lado|cuenta|empresa|negocio|"
    r"cliente|clienta|company|business|client))")


def _doc_job_words(text: str) -> bool:
    """Fix 8: her message speaks of a job (a job word, or her fit / her profile / what a job asks)."""
    low = _doc_fold(text)
    return bool(_DOC_JOB_WORD.search(low) or _DOC_FIT_HER.search(low) or _DOC_FIT_SELF.search(low)
                or _DOC_FIT_ME.search(low) or _DOC_FIT_REQ.search(low) or _DOC_FIT_JOB.search(low))


def _doc_fit_question(question: str) -> bool:
    """A chat question about HER FIT for a job ("¿Encajo en la vacante de Pitz?", "Should I apply to Clara?",
    "¿Cumplo los requisitos?"): its answer is a fit evaluation written in the chat, never filed (19 Oct review)."""
    low = _doc_fold(question)
    # fix 8 (review: "¿Qué opinas de la vacante de Pitz?" → "PDF" filed "Eres una candidata muy fuerte para Pitz"):
    # a question that names one of her live job deals is about that job, whatever it asks ("Háblame de Pitz", "Vale la
    # pena Pitz?", "Analiza la oferta de Clara") — its answer reads her fit, and a fit document is only the deal's kit
    if _doc_companies(question):
        return True
    job = _DOC_FIT_JOB.search(low)
    return bool(_DOC_FIT_HER.search(low) or _DOC_FIT_SELF.search(low) or (_DOC_FIT_REQ.search(low) and job)
                or (_DOC_FIT_ME.search(low) and job) or (job and re.search(
                    r"\b(?:aplic\w*|apply\w*|postul\w*|estoy|soy|tengo|deberia|should\s+i|am\s+i|do\s+i|me\s+conviene|"
                    r"worth|opinas|piensas|te\s+parece|think|thoughts|vale\s+la\s+pena|merece\s+la\s+pena|hablame|"
                    r"cuentame|tell\s+me\s+about|analiza\w*|analy[sz]e\w*|evalua\w*|evaluate\w*|assess\w*|review\w*|"
                    r"revisa\w*|recomiendas|recommend\w*|chances?|posibilidades|probabilidad\w*)\b", low)))


def _letter_company(text: str) -> Optional[str]:
    """The company of "letter for X" / "carta (de presentación) para X": a live-deal company she names, else the words
    after the preposition (up to 4, until a format word or punctuation)."""
    low = _doc_fold(text)
    m = re.search(r"\b(cover\s+letter|(?<!offer )letter|carta(?:\s+de\s+presentacion)?)(?!\s+(?:de\s+)?oferta)\b"
                  r"(?P<rest>[^.!?\n]{0,90})", low)
    if not m:
        return None
    named = _named_company(text)
    if named:
        return named
    preps = r"for|to|para|al|a" + (r"|de|del" if m.group(1) in ("cover letter", "carta de presentacion") else "")
    c = re.search(rf"\b(?:{preps})\s+(?P<co>.+)", m.group("rest"))
    if not c:
        return None
    words = []
    for w in re.findall(r"[\w&'.+-]+|[,;:]", c.group("co")):
        if w in (",", ";", ":") or w in ("in", "as", "en", "como", "y", "and", "please", "por", "pdf", "word", "docx",
                                         "ms", "microsoft", "que", "that", "with", "con", "now", "ahora", "hoy",
                                         "today", "formato", "format", "file", "archivo"):
            break
        words.append(w)
        if len(words) == 4:
            break
    while words and words[0] in ("la", "el", "los", "las", "the", "a", "an", "empresa", "company", "compania", "deal",
                                 "de", "del", "puesto", "role", "job", "trabajo", "vacante"):
        words.pop(0)
    co = " ".join(words).strip(" .'-")
    if len(co) < 2 or co in ("me", "mi", "mí", "it", "him", "her", "them", "elena", "you", "ti", "ella", "ellos"):
        return None
    return co[:60]


def _role_words(text: str, company: str) -> str:
    """Her words beyond the request and the company — the role, when one company has several job deals."""
    co = set(re.findall(r"\w+", _doc_fold(company)))
    words = [w for w in re.findall(r"[a-z0-9ñ]+", _doc_fold(text))
             if len(w) > 2 and w not in co and w not in _DOC_ROLE_STOP and w not in _DOC_FMT_TOKENS]
    return " ".join(words)[:80]


def _is_bare_doc_request(text: str) -> bool:
    """Only "make it a file" ("PDF", "send me that as a Word file", "hazme un PDF de eso"): no task of its own, so the
    file is the answer she just got."""
    toks = _doc_bare_tokens(_doc_fold(text))
    return bool(toks) and all(t in _DOC_BARE_WORDS or t in _DOC_FMT_TOKENS for t in toks)


# Lines the bot writes itself around a document: (English, Spanish), in the language she wrote in.
_DOC_SAY = {
    "no_renderer": ("📄 No file this time: the document renderer (cto-aipa build-lane-cv.cjs, text mode) is not on this "
                    "server yet{extra}.",
                    "📄 Esta vez sin archivo: el generador de documentos (build-lane-cv.cjs de cto-aipa, modo texto) "
                    "todavía no está en este servidor{extra}."),
    "no_renderer_extra": (" — the text above is the whole document", " — el texto de arriba es el documento completo"),
    "marks": ("📄 No {fmt} file: the text uses marks the file's font cannot draw ({marks}), and leaving them out could "
              "erase a verdict. Ask again with \"in PDF\" in the same message and I will write them in words.",
              "📄 Sin archivo {fmt}: el texto usa marcas que la fuente del archivo no puede dibujar ({marks}) y quitarlas "
              "podría borrar un veredicto. Pídelo otra vez con \"en PDF\" en el mismo mensaje y lo escribo con palabras."),
    # 18 Oct review: no model rewrites a letter or a last answer, so asking again with "PDF" alone changes nothing
    "marks_letter": ("📄 No {fmt} file: the letter note itself contains marks the file's font cannot draw ({marks}). "
                     "I never edit a letter: change the note in HubSpot, then ask again.",
                     "📄 Sin archivo {fmt}: la nota de la carta contiene marcas que la fuente del archivo no puede dibujar "
                     "({marks}). Nunca edito una carta: cambia la nota en HubSpot y pídelo otra vez."),
    "marks_last": ("📄 No {fmt} file: your last answer uses marks the file's font cannot draw ({marks}), and leaving "
                   "them out could erase a verdict. Send the question itself again with \"in PDF\" in the same message: "
                   "a document answer writes them in words.",
                   "📄 Sin archivo {fmt}: tu última respuesta usa marcas que la fuente del archivo no puede dibujar "
                   "({marks}) y quitarlas podría borrar un veredicto. Envía la pregunta misma otra vez con \"en PDF\" en "
                   "el mismo mensaje: la respuesta como documento las escribe con palabras."),
    "cv": ("📄 I do not write CVs here. The CV for a job is the posting-shaped one the apply kit builds on that job's "
           "HubSpot deal (the 🛡️ note); open it there. Nothing was made.",
           "📄 Aquí no escribo CVs. El CV para un empleo es el que el kit arma a la medida de la oferta en el deal de "
           "HubSpot (la nota 🛡️); ábrelo allí. No hice nada."),
    "cover_which": ("✉️ Which job? A cover letter comes word for word from that job deal's ✅ note: say \"cover letter "
                    "for <company> in PDF\". I do not write cover letters myself.",
                    "✉️ ¿Para qué empleo? La carta de presentación sale palabra por palabra de la nota ✅ de ese deal: "
                    "dime \"carta de presentación para <empresa> en PDF\". No escribo cartas de presentación por mi "
                    "cuenta."),
    "refused": ("📄 No {fmt} file: the renderer refused it ({why}). Nothing was saved.",
                "📄 Sin archivo {fmt}: el generador lo rechazó ({why}). No se guardó nada."),
    "no_libs": ("📄 No {fmt} file: the renderer is here but its libraries or fonts are missing ({why}).",
                "📄 Sin archivo {fmt}: el generador está, pero faltan sus librerías o fuentes ({why})."),
    "timeout": ("📄 No {fmt} file: the renderer took too long, so I stopped.",
                "📄 Sin archivo {fmt}: el generador tardó demasiado y paré."),
    "no_file": ("📄 No {fmt} file: the renderer finished but no valid file came out, so I sent nothing.",
                "📄 Sin archivo {fmt}: el generador terminó pero no salió un archivo válido, así que no envié nada."),
    "no_last": ("📄 I have no recent answer to turn into a file (nothing asked in the last 6 hours, or the bot restarted). "
                "Send the task itself with \"in PDF\" or \"in Word\".",
                "📄 No tengo una respuesta reciente para convertir en archivo (nada en las últimas 6 horas, o el bot se "
                "reinició). Envíame la tarea misma con \"en PDF\" o \"en Word\"."),
    "from_last": ("📄 Your last answer as a file{q}.", "📄 Tu última respuesta como archivo{q}."),
    "fill": ("To fill in before you send it: {items}", "Por completar antes de enviarlo: {items}"),
    "emoji": ("{n} emoji left out of the file (its font cannot draw them).",
              "{n} emojis quedaron fuera del archivo (su fuente no los dibuja)."),
    "deal_used": ("{n} job deals match \"{co}\"; I used {deal}.", "{n} deals de empleo coinciden con \"{co}\"; usé {deal}."),
    "not_attached": ("Not attached in HubSpot: VJH cannot upload files there{why}.",
                     "No lo adjunté en HubSpot: VJH no puede subir archivos allí{why}."),
    "not_attached_letter": (" (the letter note itself is already on the deal)", " (la nota de la carta ya está en el deal)"),
    "letter_ok": ("✉️ The ✅ cover letter on {deal} ({stage}), word for word as it is in HubSpot.",
                  "✉️ La carta ✅ del deal {deal} ({stage}), palabra por palabra como está en HubSpot."),
    "letter_en_only": ("", "La carta está en inglés, como la escribió el kit; no traduzco cartas."),
    "letter_lang": ("✉️ The ✅ letter on {deal} is in {have}, and I do not translate letters, so I made no file.",
                    "✉️ La carta ✅ del deal {deal} está en {have} y no traduzco cartas, así que no hice el archivo."),
    "no_deal": ("✉️ No job deal in HubSpot matches \"{co}\", so there is no letter to put in a file. I do not write "
                "cover letters myself.",
                "✉️ Ningún deal de empleo en HubSpot coincide con \"{co}\", así que no hay carta para el archivo. No "
                "escribo cartas de presentación por mi cuenta."),
    "deal_unavailable": ("✉️ I could not read the HubSpot deals just now (cto-aipa hs-deal-prep did not answer). Nothing "
                         "was made; ask again in a minute.",
                         "✉️ No pude leer los deals de HubSpot ahora (hs-deal-prep de cto-aipa no respondió). No hice "
                         "nada; pídelo otra vez en un minuto."),
    "wrong_co": ("✉️ \"{co}\" matches {deal} by its role, not its company, so I made no file. Name the company.",
                 "✉️ \"{co}\" coincide con {deal} por el puesto, no por la empresa, así que no hice el archivo. Dime la "
                 "empresa."),
    "pick_role": ("✉️ {n} 🔥 job deals at {co}: {names}. Say which role, for example \"{doc} for {co} {role} in PDF\". "
                  "Nothing was made.",
                  "✉️ {n} deals 🔥 de empleo en {co}: {names}. Dime qué puesto, por ejemplo \"{doc} para {co} {role} en "
                  "PDF\". No hice nada."),
    "no_letter": ("✉️ {deal} has no ✅ READY TO SEND letter yet. The apply kit writes one for 🔥 I act today deals every "
                  "10 minutes; I do not write cover letters myself.",
                  "✉️ {deal} todavía no tiene la carta ✅ READY TO SEND. El kit la escribe cada 10 minutos para los deals "
                  "en 🔥 I act today; yo no escribo cartas de presentación por mi cuenta."),
    "several": ("✉️ {deal} carries {n} ✅ letters and I cannot tell which one is current, so I made no file.",
                "✉️ {deal} tiene {n} cartas ✅ y no sé cuál es la vigente, así que no hice el archivo."),
    "cut": ("✉️ The ✅ letter on {deal} is longer than the deal reader returns, so I made no file rather than a cut one.",
            "✉️ La carta ✅ de {deal} es más larga de lo que devuelve el lector del deal: no hago un archivo cortado."),
    "short": ("✉️ The ✅ note on {deal} holds too little text to be the letter, so I made no file.",
              "✉️ La nota ✅ de {deal} tiene muy poco texto para ser la carta, así que no hice el archivo."),
    # 19 Oct 2026: fit documents (rule 2), two documents at once, HTML, the honesty gate, a swipe-reply to an older answer
    "two_docs": ("📄 One document per message, please: this one asks for two, and I would make only the first. Send "
                 "them one at a time.",
                 "📄 Un documento por mensaje, por favor: este pide dos y solo haría el primero. Envíamelos uno por uno."),
    "kit_which": ("🎯 Which job? An evaluation or interview prep is that job deal's own kit notes (🎯 role defense, 🛡️ "
                  "technical defense), word for word: say \"evaluation for <company> in PDF\". I never evaluate your fit "
                  "myself. A posting that is not a deal yet: stage it first, and the apply kit writes those notes.",
                  "🎯 ¿Para qué empleo? Una evaluación o una preparación de entrevista son las notas del kit de ese deal "
                  "(🎯 defensa del rol, 🛡️ defensa técnica), palabra por palabra: dime \"evaluación para <empresa> en "
                  "PDF\". Nunca evalúo tu encaje por mi cuenta. Si la oferta aún no es un deal, primero hay que crearlo y "
                  "el kit escribe esas notas."),
    "kit_no_deal": ("🎯 No job deal in HubSpot matches \"{co}\", so there are no kit notes to put in a file. I never "
                    "evaluate your fit myself.",
                    "🎯 Ningún deal de empleo en HubSpot coincide con \"{co}\", así que no hay notas del kit para el "
                    "archivo. Nunca evalúo tu encaje por mi cuenta."),
    "kit_none": ("🎯 {deal} has no 🎯 role defense or 🛡️ technical defense note yet. The apply kit writes them for 🔥 I "
                 "act today deals; I never evaluate your fit myself, so I made no file.",
                 "🎯 {deal} todavía no tiene la nota 🎯 de defensa del rol ni la 🛡️ de defensa técnica. El kit las escribe "
                 "para los deals en 🔥 I act today; nunca evalúo tu encaje por mi cuenta, así que no hice el archivo."),
    "kit_several": ("🎯 {deal} carries more than one {note} note and I cannot tell which one is current, so I made no "
                    "file.",
                    "🎯 {deal} tiene más de una nota {note} y no sé cuál es la vigente, así que no hice el archivo."),
    "kit_cut": ("🎯 A kit note on {deal} is longer than the deal reader returns, so I made no file rather than a cut one.",
                "🎯 Una nota del kit de {deal} es más larga de lo que devuelve el lector del deal: no hago un archivo "
                "cortado."),
    "kit_ok": ("🎯 {what} — {deal} ({stage}): the kit's {notes} notes, word for word as they are in HubSpot. Nothing was "
               "evaluated again.",
               "🎯 {what} — {deal} ({stage}): las notas {notes} del kit, palabra por palabra como están en HubSpot. No "
               "evalué nada de nuevo."),
    "kit_missing": ("This deal has no {notes} note, so it is not in the file.",
                    "Este deal no tiene la nota {notes}, así que no está en el archivo."),
    "kit_en_only": ("", "Las notas están en inglés, como las escribió el kit; no las traduzco."),
    "marks_kit": ("📄 No {fmt} file: the kit notes themselves contain marks the file's font cannot draw ({marks}). I "
                  "never edit them: ask for it \"in HTML\" (a browser draws them) or change the note in HubSpot.",
                  "📄 Sin archivo {fmt}: las notas del kit contienen marcas que la fuente del archivo no puede dibujar "
                  "({marks}). Nunca las edito: pídelo \"en HTML\" (un navegador las dibuja) o cambia la nota en HubSpot."),
    "no_html": ("📄 No HTML file: the renderer on this server has no HTML mode yet. Ask for PDF or Word.",
                "📄 Sin archivo HTML: el generador de este servidor todavía no tiene modo HTML. Pídelo en PDF o Word."),
    "gated": ("Left blank because no source holds it (your verified facts, your request, your offers): {items}",
              "Lo dejé en blanco porque ninguna fuente lo respalda (tus datos verificados, tu pedido, tus ofertas): "
              "{items}"),
    "quote_ok": ("📄 Quote from your published catalog, word for word (titles, prices, descriptions); no model wrote it.",
                 "📄 Cotización con tu catálogo publicado, palabra por palabra (títulos, precios, descripciones); ningún "
                 "modelo la escribió."),
    "from_reply": ("📄 The message you replied to, as a file.", "📄 El mensaje al que respondiste, como archivo."),
    # 19 Oct review: two job deals in one request (the wrong deal was filed); a fit answer asked "as a file"
    "which_co": ("📄 This names {cos}: which one is it for? Ask again naming only one company. Nothing was made.",
                 "📄 Esto nombra {cos}: ¿para cuál es? Pídelo otra vez nombrando solo una empresa. No hice nada."),
    "job_or_client": ("📄 {co} is one of your job deals. For that job, the document is its kit, word for word: say "
                      "\"evaluation for {co} in PDF\" (or \"interview prep\"). If this one is for a client or another "
                      "company with that name, write its full name (\"Clara Dental\", \"the clinic …\"). Nothing was made.",
                      "📄 {co} es uno de tus deals de empleo. Para ese empleo, el documento es su kit, palabra por palabra: "
                      "dime \"evaluación para {co} en PDF\" (o \"preparación de entrevista\"). Si este es para un cliente "
                      "u otra empresa con ese nombre, escribe su nombre completo (\"Clara Dental\", \"la clínica …\"). "
                      "No hice nada."),
    "kit_last": ("🎯 That answer is about your fit for a job, and I never file a fit evaluation written in the chat: a "
                 "fit document is that job deal's own kit notes (🎯 role defense, 🛡️ technical defense), word for word. "
                 "Say \"evaluation for <company> in PDF\". Nothing was made.",
                 "🎯 Esa respuesta trata de tu encaje en un empleo, y nunca archivo una evaluación de encaje escrita en el "
                 "chat: un documento de encaje son las notas del kit de ese deal (🎯 defensa del rol, 🛡️ defensa "
                 "técnica), palabra por palabra. Dime \"evaluación para <empresa> en PDF\". No hice nada."),
}


def _doc_say(lang: str, key: str, **kw) -> str:
    en, es = _DOC_SAY[key]
    return (es if lang == "es" else en).format(**kw)


def _doc_why(lang: str, detail: str) -> str:
    """The renderer's refusal in plain words."""
    d = detail or ""
    es = lang == "es"
    if "tables are not supported" in d:
        return "tiene una tabla; el archivo usa líneas numeradas" if es else "it has a table; the file takes numbered lines"
    if "no glyph" in d:
        chars = ", ".join(re.findall(r'U\+[0-9A-F]{4,6} "[^"]*"', d)[:6])
        return (f"caracteres que su fuente no dibuja: {chars}" if es else f"characters its font cannot draw: {chars}")
    if "over the 20-page cap" in d or re.search(r"\d+ pages — over", d):
        return "más de 20 páginas" if es else "over 20 pages"
    if "python-docx" in d:
        return "Word necesita python-docx, que falta aquí; pide PDF" if es else "Word needs python-docx, missing here; ask for PDF"
    mod = re.search(r"Cannot find module '([^']+)'", d)
    if mod:
        return f"falta {mod.group(1)}" if es else f"{mod.group(1)} is missing"
    if "ENOENT" in d:
        return "falta un archivo de fuente" if es else "a font file is missing"
    if "pass/fail marks" in d:
        return "marcas de sí/no que su fuente no dibuja" if es else "yes/no marks its font cannot draw"
    return _doc_clean_error(d) or ("sin motivo" if es else "no reason given")


class EnhancedTelegramBot:
    """
    Interactive Telegram bot for VibeJobHunter
    
    🆕 UPGRADED January 2026: Comprehensive menu system
    
    MAIN MENU COMMANDS:
    /menu - Interactive menu with all options
    /workflow - Visual explanation of how the engine works
    /manual - What YOU need to do (steps that can't be automated)
    /today - Today's complete activity summary
    /outreach - Pending LinkedIn/Twitter messages to send manually
    
    ORIGINAL COMMANDS:
    /start - Welcome message and help
    /jobs - Show top jobs found today
    /stats - Show application statistics
    /apply <company> - Generate materials for a company
    /pause - Pause autonomous job hunting
    /resume - Resume autonomous job hunting
    /keywords <keywords> - Update job search keywords
    /companies - Show tracked companies
    /recent - Show recent applications
    /status - Show system status

    VOICE & QUESTIONS (added, additive):
    - Send a voice message → transcribed with Whisper → answered from real data
    - Type: "What companies did I apply to today?" → honest list from logs
    """
    
    def __init__(self, token: str, chat_id: str, db_helper=None):
        self.token = token
        self.chat_id = chat_id
        self.db_helper = db_helper
        self.is_paused = False
        
        if not TELEGRAM_AVAILABLE:
            raise ImportError("python-telegram-bot not installed")
        
        self.app = Application.builder().token(token).build()
        self._register_handlers()
    
    def _chat_lock_filter(self):
        """ONE chat lock built from TELEGRAM_CHAT_ID (16 Oct 2026): answers now carry her pipeline, screenshots and
        voice cost money, and the 📘 button posts publicly. A missing or non-numeric id applies NO lock, so she can
        never be locked out of her own bot."""
        raw = self.chat_id if self.chat_id not in (None, "") else os.getenv("TELEGRAM_CHAT_ID", "")
        try:
            self._allowed_chat_id = int(str(raw).strip())
        except (TypeError, ValueError):
            self._allowed_chat_id = None
            logger.warning("⚠️ TELEGRAM_CHAT_ID is missing or not a number: the VJH bot chat lock is OFF")
            return None
        return filters.Chat(chat_id=self._allowed_chat_id)

    def _chat_allowed(self, update) -> bool:
        """The same lock for updates a filter cannot gate (CallbackQueryHandler takes no filters in PTB 20.7)."""
        allowed = getattr(self, "_allowed_chat_id", None)
        if allowed is None:
            return True
        chat = getattr(update, "effective_chat", None)
        return bool(chat) and chat.id == allowed

    def _register_handlers(self):
        """Register all command handlers"""
        lock = self._chat_lock_filter()
        # CommandHandler's own default is UpdateType.MESSAGES; the lock is added on top of it. No lock → None (default).
        cmd = (filters.UpdateType.MESSAGES & lock) if lock else None

        def locked(f):
            return (f & lock) if lock else f

        # 🆕 NEW MAIN MENU COMMANDS (January 2026)
        self.app.add_handler(CommandHandler("menu", self.cmd_menu, filters=cmd))
        self.app.add_handler(CommandHandler("workflow", self.cmd_workflow, filters=cmd))
        self.app.add_handler(CommandHandler("manual", self.cmd_manual, filters=cmd))
        self.app.add_handler(CommandHandler("today", self.cmd_today, filters=cmd))
        self.app.add_handler(CommandHandler("outreach", self.cmd_outreach, filters=cmd))
        
        # Original commands
        self.app.add_handler(CommandHandler("start", self.cmd_start, filters=cmd))
        self.app.add_handler(CommandHandler("help", self.cmd_help, filters=cmd))
        self.app.add_handler(CommandHandler("jobs", self.cmd_jobs, filters=cmd))
        self.app.add_handler(CommandHandler("stats", self.cmd_stats, filters=cmd))
        self.app.add_handler(CommandHandler("apply", self.cmd_apply, filters=cmd))
        self.app.add_handler(CommandHandler("pause", self.cmd_pause, filters=cmd))
        self.app.add_handler(CommandHandler("resume", self.cmd_resume, filters=cmd))
        self.app.add_handler(CommandHandler("keywords", self.cmd_keywords, filters=cmd))
        self.app.add_handler(CommandHandler("companies", self.cmd_companies, filters=cmd))
        self.app.add_handler(CommandHandler("recent", self.cmd_recent, filters=cmd))
        self.app.add_handler(CommandHandler("status", self.cmd_status, filters=cmd))
        # Priority companies (additive — Feb 2026)
        self.app.add_handler(CommandHandler("priority", self.cmd_priority, filters=cmd))

        # Callback query handler for inline buttons (chat-locked inside handle_callback)
        self.app.add_handler(CallbackQueryHandler(self.handle_callback))

        # ADDED: Voice + free-form text (additive — commands unchanged)
        self.app.add_handler(MessageHandler(locked(filters.VOICE), self.handle_voice))
        self.app.add_handler(MessageHandler(locked(filters.TEXT & ~filters.COMMAND), self.handle_freeform_text))
        # ADDED 16 Oct 2026: screenshots (a photo, or an image sent as a file for full resolution)
        self.app.add_handler(MessageHandler(locked(filters.PHOTO | filters.Document.IMAGE), self.handle_photo))
        # ADDED 16 Oct 2026: PDF and Word .docx documents; any other file gets an honest "not read" line instead of
        # silence. After the screenshot handler, so an image sent as a file still goes to handle_photo.
        self.app.add_handler(MessageHandler(
            locked(filters.Document.ALL & ~filters.Document.IMAGE & ~filters.ANIMATION), self.handle_document))
    
    async def cmd_start(self, update: Update, context: ContextTypes.DEFAULT_TYPE):
        """Welcome message with main menu"""
        await update.message.reply_text(START_TEXT, parse_mode='Markdown', reply_markup=_menu_keyboard())

    # ═══════════════════════════════════════════════════════════════════════════
    # 🆕 NEW MENU COMMANDS (January 2026 Upgrade)
    # ═══════════════════════════════════════════════════════════════════════════

    async def cmd_menu(self, update: Update, context: ContextTypes.DEFAULT_TYPE):
        """Interactive main menu with all options"""
        import logging
        logger = logging.getLogger(__name__)
        logger.info("📋 /menu command received!")

        try:
            # 16 Oct 2026 (Elena): everything about jobs lives in THIS bot, whichever agent does the work behind it.
            await update.message.reply_text(MENU_TEXT, parse_mode='Markdown', reply_markup=_menu_keyboard())
            logger.info("📋 /menu response sent successfully!")
        except Exception as e:
            logger.error(f"❌ /menu command error: {e}", exc_info=True)

    async def cmd_workflow(self, update: Update, context: ContextTypes.DEFAULT_TYPE):
        """How the engine really works today"""
        keyboard = [[InlineKeyboardButton("👤 What is mine", callback_data="manual")]]
        await update.message.reply_text(WORKFLOW_TEXT, parse_mode='Markdown', reply_markup=InlineKeyboardMarkup(keyboard))

    async def cmd_manual(self, update: Update, context: ContextTypes.DEFAULT_TYPE):
        """Show what the user needs to do manually"""
        keyboard = [[InlineKeyboardButton("📋 Back to Menu", callback_data="menu")]]
        await update.message.reply_text(MANUAL_TEXT, parse_mode='Markdown', reply_markup=InlineKeyboardMarkup(keyboard))

    async def cmd_today(self, update: Update, context: ContextTypes.DEFAULT_TYPE):
        """Show today's complete activity summary"""
        try:
            # Load today's data
            jobs_count = self._count_todays_jobs()
            outreach = self._load_pending_outreach()
            
            now = datetime.utcnow()
            panama_hour = (now.hour - 5) % 24  # UTC-5
            
            message = f"""📊 *TODAY'S ACTIVITY SUMMARY*
📅 {now.strftime('%B %d, %Y')} | {panama_hour}:{now.strftime('%M')} Panama

━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
*🔍 JOB DISCOVERY*
━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
• Jobs Found: {jobs_count}
• Sources Checked: 7
• Companies Scanned: 200+

━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
*🚀 APPLICATIONS*
━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
• Auto-Applied: Check /stats
• Materials Generated: ✅
• ATS Forms Filled: ✅

━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
*📨 YOUR ACTION NEEDED*
━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
• LinkedIn Messages: {len(outreach)} pending
• Use /outreach to see them

━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
*⏰ NEXT CYCLE*
━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
• In ~{60 - now.minute} minutes
"""
            keyboard = [
                [
                    InlineKeyboardButton("💼 See Jobs", callback_data="jobs"),
                    InlineKeyboardButton("📨 Outreach", callback_data="outreach")
                ],
                [InlineKeyboardButton("🔄 Refresh", callback_data="today")]
            ]
            reply_markup = InlineKeyboardMarkup(keyboard)
            
            await update.message.reply_text(message, parse_mode='Markdown', reply_markup=reply_markup)
            
        except Exception as e:
            await update.message.reply_text(f"❌ Error: {str(e)}")
    
    async def cmd_outreach(self, update: Update, context: ContextTypes.DEFAULT_TYPE):
        """Show pending LinkedIn/Twitter messages to send manually"""
        try:
            outreach = self._load_pending_outreach()
            
            if not outreach:
                message = """📨 *PENDING OUTREACH*

✅ No pending messages right now!

When the engine finds good opportunities:
• LinkedIn messages will appear here
• You copy-paste and send them
• This is the human touch that wins!

Check back after the next job cycle (~hourly).
"""
            else:
                message = f"""📨 *PENDING OUTREACH*

You have *{len(outreach)}* message(s) to send:

"""
                for i, item in enumerate(outreach[:5], 1):  # Show max 5
                    company = item.get('company', 'Unknown')
                    platform = item.get('platform', 'LinkedIn')
                    contact = item.get('contact', 'Team')
                    url = item.get('url', '#')
                    msg_preview = item.get('message', '')[:100]
                    
                    message += f"""━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
{i}. *{company}* ({platform})
👤 Contact: {contact}
🔗 {url}

📝 Message Preview:
"{msg_preview}..."
"""
                
                message += """
━━━━━━━━━━━━━━━━━━━━━━━━━━━━━

*HOW TO SEND:*
1. Click the LinkedIn URL above
2. Send connection request
3. Copy the full message from logs
4. Paste as connection note

💡 Tip: Personalize if you can!
"""
            
            keyboard = [[InlineKeyboardButton("📋 Back to Menu", callback_data="menu")]]
            reply_markup = InlineKeyboardMarkup(keyboard)
            
            await update.message.reply_text(
                message, 
                parse_mode='Markdown', 
                reply_markup=reply_markup,
                disable_web_page_preview=True
            )
            
        except Exception as e:
            await update.message.reply_text(f"❌ Error loading outreach: {str(e)}")
    
    def _load_pending_outreach(self) -> list:
        """Load pending outreach messages from the ACTUAL outreach files"""
        messages = []

        # Source 1: manual_outreach_queue.json (primary — generated by founder_finder)
        try:
            queue_file = "autonomous_data/manual_outreach_queue.json"
            if os.path.exists(queue_file):
                with open(queue_file, 'r') as f:
                    data = json.load(f)
                    items = data if isinstance(data, list) else data.get('messages', [])
                    for m in items:
                        status = m.get('status', '')
                        if status in ('pending_manual_send', 'pending', ''):
                            messages.append({
                                'company': m.get('company', 'Unknown'),
                                'platform': m.get('platform', 'LinkedIn'),
                                'contact': m.get('contact_name', m.get('contact', 'Founder')),
                                'url': m.get('linkedin_url', m.get('url', '')),
                                'message': m.get('message', m.get('content', '')),
                            })
        except Exception:
            pass

        # Source 2: outreach_log.jsonl (backup — all outreach history)
        if not messages:
            try:
                log_file = "autonomous_data/outreach_log.jsonl"
                if os.path.exists(log_file):
                    pending = []
                    with open(log_file, 'r') as f:
                        for line in f:
                            line = line.strip()
                            if not line:
                                continue
                            try:
                                entry = json.loads(line)
                                if entry.get('status') in ('pending_manual_send', 'pending'):
                                    pending.append({
                                        'company': entry.get('company', 'Unknown'),
                                        'platform': 'LinkedIn',
                                        'contact': entry.get('contact_name', 'Founder'),
                                        'url': entry.get('linkedin_url', ''),
                                        'message': entry.get('message', '')[:300],
                                    })
                            except json.JSONDecodeError:
                                continue
                    # Show most recent 10
                    messages = pending[-10:]
            except Exception:
                pass

        return messages
    
    async def cmd_help(self, update: Update, context: ContextTypes.DEFAULT_TYPE):
        """Show detailed help"""
        await update.message.reply_text(HELP_TEXT, parse_mode='Markdown')

    async def cmd_jobs(self, update: Update, context: ContextTypes.DEFAULT_TYPE):
        """Show today's top jobs"""
        try:
            # Load today's jobs from cache
            jobs = self._load_recent_jobs()
            
            if not jobs:
                await update.message.reply_text(" No jobs found today yet. Check back in an hour!")
                return
            
            # Sort by match score and take top 5
            top_jobs = sorted(jobs, key=lambda x: x.get('match_score', 0), reverse=True)[:5]
            
            message = f" *Top {len(top_jobs)} Jobs Found Today*\n\n"
            
            for i, job in enumerate(top_jobs, 1):
                company = job.get('company', 'Unknown')
                title = job.get('title', 'Unknown')
                location = job.get('location', 'Remote')
                match_score = job.get('match_score', 0) * 100
                url = job.get('url', '#')
                
                message += f"{i}. *{company}* - {title}\n"
                message += f"    {location} |  {match_score:.0f}% match\n"
                message += f"    [Apply Here]({url})\n\n"
            
            # Add quick action buttons
            keyboard = [
                [
                    InlineKeyboardButton(" Stats", callback_data="stats"),
                    InlineKeyboardButton(" Refresh", callback_data="refresh_jobs")
                ]
            ]
            reply_markup = InlineKeyboardMarkup(keyboard)
            
            await update.message.reply_text(
                message, 
                parse_mode='Markdown',
                reply_markup=reply_markup,
                disable_web_page_preview=True
            )
            
        except Exception as e:
            await update.message.reply_text(f" Error loading jobs: {str(e)}")
    
    async def cmd_stats(self, update: Update, context: ContextTypes.DEFAULT_TYPE):
        """Show application statistics"""
        try:
            if not self.db_helper:
                await update.message.reply_text(" Database not connected")
                return
            
            # Get weekly stats
            this_week = self.db_helper.get_weekly_stats(0)
            last_week = self.db_helper.get_weekly_stats(1)
            
            message = " *Your Job Hunting Stats*\n\n"
            message += "*This Week:*\n"
            message += f" Applications: {this_week['applications_sent']}\n"
            message += f" Responses: {this_week['responses_received']}\n"
            message += f" Interviews: {this_week['interviews_scheduled']}\n"
            message += f" Offers: {this_week['offers_received']}\n"
            message += f" Response Rate: {this_week['response_rate']*100:.1f}%\n\n"
            
            message += "*Last Week:*\n"
            message += f" Applications: {last_week['applications_sent']}\n"
            message += f" Responses: {last_week['responses_received']}\n\n"
            
            # Get top companies
            responsive = self.db_helper.get_responsive_companies()[:3]
            if responsive:
                message += "* Most Responsive Companies:*\n"
                for company, rate in responsive:
                    message += f"   {company} ({rate*100:.0f}% response)\n"
            
            # Get ghost companies
            ghosts = self.db_helper.get_ghost_companies()
            if ghosts:
                message += f"\n *Ghost Companies* ({len(ghosts)})\n"
                message += "These companies never respond - avoid them!\n"
            
            await update.message.reply_text(message, parse_mode='Markdown')
            
        except Exception as e:
            await update.message.reply_text(f" Error loading stats: {str(e)}")
    
    async def cmd_apply(self, update: Update, context: ContextTypes.DEFAULT_TYPE):
        """Generate materials for a specific company"""
        if not context.args:
            await update.message.reply_text(
                " Please specify a company name\n"
                "Example: /apply anthropic"
            )
            return
        
        company_name = " ".join(context.args)
        
        await update.message.reply_text(
            f" Generating materials for *{company_name}*...\n"
            "This will take ~30 seconds",
            parse_mode='Markdown'
        )
        
        try:
            # TODO: Call material generation service
            # For now, send placeholder
            message = f" *Materials Ready for {company_name}*\n\n"
            message += " Resume: Tailored for this role\n"
            message += " Cover Letter: Personalized\n"
            message += " Portfolio: Updated links\n\n"
            message += "Next: Apply manually or I can send via email"
            
            keyboard = [
                [
                    InlineKeyboardButton(" Send Email", callback_data=f"send_email_{company_name}"),
                    InlineKeyboardButton(" Save Only", callback_data="save_materials")
                ]
            ]
            reply_markup = InlineKeyboardMarkup(keyboard)
            
            await update.message.reply_text(
                message,
                parse_mode='Markdown',
                reply_markup=reply_markup
            )
            
        except Exception as e:
            await update.message.reply_text(f" Error: {str(e)}")
    
    async def cmd_pause(self, update: Update, context: ContextTypes.DEFAULT_TYPE):
        """Pause autonomous job hunting"""
        self.is_paused = True
        await update.message.reply_text(
            " *Autonomous job hunting PAUSED*\n\n"
            "I'll stop checking for new jobs.\n"
            "Use /resume to restart.",
            parse_mode='Markdown'
        )
    
    async def cmd_resume(self, update: Update, context: ContextTypes.DEFAULT_TYPE):
        """Resume autonomous job hunting"""
        self.is_paused = False
        await update.message.reply_text(
            " *Autonomous job hunting RESUMED*\n\n"
            "I'm back to checking jobs hourly!\n"
            "You'll get alerts for high-match jobs.",
            parse_mode='Markdown'
        )
    
    async def cmd_keywords(self, update: Update, context: ContextTypes.DEFAULT_TYPE):
        """Update job search keywords"""
        if not context.args:
            # Show current keywords
            keywords = self._load_keywords()
            message = " *Current Keywords:*\n" + ", ".join(keywords)
            message += "\n\nTo update: /keywords ML engineer, AI researcher"
            await update.message.reply_text(message, parse_mode='Markdown')
            return
        
        # Update keywords
        new_keywords = " ".join(context.args).split(",")
        new_keywords = [k.strip() for k in new_keywords]
        
        self._save_keywords(new_keywords)
        
        message = f" *Keywords Updated!*\n\n"
        message += "New keywords:\n" + "\n".join(f"   {k}" for k in new_keywords)
        await update.message.reply_text(message, parse_mode='Markdown')
    
    async def cmd_companies(self, update: Update, context: ContextTypes.DEFAULT_TYPE):
        """Show tracked companies"""
        try:
            # Load company list
            companies = self._load_target_companies()
            
            message = " *Tracked Companies*\n\n"
            
            for category, data in companies.items():
                if category == 'filters' or category == 'scraping_config':
                    continue
                
                company_list = data.get('companies', [])
                message += f"*{data.get('description', category)}*\n"
                message += f"  Companies: {len(company_list)}\n\n"
            
            total = sum(len(data.get('companies', [])) for key, data in companies.items() 
                       if key not in ['filters', 'scraping_config'])
            message += f" *Total: {total} companies tracked*"
            
            await update.message.reply_text(message, parse_mode='Markdown')
            
        except Exception as e:
            await update.message.reply_text(f" Error: {str(e)}")
    
    async def cmd_recent(self, update: Update, context: ContextTypes.DEFAULT_TYPE):
        """Show recent applications"""
        try:
            if not self.db_helper:
                await update.message.reply_text(" Database not connected")
                return
            
            # Placeholder - database integration coming soon
            await update.message.reply_text(" No applications tracked yet! Database integration in progress.")
            
        except Exception as e:
            await update.message.reply_text(f" Error: {str(e)}")
    
    async def cmd_status(self, update: Update, context: ContextTypes.DEFAULT_TYPE):
        """Show system status"""
        status_emoji = "" if not self.is_paused else ""
        status_text = "RUNNING" if not self.is_paused else "PAUSED"
        
        message = f" *System Status*\n\n"
        message += f"{status_emoji} Mode: *{status_text}*\n"
        message += f" Last Check: {datetime.utcnow().strftime('%H:%M UTC')}\n"
        message += f" Jobs Today: {self._count_todays_jobs()}\n"
        message += f" Next Check: In ~{60 - datetime.utcnow().minute} minutes\n\n"
        
        if self.is_paused:
            message += "Use /resume to restart"
        else:
            message += "Use /pause to stop"
        
        await update.message.reply_text(message, parse_mode='Markdown')

    async def cmd_priority(self, update: Update, context: ContextTypes.DEFAULT_TYPE):
        """Priority companies: list, add, remove, sync yc (additive — Feb 2026)"""
        if not self.db_helper:
            await update.message.reply_text("Priority list not available (no database).")
            return
        args = (context.args or []) if context.args else []
        sub = (args[0] or "list").lower() if args else "list"

        try:
            if sub == "list":
                # Sync from export file first (fresh data on manual request)
                path = self._get_priority_export_path()
                if os.path.exists(path):
                    try:
                        self.db_helper.sync_priority_from_yc_file(path)
                    except Exception:
                        pass
                rows = self.db_helper.list_priority_companies()
                total = len(rows)
                ts = rows[0]["created_at"].strftime("%Y-%m-%d %H:%M") if rows else "—"
                msg = f"🎯 *Priority Companies* ({total})\n\n"
                for r in rows[:30]:
                    msg += f"• {r['company_name']} ({r['source']})\n"
                if total > 30:
                    msg += f"\n... and {total - 30} more"
                msg += f"\n\n_Last updated: {ts}_"
                await update.message.reply_text(msg, parse_mode='Markdown')
            elif sub == "add":
                name = " ".join(args[1:]).strip() if len(args) > 1 else ""
                if not name:
                    await update.message.reply_text("Usage: /priority add <company>")
                    return
                if self.db_helper.add_priority_company(name, "manual"):
                    await update.message.reply_text(f"✅ Added: {name}")
                else:
                    await update.message.reply_text(f"Already in list or invalid: {name}")
            elif sub == "remove":
                name = " ".join(args[1:]).strip() if len(args) > 1 else ""
                if not name:
                    await update.message.reply_text("Usage: /priority remove <company>")
                    return
                if self.db_helper.remove_priority_company(name):
                    await update.message.reply_text(f"✅ Removed: {name}")
                else:
                    await update.message.reply_text(f"Not found or invalid: {name}")
            elif sub == "sync":
                sync_source = (args[1] or "yc").lower() if len(args) > 1 else "yc"
                if sync_source == "yc":
                    path = self._get_priority_export_path()
                    if not os.path.exists(path):
                        await update.message.reply_text(
                            f"📂 Export file not found at {path}. "
                            "Ask OpenClaw for your job shortlist first — it creates the file. Then run /priority sync yc again."
                        )
                        return
                    added, skipped = self.db_helper.sync_priority_from_yc_file(path)
                    await update.message.reply_text(f"✅ Sync YC: {added} added, {skipped} skipped (from {path})")
                else:
                    await update.message.reply_text("Usage: /priority sync yc")
            elif sub == "refresh":
                # Run pipeline + sync for truly fresh data on demand
                await update.message.reply_text("🔄 Refreshing from YC (30–60 sec)...")
                import subprocess
                path = self._get_priority_export_path()
                base = os.path.dirname(path)
                script = os.path.join(base, "run_shortlist.sh")
                if os.path.exists(script):
                    try:
                        subprocess.run(
                            ["bash", script],
                            cwd=base,
                            capture_output=True,
                            timeout=120,
                            env={**os.environ}
                        )
                    except subprocess.TimeoutExpired:
                        await update.message.reply_text("⏱️ Pipeline timed out. Try again.")
                        return
                    except Exception as e:
                        await update.message.reply_text(f"❌ Pipeline error: {str(e)[:100]}")
                        return
                if os.path.exists(path):
                    added, skipped = self.db_helper.sync_priority_from_yc_file(path)
                    rows = self.db_helper.list_priority_companies()
                    total = len(rows)
                    ts = rows[0]["created_at"].strftime("%Y-%m-%d %H:%M") if rows else "—"
                    msg = f"✅ Refreshed! {added} added, {skipped} skipped.\n\n🎯 *Priority Companies* ({total})\n\n"
                    for r in rows[:30]:
                        msg += f"• {r['company_name']} ({r['source']})\n"
                    if total > 30:
                        msg += f"\n... and {total - 30} more"
                    msg += f"\n\n_Last updated: {ts}_"
                    await update.message.reply_text(msg, parse_mode='Markdown')
                else:
                    await update.message.reply_text("❌ Export file not created. Check job-list-filter path.")
            else:
                await update.message.reply_text(
                    "Usage:\n/priority list\n/priority add <company>\n/priority remove <company>\n/priority sync yc\n/priority refresh"
                )
        except Exception as e:
            logger.exception("cmd_priority error")
            await update.message.reply_text(f"Error: {str(e)[:200]}")

    # ─────────────────────────────────────────────────────────────────────────
    # VOICE + FREE-FORM QUESTIONS (additive)
    # ─────────────────────────────────────────────────────────────────────────

    def _detect_language(self, text: str) -> str:
        """
        Very lightweight language detector for voice questions.
        Only distinguishes between EN / ES / RU using simple heuristics.
        """
        if not text:
            return "en"
        # Cyrillic characters → assume Russian
        if any("\u0400" <= ch <= "\u04FF" for ch in text):
            return "ru"
        lower = text.lower()
        # Basic Spanish markers
        spanish_markers = [
            " hola ",
            " buenos dias",
            " buenas tardes",
            " buenas noches",
            " gracias",
            " por favor",
            " que ",
            " porque",
            " para ",
            " contigo",
            " estoy",
            " tengo",
            " quiero",
            " trabajo",
        ]
        if any(m in lower for m in spanish_markers) or any(c in lower for c in ["ñ", "á", "é", "í", "ó", "ú", "ü"]):
            return "es"
        return "en"

    async def _translate_to_english(self, text: str, source_lang: str) -> str:
        """
        Translate ES/RU → EN through the shared provider chain (Gemini → Groq → Grok → OpenAI).
        25 Oct 2026: it was OpenAI-only, and with OpenAI out of credit since 23 Oct every Spanish or Russian question
        was silently passed on untranslated. Falls back to the original text only if every provider fails.
        """
        if not text or source_lang == "en":
            return text
        try:
            from src.utils.llm_chain import PROFILE_QUALITY, complete
            prompt = (
                "Translate the following user question into natural English, keeping all job-search-specific details. "
                "Return ONLY the translated English text, no explanations.\n\n"
                f"Original ({source_lang}): {text}"
            )
            # off the event loop: the bot shares it with the hourly job hunt (orchestrator.py create_task)
            translated, errors = await asyncio.to_thread(
                complete, [{"role": "user", "content": prompt}], 600, PROFILE_QUALITY)
            if not translated:
                logger.warning(f"Translation failed on every provider, using original text: {'; '.join(errors)[:200]}")
            return (translated or "").strip() or text
        except Exception as e:
            logger.warning(f"Translation failed, using original text. Error: {e}")
            return text

    async def handle_voice(self, update: Update, context: ContextTypes.DEFAULT_TYPE):
        """Transcribe voice message with Whisper, then answer from real job data."""
        if not update.message or not update.message.voice:
            return
        api_key = os.getenv("OPENAI_API_KEY")
        groq_key = os.getenv("GROQ_API_KEY")
        if not api_key and not groq_key:
            await update.message.reply_text(
                "🎤 Voice works! Add GROQ_API_KEY or OPENAI_API_KEY to .env for transcription. "
                "Or type your question — I answer job questions from real data."
            )
            return
        try:
            voice = update.message.voice
            tg_file = await context.bot.get_file(voice.file_id)
            tmp_path = None
            with tempfile.NamedTemporaryFile(suffix=".ogg", delete=False) as tmp:
                tmp_path = tmp.name
            await tg_file.download_to_drive(tmp_path)

            def _transcribe():
                # verbose_json: Whisper names the language it heard (16 Oct 2026). Sync SDK → run in a thread,
                # because the bot shares one event loop with the hourly job hunt. Client built as before (SDK
                # default timeout and retries), so a long voice note that used to transcribe still does.
                # 23 Oct 2026: Groq Whisper first, OpenAI as the backup. It was OpenAI only, and when OpenAI ran out of
                # credits (429 insufficient_quota) every voice note failed. Groq speaks the same API and also names
                # the language in verbose_json. Same order as CTO AIPA, EspaLuz and OpenClaw.
                from openai import OpenAI
                last_err = None
                for key, base, model in ((groq_key, "https://api.groq.com/openai/v1", "whisper-large-v3-turbo"),
                                         (api_key, None, "whisper-1")):
                    if not key:
                        continue
                    try:
                        client = OpenAI(api_key=key, base_url=base) if base else OpenAI(api_key=key)
                        with open(tmp_path, "rb") as f:
                            return client.audio.transcriptions.create(model=model, file=f,
                                                                      response_format="verbose_json")
                    except Exception as e:
                        last_err = e
                        logger.warning(f"🎤 transcription via {base or 'openai'} failed: {str(e)[:120]}")
                raise last_err or RuntimeError("no transcription provider")
            try:
                transcript = await asyncio.to_thread(_transcribe)
                text = transcript.text.strip() if hasattr(transcript, 'text') else str(transcript)
                heard = str(getattr(transcript, "language", "") or "").strip().lower()
            finally:
                if tmp_path and os.path.exists(tmp_path):
                    os.unlink(tmp_path)
            if not text:
                await update.message.reply_text("Couldn't transcribe the voice message. Try typing your question! / "
                                                "No pude transcribir la nota de voz. ¡Escribe tu pregunta!")
                return
            # Whisper's language first; the old keyword heuristic when Whisper names none, or names a language
            # outside es/en/ru (short Spanish notes come back as "galician", "catalan" or "portuguese").
            lang = _WHISPER_LANG.get(heard) or self._detect_language(text)
            logger.info(f"🎤 Voice transcribed ({lang}): {text[:100]}{'...' if len(text) > 100 else ''}")

            # 16 Oct 2026: Spanish is answered in Spanish (no translation). Russian is still normalised to English,
            # as before (she works in English and Spanish). The echo is plain text: a "_" or "*" in a transcript
            # made Markdown fail and she read "Voice transcription failed" for a formatting error.
            if lang == "ru":
                question = await self._translate_to_english(text, lang)
                echo, q_lang = f"📝 I heard (translated from RU):\n{question}", "en"
            elif lang == "es":
                question, echo, q_lang = text, f"📝 Escuché:\n{text}", "es"
            else:
                question, echo, q_lang = text, f"📝 I heard:\n{text}", ("en" if lang == "en" else None)
            # _reply_plain splits a long transcript under Telegram's 4,096 cap and never raises (it logs).
            if not await self._reply_plain(update, echo):
                logger.warning("Voice echo failed (answering anyway)")

            # Same grounded path as text questions; a voice note in → a voice reply too
            await self._reply_job_question(update, context, question, speak=True, lang=q_lang)
        except Exception as e:
            logger.warning(f"Voice transcription error: {e}")
            await update.message.reply_text(
                f"Voice transcription failed. Try typing: \"What companies did I apply to today?\" / "
                f"No pude transcribir la nota de voz. Escribe, por ejemplo: \"¿A qué empresas apliqué hoy?\""
            )

    async def handle_freeform_text(self, update: Update, context: ContextTypes.DEFAULT_TYPE):
        """Answer free-form job questions from real data (honest answers)."""
        if not update.message or not update.message.text:
            return
        text = (update.message.text or "").strip()
        if not text:
            return
        # 18 Oct 2026: a file request while a ❌ question waits is never the No-fit reason (it would move the deal to No
        # fit with "PDF" or "hazme la cotización en Word" as her reason). A swipe-reply to the question is always a
        # reason. 19 Oct: any other text within 15 minutes of ❌ IS the reason unless it is an explicit request by the
        # rule of _doc_request (a request verb opening the message with a document or format as its object, or a
        # message that is only a format) — so a short reason with no verb ("Cover letter in Word, on-site, no thanks",
        # "Reporte semanal en PDF, nada de IA") is saved, as before the document feature.
        doc_ask = (_asks_for_document(text) and getattr(update.message, "reply_to_message", None) is None)
        if not doc_ask and await self._take_nofit_reason(update, context, text):     # 18 Oct 2026: the ❌ question
            return
        await self._reply_job_question(update, context, text)

    async def _reply_job_question(self, update: Update, context: ContextTypes.DEFAULT_TYPE, text: str,
                                  speak: Optional[bool] = None, lang: Optional[str] = None):
        """Shared: answer job question and reply.

        16 Oct 2026: the answer is grounded (telegram_qa_helper) and computed in a thread, off the event loop the
        job hunt shares. It goes out as PLAIN text: a model's stray _ or * made Markdown fail and a canned line
        replaced the answer. speak=None applies THE SAME VOICE RULE to the text; handle_voice passes True.
        18 Oct 2026: a request for a PDF or Word file goes to _reply_with_document; every other message takes exactly
        the path below (its answer is also kept in memory, for a bare "PDF" that may follow)."""
        formats = _document_formats(text)
        if formats:
            await self._reply_with_document(update, context, text, formats, speak=speak, lang=lang)
            return
        try:
            from src.notifications.telegram_qa_helper import answer_job_question
            logger.info(f"📝 QA question: {text[:80]}{'...' if len(text) > 80 else ''}")
            reply = await asyncio.to_thread(answer_job_question, text, lang)
        except Exception as e:
            logger.warning(f"QA reply error: {e}")
            reply = ("Something went wrong answering that. Please ask again in a minute. / "
                     "Algo falló al responder. Vuelve a preguntar en un minuto.")
        sent = await self._reply_plain(update, reply)
        if sent:
            self._remember_answer(update, text, reply)
        want_voice = _asks_for_voice(text) if speak is None else bool(speak)
        if sent and want_voice:
            await self._speak_reply(update, reply)

    async def _reply_plain(self, update: Update, text: str) -> bool:
        """Send model- or user-written text with NO parse_mode, in pieces under Telegram's 4,096-character cap."""
        try:
            text = text or ""
            for i in range(0, max(len(text), 1), 3900):
                await update.message.reply_text(text[i:i + 3900] or "…")
            return True
        except Exception as e:
            logger.warning(f"Plain reply failed: {e}")
            return False

    async def _speak_reply(self, update: Update, text: str):
        """VJH-4 voice reply. Called only AFTER the text answer went out and fully inside try/except, so a TTS failure
        never touches the text. Off switch: VJH_SPEAK_REPLIES=0.
        25 Oct 2026: it was OpenAI tts-1 only — out of credit since 23 Oct, so every voice reply failed silently (a
        warning in the log). Now the fleet's free voice first (scripts/speak-free.py in cto-aipa: Microsoft Edge in
        sentence chunks, Google gTTS if Edge fails — the same one PrepClaw and CTO AIPA speak with), OpenAI second."""
        try:
            if os.getenv("VJH_SPEAK_REPLIES", "1").strip() == "0":
                return
            speech = _for_speech(text)
            if not speech:
                return
            audio, by = await self._free_voice_opus(speech), "free"
            api_key = os.getenv("OPENAI_API_KEY")
            if not audio and api_key:
                from openai import AsyncOpenAI
                # "async with" closes the client's connection pool after each reply (a long-running process).
                async with AsyncOpenAI(api_key=api_key, timeout=30, max_retries=1) as client:
                    # The SDK timeout covers the response headers, not the audio body, so the whole call has its own
                    # limit.
                    resp = await asyncio.wait_for(
                        client.audio.speech.create(model="tts-1", voice="nova", input=speech, response_format="opus"),
                        timeout=45)
                    audio = getattr(resp, "content", None)
                    if not isinstance(audio, (bytes, bytearray)):
                        audio = await resp.aread()
                by = "openai"
            if not audio:
                return
            from telegram import InputFile
            await update.message.reply_voice(voice=InputFile(bytes(audio), filename="reply.ogg"))
            logger.info(f"🔊 Voice reply sent ({len(speech)} characters spoken, {by} voice)")
        except Exception as e:
            logger.warning(f"🔊 Voice reply failed (the text answer already went out): {e}")

    async def _free_voice_opus(self, speech: str):
        """Ogg/Opus bytes from the fleet's free voice script, or None. Never raises."""
        import tempfile
        home = os.path.expanduser("~")
        py = os.getenv("EDGE_TTS_PY") or os.path.join(home, ".openclaw", "edge-tts-venv", "bin", "python")
        script = os.getenv("SPEAK_FREE") or os.path.join(home, "cto-aipa", "scripts", "speak-free.py")
        if not (os.path.exists(py) and os.path.exists(script)):
            return None
        try:
            with tempfile.TemporaryDirectory(prefix="vjh-voice-") as d:
                txt, mp3, ogg = os.path.join(d, "in.txt"), os.path.join(d, "out.mp3"), os.path.join(d, "out.ogg")
                with open(txt, "w", encoding="utf-8") as f:
                    f.write(speech)
                voice = os.getenv("VJH_FREE_VOICE", "en-US-AvaMultilingualNeural")
                p = await asyncio.create_subprocess_exec(py, script, "--in", txt, "--out", mp3, "--voice", voice,
                                                         "--rate=-5%", stdout=asyncio.subprocess.DEVNULL,
                                                         stderr=asyncio.subprocess.PIPE)
                _, err = await asyncio.wait_for(p.communicate(), timeout=150)
                if p.returncode != 0 or not os.path.exists(mp3):
                    logger.warning(f"🔊 Free voice failed: {(err or b'').decode(errors='replace')[-160:]}")
                    return None
                q = await asyncio.create_subprocess_exec("ffmpeg", "-y", "-loglevel", "error", "-i", mp3, "-c:a",
                                                         "libopus", "-b:a", "48k", "-ac", "1", ogg)
                await asyncio.wait_for(q.wait(), timeout=60)
                if q.returncode != 0 or not os.path.exists(ogg):
                    return None
                with open(ogg, "rb") as f:
                    return f.read()
        except Exception as e:
            logger.warning(f"🔊 Free voice failed: {e}")
            return None

    # ─────────────────────────────────────────────────────────────────────────
    # A DOCUMENT ON REQUEST (18 Oct 2026) — the recogniser is _document_formats, the file is _render_document.
    # ─────────────────────────────────────────────────────────────────────────

    def _remember_answer(self, update, question: str, answer: str, doc: Optional[dict] = None, src: str = ""):
        """The last real answer in this chat, in memory only (a restart forgets it): a bare "PDF" / "en Word" right
        after it turns THAT answer into a file without asking any model again. Canned lines are never kept.
        doc = {title, kind, base} when the answer was itself a document (so "y en Word" keeps its title and kind).
        src = what was answered when it was a screenshot or a document she sent ("job posting", "contract"…)."""
        try:
            from src.notifications.telegram_qa_helper import is_canned_reply
            if is_canned_reply(answer):
                self._forget_answer(update)        # 19 Oct: a canned line also drops the OLDER answer ("PDF" next
                return                             # gets "no recent answer", never an older, unrelated file)
            import time as _time
            chat = getattr(update, "effective_chat", None)
            self.__dict__.setdefault("_last_answers", {})[chat.id if chat else None] = {
                "q": (question or "").strip()[:300], "a": answer, "at": _time.monotonic(), "doc": doc or {},
                "src": src or ""}
        except Exception as e:
            logger.debug(f"📄 last answer not kept: {e}")

    def _forget_answer(self, update):
        """A request that made no file drops the remembered answer: a "PDF" / "en Word" right after it then gets the
        honest "no recent answer" line, never an older, unrelated answer as the file (18 Oct review)."""
        chat = getattr(update, "effective_chat", None)
        (self.__dict__.get("_last_answers") or {}).pop(chat.id if chat else None, None)

    def _last_answer(self, update) -> Optional[dict]:
        import time as _time
        chat = getattr(update, "effective_chat", None)
        last = (self.__dict__.get("_last_answers") or {}).get(chat.id if chat else None)
        return last if last and _time.monotonic() - last["at"] <= _DOC_LAST_WINDOW_S else None

    async def _doc_no_file(self, update, line: str) -> bool:
        """No file: say why, and drop the remembered answer so "PDF" next never files an older, unrelated one (18 Oct
        review). True, for the callers that report "handled"."""
        self._forget_answer(update)
        await self._reply_plain(update, line)
        return True

    def _replied_bot_text(self, update) -> str:
        """The text of the bot message she swipe-replied to, else '' (19 Oct review: "PDF" as a reply to an OLDER
        answer filed her LAST one)."""
        r = getattr(getattr(update, "message", None), "reply_to_message", None)
        text = getattr(r, "text", None)
        return text.strip() if isinstance(text, str) and getattr(getattr(r, "from_user", None), "is_bot", False) else ""

    async def _reply_with_document(self, update, context, text: str, formats: tuple, speak: Optional[bool] = None,
                                   lang: Optional[str] = None):
        """A request for a PDF / Word / HTML file. The text has one of five sources, never a sixth:
        · "letter for <company>" → that job deal's ✅ cover letter note, read-only, word for word (no model call);
        · a document about her FIT for a job, whatever its noun (an evaluation, interview prep, a report, one-pager,
          presentation, promotion or proposal for a live job deal, "por qué encajo", "pitch me to…") → that job deal's
          own reviewed kit notes, word for word (no model call) — rule 2, decided by structure (_doc_fit);
        · a client quote → her published catalog word for word, her client, her terms (no model call);
        · only "make it a file" ("PDF", "en Word", "send me that as a PDF") → the answer she got last, or the bot
          message she swipe-replied to (no model call), held to her sources like any model text;
        · any other task (an analysis, proposal, one-pager, presentation, promotion, report…) → the grounded answer
          with the document rules (one call, the usual model chain; never a job deal's notes), held to her sources by
          gate_document, shown in the chat first, then drawn into the file.
        The file goes back to this chat only and is uploaded nowhere."""
        from src.notifications import telegram_qa_helper as qa
        want_voice = _asks_for_voice(text) if speak is None else bool(speak)
        say = lang if lang in ("es", "en") else qa._guess_lang(text)
        req = _doc_request(text) or {}
        kind = req.get("kind") or _doc_kind_asked(req.get("sentence") or text)
        if kind == "answer":
            kind = _doc_kind_asked(text)
        ask_lang = _doc_lang_asked(text)
        ready = await asyncio.to_thread(_renderer_ready)
        low = _doc_fold(text)
        logger.info(f"📄 Document request ({'+'.join(formats)}, {kind}{'' if ready else ', renderer NOT installed'})")
        if kind == "answer" and _DOC_CV_ASK_RE.search(low):                 # the kit's CV lives on the deal
            await self._doc_no_file(update, _doc_say(say, "cv"))
            return
        if _doc_two_kinds(req.get("sentence") or text):
            await self._doc_no_file(update, _doc_say(say, "two_docs"))
            return
        if req.get("bare") or _is_bare_doc_request(text):
            await self._file_last_answer(update, text, formats, say, lang, want_voice, ready)
            return
        # the live-deal snapshot is a file read: off the event loop, like every other blocking step here. Companies are
        # taken from HER request sentence, never from a posting she pasted under it.
        sentence = req.get("raw") or text
        cos = await asyncio.to_thread(_doc_companies, sentence)
        if len(cos) > 1:
            # 19 Oct review: "…para Pitz, la de Clara ya la tengo" filed Clara's prep — two job deals named: ask
            await self._doc_no_file(update, _doc_say(say, "which_co", cos=" / ".join(cos)))
            return
        caps = [c for c in _doc_caps(sentence) if not cos or _doc_fold(c) not in _doc_fold(cos[0])]
        if cos and kind in _DOC_CLIENT_KINDS and not _doc_job_words(text):
            # fix 8 (wrong deal): "una propuesta para Clara" names a job deal but asks for a client document — which?
            await self._doc_no_file(update, _doc_say(say, "job_or_client", co=cos[0]))
            return
        fit = _doc_fit(text, kind, cos, caps)
        if kind == "evaluation" and not fit:
            kind = "analysis"                 # an evaluation of something else (a market, an idea): written, then gated
        if kind == "quote":
            pass                              # her catalog, word for word: no model, no fit claim possible
        elif kind == "letter":
            company = cos[0] if cos else (caps[0] if fit and len(caps) == 1 else
                                          None if fit else await asyncio.to_thread(_letter_company, sentence))
            if fit and not cos and len(caps) > 1:
                await self._doc_no_file(update, _doc_say(say, "which_co", cos=" / ".join(caps)))
                return
            if not company and (fit or _DOC_COVER_RE.search(low)):
                await self._doc_no_file(update, _doc_say(say, "cover_which"))   # never model-written
                return
            if company:
                if not ready:
                    await self._doc_no_file(update, _doc_say(say, "no_renderer", extra=""))
                    return
                if await self._send_deal_letter(update, text, company, formats, say, ask_lang, want_voice):
                    return
                if fit:                       # her fit, and no job deal holds a letter: never written here
                    await self._doc_no_file(update, _doc_say(say, "no_deal", co=company))
                    return
            # no job deal at that company, and no fit or cover letter: a letter to a client, written below
        elif fit:
            company = cos[0] if cos else (caps[0] if len(caps) == 1 else None)
            if not cos and len(caps) > 1:
                await self._doc_no_file(update, _doc_say(say, "which_co", cos=" / ".join(caps)))
            elif not company:
                await self._doc_no_file(update, _doc_say(say, "kit_which"))
            elif not ready:
                await self._doc_no_file(update, _doc_say(say, "no_renderer", extra=""))
            else:
                kit = "prep" if kind == "prep" or _DOC_FIT_INTERVIEW.search(low) else "evaluation"
                await self._send_deal_kit(update, text, company, kit, formats, say, want_voice)
            return
        today = _panama_today()
        notes, replaced = [], []
        if kind == "quote":
            # a client quote is her catalog word for word (no model): titles, prices, descriptions, her terms or blanks
            dlang = qa._doc_lang_of(text, lang, ask_lang)
            title, body, qnotes = await asyncio.to_thread(qa.quote_document, text, dlang, today)
            notes = [_doc_say(say, "quote_ok")] + qnotes
        else:
            # A task of its own: never a job deal's notes (rule 2 sent every job deal above), never their needs lines
            label = " and ".join(_DOC_FMT_LABEL[f] for f in formats)
            try:
                answer = await asyncio.to_thread(qa.answer_document_request, text, kind, lang, ask_lang, "", label)
            except Exception as e:
                logger.warning(f"📄 Document answer error: {e}")
                answer = ("Something went wrong answering that. Please ask again in a minute. / "
                          "Algo falló al responder. Vuelve a preguntar en un minuto.")
            if qa.is_canned_reply(answer):             # no model answered: that line is no document
                self._forget_answer(update)            # and "PDF" next must not file an older answer (19 Oct proof)
                sent = await self._reply_plain(update, answer)
                if sent and want_voice:
                    await self._speak_reply(update, answer)
                return
            dlang = qa._guess_lang(answer)
            title, body = qa.split_document_title(answer, kind, ask_lang or dlang)
            # THE HONESTY LINE, in code: what no source holds becomes a visible blank, a fit claim is dropped
            sources = await asyncio.to_thread(qa._gate_sources)
            title, rep_t = qa.gate_document(title, text, "", dlang, today, sources, kind)
            body, rep_b = qa.gate_document(body, text, "", dlang, today, sources, kind)
            title = title or qa._DOCREQ_TITLES.get(kind, qa._DOCREQ_TITLES["answer"])[1 if dlang == "es" else 0]
            replaced = list(dict.fromkeys(rep_t + rep_b))
            if replaced:
                items = "; ".join(f"\"{r[:60]}\"" for r in replaced[:8]) + (" …" if len(replaced) > 8 else "")
                notes.append(_doc_say(say, "gated", items=items))
        word = qa._DOCREQ_TITLES.get(kind, qa._DOCREQ_TITLES["answer"])[1 if dlang == "es" else 0]
        base = _doc_file_base(word, today)
        sent = await self._reply_plain(update, f"📄 {title}\n\n{body}")
        if sent:
            self._remember_answer(update, text, body, doc={"title": title, "kind": kind, "base": base})
        if not ready:
            items = _placeholders(f"{title}\n{body}")
            lines = [_doc_say(say, "no_renderer", extra=_doc_say(say, "no_renderer_extra"))]
            lines += [_doc_say(say, "fill", items=", ".join(items))] if items else []
            await self._reply_plain(update, "\n".join(lines + notes))
        else:
            spec = {"kind": _DOC_RENDER_KIND.get(kind, kind), "lang": dlang, "title": title,
                    "text": _doc_markdown(body, kind)}
            await self._deliver_documents(update, spec, formats, say, [], base, notes)
        if sent and want_voice:
            await self._speak_reply(update, body)

    async def _file_last_answer(self, update, text: str, formats: tuple, say: str, lang: Optional[str],
                                want_voice: bool, ready: bool):
        """Only "make it a file": the answer she got last — or, when she swipe-replies to an OLDER bot message, that
        message — drawn into the file with no model call. A letter, kit notes and a quote are filed as they are (built
        from HubSpot notes and her catalog); any other text (a chat, screenshot or document answer, a swipe-replied
        message) was written by a model, so it passes gate_document first (19 Oct review blocker: "PDF" filed "cumples
        con todos los requisitos … 12 años liderando un equipo de 40 ingenieros" word for word). An answer about her FIT
        for a job, or about a job posting she sent, is never filed: a fit document is only the deal's kit notes."""
        from src.notifications import telegram_qa_helper as qa
        last = self._last_answer(update)
        replied = self._replied_bot_text(update)
        if replied and not (last and (last["a"].strip() in replied or replied in last["a"])):
            last = None if qa.is_canned_reply(replied) else {"q": "", "a": replied, "doc": {}, "from_reply": True}
        if lang not in ("es", "en"):
            say = qa._clear_lang(text) or qa._guess_lang((last or {}).get("q") or (last or {}).get("a") or text)
        if not last:
            await self._reply_plain(update, _doc_say(say, "no_last"))
            return
        body, blang, doc = last["a"], qa._guess_lang(last["a"]), last.get("doc") or {}
        dkind = doc.get("kind") or "answer"
        notes = []
        if dkind not in ("letter", "evaluation", "prep", "quote"):
            fit_q = await asyncio.to_thread(_doc_fit_question, last.get("q") or "")   # the snapshot is a file read
            if last.get("src") == "job posting" or fit_q:
                await self._doc_no_file(update, _doc_say(say, "kit_last"))
                return
            sources = await asyncio.to_thread(qa._gate_sources)
            body, replaced = qa.gate_document(body, last.get("q") or "", "", blang, _panama_today(), sources,
                                              dkind if dkind != "answer" else "")
            if not body.strip():
                await self._doc_no_file(update, _doc_say(say, "kit_last"))
                return
            if replaced:
                items = "; ".join(f"\"{r[:60]}\"" for r in replaced[:8]) + (" …" if len(replaced) > 8 else "")
                notes.append(_doc_say(say, "gated", items=items))
        if not ready:
            await self._reply_plain(update, _doc_say(say, "no_renderer", extra=""))
            return
        spec = {"kind": _DOC_RENDER_KIND.get(dkind, dkind), "lang": blang, "text": _doc_markdown(body, dkind),
                "title": doc.get("title") or qa._DOCREQ_TITLES["answer"][1 if blang == "es" else 0]}
        q_line = ""
        if last["q"]:
            q_line = f' ("{last["q"][:80]}{"…" if len(last["q"]) > 80 else ""}")'
            if not doc:
                spec["subtitle"] = last["q"][:200]
        head = [_doc_say(say, "from_reply") if last.get("from_reply") else _doc_say(say, "from_last", q=q_line)]
        base = doc.get("base") or _doc_file_base("Respuesta" if blang == "es" else "Answer", "Elena_Revicheva",
                                                 _panama_today())
        marks_key = ("marks_letter" if dkind == "letter" else "marks_kit" if dkind in ("evaluation", "prep")
                     else "marks_last")
        await self._deliver_documents(update, spec, formats, say, head, base, notes, marks_key=marks_key)
        if want_voice:
            await self._speak_reply(update, head[0])

    def _deal_match_lines(self, say: str, prep: dict, company: str, letter: bool) -> List[str]:
        """One job deal matched → it says the file is NOT on that deal (VJH has no file-upload path to HubSpot).
        Two or more → which one was used. None → nothing."""
        n = len(prep.get("listed") or []) or (1 if prep.get("status") == "ok" else 0)
        if n == 1:
            return [_doc_say(say, "not_attached", why=_doc_say(say, "not_attached_letter") if letter else "")]
        if n > 1:
            return [_doc_say(say, "deal_used", n=n, co=company, deal=_deal_title(prep.get("name", "")))]
        return []

    async def _job_deal_for(self, update, text: str, company: str, say: str, doc: str = "letter") -> Optional[dict]:
        """hs-deal-prep (read-only) for the job deal she names, with her role words (one company can carry several job
        deals: Puente has seven). The parsed printout when it can be used, or {status: none | not_job}; None when an
        honest line already went out instead: HubSpot unreadable, a deal that shares only the role's words, or two or
        more 🔥 deals at that company that her words do not tell apart (19 Oct review: an evaluation for Puente's
        Chief of Staff role was built from its RevOps deal, and Clara SDET prep from the Growth Automation deal)."""
        role = _role_words(text, company)
        prep = _parse_deal_prep(await asyncio.to_thread(_run_deal_prep, company, role))
        if prep["status"] == "unavailable":
            await self._doc_no_file(update, _doc_say(say, "deal_unavailable"))
            return None
        if prep["status"] != "ok":
            return prep
        deal = _deal_title(prep["name"])
        co_part = deal.rpartition(" @ ")[2] or deal
        if not set(re.findall(r"\w+", _doc_fold(company))) <= set(re.findall(r"\w+", _doc_fold(co_part))):
            await self._doc_no_file(update, _doc_say(say, "wrong_co", co=company, deal=deal))
            return None
        acting = [d for d in prep["listed"] if d[1] in _DOC_ACTING]
        if len(acting) > 1:
            def hits(name):
                title = _doc_fold(_deal_title(name))
                return sum(1 for w in role.split() if w in title)
            best = max(hits(d[0]) for d in acting) if role else 0
            top = [d for d in acting if best > 0 and hits(d[0]) == best]
            if len(top) != 1 or top[0][2] != prep["id"]:
                names = "; ".join(_deal_title(d[0]) for d in acting[:6])
                example = _deal_title(acting[0][0]).rpartition(" @ ")[0] or "<role>"
                word = _DOC_EXAMPLE_WORD.get(doc, _DOC_EXAMPLE_WORD["letter"])[1 if say == "es" else 0]
                await self._doc_no_file(update, _doc_say(say, "pick_role", n=len(acting), co=co_part, names=names,
                                                         role=example, doc=word))
                return None
        return prep

    async def _send_deal_letter(self, update, text: str, company: str, formats: tuple, say: str,
                                ask_lang: Optional[str], want_voice: bool) -> bool:
        """"letter for <company>" → the deal's own ✅ cover letter note (hs-deal-prep.cjs, read-only), word for word.
        No ✅ note, two notes, a cut note, a role that only shares the word, or two 🔥 deals at that company and no
        role to tell them apart → an honest line and no file. A cover letter is never written or translated here.
        False = no job deal at that company and she did not say "cover letter / carta de presentación": a letter to
        a client, which the caller writes like any other task. True = handled here."""
        from src.notifications import telegram_qa_helper as qa
        prep = await self._job_deal_for(update, text, company, say, "letter")
        if prep is None:
            return True
        if prep["status"] != "ok":
            if not _DOC_COVER_RE.search(_doc_fold(text)):
                return False
            return await self._doc_no_file(update, _doc_say(say, "no_deal", co=company))
        deal = _deal_title(prep["name"])
        letter = _deal_letter(prep)
        if letter["status"] != "ok":
            n = sum(1 for _l, t in prep["sections"] if t.lstrip().startswith(_DOC_READY_MARK))
            key = {"none": "no_letter", "several": "several", "cut": "cut"}.get(letter["status"], "short")
            return await self._doc_no_file(update, _doc_say(say, key, deal=deal, n=n))
        llang = qa._guess_lang(letter["text"])
        if ask_lang and ask_lang != llang:
            have = {"es": ("Spanish", "español"), "en": ("English", "inglés")}[llang][1 if say == "es" else 0]
            return await self._doc_no_file(update, _doc_say(say, "letter_lang", deal=deal, have=have))
        stage = _ALERT_LABEL.get(prep["stage"], prep["stage"] or "?")
        head = [_doc_say(say, "letter_ok", deal=deal, stage=stage)]
        if say == "es" and llang == "en":
            head.append(_doc_say(say, "letter_en_only"))
        spec = {"kind": "letter", "lang": llang, "text": _letter_markdown(letter["text"]),
                "title": f"{'Carta de presentación' if llang == 'es' else 'Cover letter'} — Elena Revicheva — {deal}"}
        role_t, _, co_t = deal.rpartition(" @ ")
        base = _doc_file_base("Cover_Letter_Elena_Revicheva", role_t or deal, co_t)
        await self._deliver_documents(update, spec, formats, say, head,
                                      base, self._deal_match_lines(say, prep, company, letter=True),
                                      marks_key="marks_letter")
        # the letter is now her last answer: "y en Word también" files THIS letter, word for word, with no model
        self._remember_answer(update, text, spec["text"], doc={"title": spec["title"], "kind": "letter", "base": base})
        if want_voice:
            await self._speak_reply(update, head[0])
        return True

    async def _send_deal_kit(self, update, text: str, company: str, kind: str, formats: tuple, say: str,
                             want_voice: bool):
        """A fit document (rule 2, 19 Oct 2026): the job deal's own reviewed kit notes, word for word — never a fresh
        evaluation of her fit. "evaluación / requisito por requisito / my fit for <company>" → 🎯 ROLE DEFENSE + 🛡️
        TECHNICAL DEFENSE; interview prep → 🔎 COMPANY BRIEF + both + the 🎤 likely questions of the deep-dive. A deal
        with neither defense note, two of one, or a cut note → an honest line and no file. Notes are never translated:
        the bot's own lines are in her language."""
        from src.notifications import telegram_qa_helper as qa
        prep = await self._job_deal_for(update, text, company, say, kind)
        if prep is None:
            return
        if prep["status"] != "ok":
            await self._doc_no_file(update, _doc_say(say, "kit_no_deal", co=company))
            return
        deal = _deal_title(prep["name"])
        found = _deal_kit(prep, kind)
        es = say == "es"
        if found["status"] != "ok":
            note = _DOC_KIT_NAMES.get(found.get("several") or "", ("", ""))[1 if es else 0]
            key = {"none": "kit_none", "several": "kit_several", "cut": "kit_cut"}[found["status"]]
            await self._doc_no_file(update, _doc_say(say, key, deal=deal, note=note))
            return
        # the notes' own language — not the posting's requirements they quote ("(because: …)", "They ask: …")
        klang = qa._guess_lang("\n".join(ln for _w, t in found["parts"] for ln in t.split("\n")
                                         if not re.search(r"\(because:|They ask:", ln)))
        what = {"evaluation": ("Evaluation", "Evaluación"), "prep": ("Interview prep", "Preparación de entrevista")}[
            kind][1 if es else 0]
        used = ", ".join(_DOC_KIT_NAMES[w][1 if es else 0] for w, _t in found["parts"] if w in _DOC_KIT_NAMES)
        stage = _ALERT_LABEL.get(prep["stage"], prep["stage"] or "?")
        head = [_doc_say(say, "kit_ok", what=what, deal=deal, stage=stage, notes=used)]
        missing = [_DOC_KIT_NAMES[w][1 if es else 0] for w in found["missing"]]
        if missing:
            head.append(_doc_say(say, "kit_missing", notes=", ".join(missing)))
        if es and klang == "en":
            head.append(_doc_say(say, "kit_en_only"))
        spec = {"kind": kind, "lang": klang, "title": f"{what} — {deal}",
                "subtitle": ("Notas del kit en HubSpot, palabra por palabra — no es una evaluación nueva" if es else
                             "The kit's notes in HubSpot, word for word — not a new evaluation"),
                "text": _kit_markdown(found["parts"], say)}
        role_t, _, co_t = deal.rpartition(" @ ")
        word = {"evaluation": ("Evaluation", "Evaluacion"), "prep": ("Interview_prep", "Preparacion_entrevista")}[kind]
        base = _doc_file_base(word[1 if es else 0], role_t or deal, co_t)
        await self._deliver_documents(update, spec, formats, say, head, base,
                                      self._deal_match_lines(say, prep, company, letter=False), marks_key="marks_kit")
        self._remember_answer(update, text, spec["text"], doc={"title": spec["title"], "kind": kind, "base": base})
        if want_voice:
            await self._speak_reply(update, head[0])

    async def _deliver_documents(self, update, spec: dict, formats: tuple, say: str, head: List[str], base: str,
                                 notes: Optional[List[str]] = None, marks_key: str = "marks") -> int:
        """Render each format (in a thread), send the files to THIS chat as documents, then one status: what to fill
        in, emoji left out, any refusal, the deal lines. Returns how many files went out. marks_key: the refusal line
        for verdict marks — "marks" only where a model writes the text (asking again rewrites them in words),
        "marks_letter" / "marks_last" / "marks_kit" where nothing rewrites it. HTML is never refused for a mark: a
        browser draws it (the renderer keeps them in HTML)."""
        if not self._chat_allowed(update):                # the handlers are locked; this is the same lock once more
            return 0
        marks = sorted(set(_DOC_VERDICT_MARKS.findall(" ".join(str(spec.get(k) or "")
                                                                for k in ("title", "subtitle", "text")))))
        lines, files, places, dropped = [], [], [], 0
        blocked = [f for f in formats if f != "html"] if marks else []
        if blocked:
            lines.append(_doc_say(say, marks_key, fmt=" / ".join(_DOC_FMT_NAME[f] for f in blocked),
                                  marks=" ".join(marks)))
        for fmt in [f for f in formats if f not in blocked]:
            fname = _DOC_FMT_NAME[fmt]
            try:
                res = await asyncio.wait_for(asyncio.to_thread(_render_document, spec, fmt, base),
                                             timeout=_DOC_RENDER_TIMEOUT_S + 30)
            except asyncio.TimeoutError:
                res = {"ok": False, "problem": "timeout", "detail": ""}
            except Exception as e:
                res = {"ok": False, "problem": "refused", "detail": str(e)}
            if not res.get("ok"):
                logger.warning(f"📄 {fname} not made: {res.get('problem')} {_doc_clean_error(res.get('detail', ''))}")
                if res.get("problem") == "no_renderer":
                    lines.append(_doc_say(say, "no_renderer", extra=""))
                elif res.get("problem") == "no_html":
                    lines.append(_doc_say(say, "no_html"))
                else:
                    lines.append(_doc_say(say, res.get("problem", "refused"), fmt=fname,
                                          why=_doc_why(say, res.get("detail", ""))))
                continue
            files.append((res["name"], res["bytes"]))
            places += [p for p in res["placeholders"] if p not in places]
            dropped = max(dropped, res["emoji_dropped"])
        if places:
            lines.append(_doc_say(say, "fill", items=", ".join(places)))
        if dropped:
            lines.append(_doc_say(say, "emoji", n=dropped))
        lines += [ln for ln in (notes or []) if ln]
        if files:                                         # "your letter / your last answer as a file" only with a file
            lines = list(head) + lines
        status = "\n".join(ln for ln in lines if ln)
        from telegram import InputFile
        sent_files = 0
        for i, (name, data) in enumerate(files):
            cap = status if (i == len(files) - 1 and status and len(status) <= 1000) else None
            try:
                await update.message.reply_document(document=InputFile(data, filename=name), caption=cap)
                sent_files += 1
                logger.info(f"📄 Document sent: {name} ({len(data)} bytes)")
                if cap:
                    status = ""
            except Exception as e:
                logger.warning(f"📄 Sending {name} failed: {e}")
                status = (status + "\n" if status else "") + (
                    f"📄 No pude enviar {name} por Telegram." if say == "es" else f"📄 Telegram did not take {name}.")
        if status:
            await self._reply_plain(update, status)
        return sent_files

    async def handle_photo(self, update: Update, context: ContextTypes.DEFAULT_TYPE):
        """VJH-3 (16 Oct 2026): a screenshot → its text is read by the vision functions judge_feedback_sync already
        runs in production → the grounded answer (never the keyword router). Never silent: every failure says so."""
        msg = update.message
        if not msg:
            return
        caption = (msg.caption or "").strip()
        try:
            item = msg.photo[-1] if msg.photo else msg.document
            if item is None:
                return
            if (getattr(item, "file_size", 0) or 0) > _SCREENSHOT_MAX_BYTES:
                await msg.reply_text("📸 That image is over 6 MB, too big to read. Please send a smaller screenshot. / "
                                     "La imagen pesa más de 6 MB. Envía una captura más pequeña.")
                return
            tg_file = await context.bot.get_file(item.file_id)
            raw = bytes(await tg_file.download_as_bytearray())
        except Exception as e:
            logger.warning(f"📸 Screenshot download failed: {e}")
            await msg.reply_text("📸 I could not download that image. Please send it again. / "
                                 "No pude descargar la imagen. Envíala otra vez.")
            return
        mime = _image_mime(raw)
        if not mime or len(raw) > _SCREENSHOT_MAX_BYTES:
            await msg.reply_text("📸 I can read PNG, JPEG, WebP or GIF screenshots under 6 MB. An iPhone HEIC file: "
                                 "send it as a photo instead. / Leo capturas PNG, JPEG, WebP o GIF de menos de 6 MB; "
                                 "un HEIC de iPhone, envíalo como foto.")
            return
        read = await self._read_screenshot_text(raw, mime)
        if not read:
            await msg.reply_text("📸 I could not read that screenshot just now (the image reader did not answer). "
                                 "Please try again in a minute. / No pude leer la captura ahora; inténtalo en un minuto.")
            return
        kind, extracted, cut = _split_screenshot_read(read)
        logger.info(f"📸 Screenshot read ({kind}, {len(extracted)} chars{', maybe cut' if cut else ''}); "
                    f"caption: {caption[:60]}")
        try:
            from src.notifications.telegram_qa_helper import answer_screenshot
            reply = await asyncio.to_thread(answer_screenshot, caption, kind, extracted, cut)
        except Exception as e:
            logger.warning(f"📸 Screenshot answer error: {e}")
            reply = ("Something went wrong answering that screenshot. Please send it again in a minute. / "
                     "Algo falló al responder la captura. Envíala otra vez en un minuto.")
        # Only when the read is empty, or short AND cut/unlabelled — a short, correctly read interviewer question
        # ("Tell me about yourself") gets no note.
        if msg.photo and (not extracted or (len(extracted) < 80 and (cut or kind == "other"))):
            reply += ("\n\n(📸 Very little text came through. Telegram shrinks photos; send the screenshot as a "
                      "file for full resolution. / Llegó muy poco texto: envía la captura como archivo.)")
        sent = await self._reply_plain(update, reply)
        if sent:                                                   # 18 Oct 2026: "PDF" next makes it a file
            self._remember_answer(update, caption, reply, src=kind)  # (never a posting's fit answer: 19 Oct review)
        if sent and _asks_for_voice(caption):
            await self._speak_reply(update, reply)

    async def _read_screenshot_text(self, raw: bytes, mime: str) -> str:
        """OpenAI vision first, Gemini second (judge_feedback_sync._vision_openai / _vision_gemini with this bot's
        prompt). Both are blocking urllib with retries, so each runs in a thread with an overall limit."""
        import base64
        try:
            jfs = _judge_sync_module()
        except Exception as e:
            logger.warning(f"📸 Screenshot reader unavailable: {e}")
            return ""
        b64 = base64.b64encode(raw).decode()
        errs = []
        for name, fn, key in (("openai", jfs._vision_openai, os.getenv("OPENAI_API_KEY", "")),
                              ("gemini", jfs._vision_gemini, os.getenv("GEMINI_API_KEY", ""))):
            if not key:
                errs.append(f"{name}: no key")
                continue
            try:
                txt = await asyncio.wait_for(
                    asyncio.to_thread(fn, key, mime, b64, _SCREENSHOT_PROMPT, 1500), timeout=150)
            except Exception as e:
                errs.append(f"{name}: {str(e)[:100]}")
                continue
            if txt and txt.strip():
                return txt
            errs.append(f"{name}: empty")
        logger.warning(f"📸 Screenshot read failed: {'; '.join(errs)}")
        return ""

    async def handle_document(self, update: Update, context: ContextTypes.DEFAULT_TYPE):
        """16 Oct 2026 (Elena: "like Chalán in Pitz"): a PDF or Word .docx → its text layer is read here (PyPDF2 /
        python-docx, in a thread) → the same grounded answer as a screenshot (answer_document). Never silent, never a
        guess: a scan with no text layer, a password, a file over 20 MB or another type (.doc, .xlsx…) each get an
        honest line. An image sent as a file never reaches here: handle_photo is registered first."""
        msg = update.message
        if not msg or not msg.document:
            return
        doc = msg.document
        caption = (msg.caption or "").strip()
        file_name = doc.file_name or ""
        ext = os.path.splitext(file_name)[1].lower()
        name = re.sub(r'[\r\n"]+', " ", file_name).strip()[:80]           # only for showing it
        label, label_es = (f'"{name}"', f'"{name}"') if name else ("that file", "ese archivo")
        gate = _document_gate(file_name, doc.mime_type or "")
        if gate != "read":
            await self._reply_plain(update, _document_problem_reply(gate, label, label_es, ext))
            return
        if (doc.file_size or 0) > _DOC_MAX_BYTES:
            await self._reply_plain(update, _document_problem_reply("too_big", label, label_es, ext))
            return
        try:
            tg_file = await context.bot.get_file(doc.file_id)
            raw = bytes(await tg_file.download_as_bytearray())
        except Exception as e:
            logger.warning(f"📄 Document download failed: {e}")
            await self._reply_plain(update, _document_problem_reply("download", label, label_es, ext))
            return
        try:
            res = await asyncio.wait_for(asyncio.to_thread(_read_document, raw), timeout=90)
        except asyncio.TimeoutError:
            res = {"status": "timeout"}
        except Exception as e:
            logger.warning(f"📄 Document read error: {e}")
            res = {"status": "unreadable"}
        if res.get("status") != "ok":
            logger.info(f"📄 Document not read: {res.get('status')} ({ext or 'no extension'}, {len(raw)} bytes)")
            await self._reply_plain(update, _document_problem_reply(
                res.get("status", ""), label, label_es, ext, res.get("pages", 0), res.get("pages_read", 0)))
            return
        kind = "other"
        try:
            from src.notifications.telegram_qa_helper import answer_document, _document_kind
            kind = _document_kind(res["text"], file_name)
            what = (f"a PDF, {res['pages']} page{'s' if res['pages'] != 1 else ''}" if res["fmt"] == "pdf"
                    else "a Word .docx") + (f', file name "{name}"' if name else "")
            logger.info(f"📄 Document read ({res['fmt']}, {kind}, {len(res['text'])} chars"
                        f"{', cut' if res['cut'] else ''}); caption: {caption[:60]}")
            reply = await asyncio.to_thread(answer_document, caption, kind, res["text"], res["cut"], what)
        except Exception as e:
            logger.warning(f"📄 Document answer error: {e}")
            reply = ("Something went wrong answering that document. Please send it again in a minute. / "
                     "Algo falló al responder el documento. Envíalo otra vez en un minuto.")
        reply += _document_notes(res, label)
        sent = await self._reply_plain(update, reply)
        if sent:                                                   # 18 Oct 2026: "PDF" next makes it a file
            self._remember_answer(update, caption, reply, src=kind)
        if sent and _asks_for_voice(caption):
            await self._speak_reply(update, reply)

    async def handle_callback(self, update: Update, context: ContextTypes.DEFAULT_TYPE):
        """Handle inline button clicks"""
        import logging
        logger = logging.getLogger(__name__)
        
        # Debug: Print to stdout AND log
        print(f"=== CALLBACK RECEIVED: {update.callback_query.data if update.callback_query else 'None'} ===")
        
        query = update.callback_query
        logger.info(f"🔘 Callback received: {query.data}")
        if not self._chat_allowed(update):
            logger.warning(f"🔒 Callback '{query.data}' from another chat ignored (chat lock)")
            return
        
        try:
            await query.answer()
            chat_id = query.message.chat_id
            logger.info(f"🔘 Processing callback '{query.data}' for chat {chat_id}")
            
            # ═══════════════════════════════════════════════════════════════
            # 🆕 NEW MENU CALLBACKS (January 2026)
            # ═══════════════════════════════════════════════════════════════
            if query.data == "menu":
                await self._send_menu(context, chat_id)
            
            elif query.data == "workflow":
                await self._send_workflow(context, chat_id)
            
            elif query.data == "manual":
                await self._send_manual(context, chat_id)
            
            elif query.data == "today":
                await self._send_today(context, chat_id)
            
            elif query.data == "outreach":
                await self._send_outreach(context, chat_id)
            
            elif query.data == "jobs":
                await self._send_jobs(context, chat_id)
            
            elif query.data == "pause":
                self.is_paused = True
                await context.bot.send_message(chat_id, "⏸️ Job hunting PAUSED. Use /resume to restart.")
            
            elif query.data == "resume_hunting":
                self.is_paused = False
                await context.bot.send_message(chat_id, "▶️ Job hunting RESUMED!")
            
            elif query.data == "status":
                await self._send_status(context, chat_id)
            
            elif query.data == "stats":
                await self._send_stats(context, chat_id)
            
            elif query.data == "priority":
                await self._send_priority_list(context, chat_id)
            
            elif query.data == "priority_sync":
                await self._do_priority_sync(context, chat_id)
            
            elif query.data == "priority_refresh":
                await self._do_priority_refresh(context, chat_id)
            
            # ═══════════════════════════════════════════════════════════════
            # 🚀 CMO AIPA CALLBACKS (LinkedIn + Instagram via Make.com)
            # ═══════════════════════════════════════════════════════════════
            elif query.data == "post_linkedin":
                # Use the actual CMO code instead of direct webhook
                try:
                    from src.notifications.linkedin_cmo_v4 import LinkedInCMO
                    cmo = LinkedInCMO()
                    success = await cmo.post_to_linkedin(post_type="random", language="random")
                    if success:
                        await context.bot.send_message(chat_id, "✅ LinkedIn + Instagram posting triggered via CMO!")
                    else:
                        await context.bot.send_message(chat_id, "❌ CMO posting failed")
                except Exception as e:
                    await context.bot.send_message(chat_id, f"❌ Error: {str(e)}")
            
            elif query.data == "post_instagram":
                import requests
                webhook_url = os.getenv('MAKE_WEBHOOK_URL_LINKEDIN', '')
                if not webhook_url:
                    await context.bot.send_message(chat_id, "❌ MAKE_WEBHOOK_URL_LINKEDIN not set in env")
                    return
                try:
                    response = requests.post(webhook_url, json={
                        "platform": "instagram",
                        "action": "trigger_posting",
                        "timestamp": datetime.utcnow().isoformat()
                    })
                    await context.bot.send_message(chat_id, "✅ Instagram posting triggered! Check Make.com scenario.")
                except Exception as e:
                    await context.bot.send_message(chat_id, f"❌ Error: {str(e)}")
            
            elif query.data in ["analytics", "campaign", "generate_image", "header_job", "header_cmo", "header_system"]:
                # Placeholder handlers for CMO features
                await context.bot.send_message(chat_id, f"🚧 Feature '{query.data}' coming soon!")
            
            # ═══════════════════════════════════════════════════════════════
            # ORIGINAL CALLBACKS
            # ═══════════════════════════════════════════════════════════════
            elif query.data == "refresh_jobs":
                await query.edit_message_text("🔄 Refreshing...")
                await self._send_jobs(context, chat_id)
            
            elif query.data.startswith("send_email_"):
                company = query.data.replace("send_email_", "")
                await query.edit_message_text(f"📧 Sending email to {company}...")
            
            elif query.data == "save_materials":
                await query.edit_message_text("💾 Materials saved!")
            
            # 18 Oct 2026: ⏳ I applied / ❌ Not a fit under a job alert (prefixes: telegram_notifier.ALERT_CB_*)
            elif query.data.startswith(("vjh_ap:", "vjh_nf:")):
                await self._alert_button(query, context)

            else:
                logger.warning(f"⚠️ Unknown callback: {query.data}")
                
        except Exception as e:
            logger.error(f"❌ Callback error: {e}", exc_info=True)
            try:
                await context.bot.send_message(query.message.chat_id, f"❌ Error: {str(e)[:100]}")
            except:
                pass
    
    # ── 18 Oct 2026: the two job-alert buttons (module notes above _ALERT_SENT) ──────────────────────────────
    def _nofit_waiting(self) -> dict:
        """{chat id: {deal id: the ❌ question waiting for her reason}}. One per deal, so ❌ on two alerts keeps
        both. In memory: a restart forgets an unanswered question."""
        return self.__dict__.setdefault("_nofit_waiting_by_chat", {})

    async def _drop_alert_buttons(self, context, chat_id, message_id):
        """The buttons leave the alert once the deal has moved; the alert's text stays as it was."""
        if not message_id:
            return
        try:
            await context.bot.edit_message_reply_markup(chat_id=chat_id, message_id=message_id, reply_markup=None)
        except Exception as e:
            logger.info(f"🔘 alert buttons left in place ({str(e)[:80]})")

    async def _alert_button(self, query, context):
        """⏳ I applied → note + ⏳ Sent at once. ❌ Not a fit → ask for her line (ForceReply), saved by
        _take_nofit_reason. callback_data = prefix + HubSpot deal id (telegram_notifier.alert_button_rows)."""
        from src.notifications.telegram_notifier import ALERT_CB_APPLIED
        chat_id = query.message.chat_id
        alert_id = query.message.message_id
        applied = query.data.startswith(ALERT_CB_APPLIED)
        deal = query.data.split(":", 1)[1]
        if not (deal.isascii() and deal.isdigit()):
            await context.bot.send_message(chat_id, "This button carries no deal id. Nothing changed.")
            return
        if applied:      # ⏳ withdraws an open ❌ question on the same job, so it cannot swallow her next message
            self._nofit_waiting().get(chat_id, {}).pop(deal, None)
        try:
            stage, name, n_notes = await asyncio.to_thread(_alert_deal, deal)
        except Exception as e:
            logger.warning(f"🔘 alert button: HubSpot read failed ({str(e)[:120]})")
            await context.bot.send_message(chat_id, _hs_problem(e))
            return
        name = name or "this job"
        target = _ALERT_SENT if applied else _ALERT_LOST
        if stage == target:
            await self._drop_alert_buttons(context, chat_id, alert_id)
            await context.bot.send_message(
                chat_id, f"{'⏳ Already in Sent' if applied else '❌ Already in No fit'}. Nothing changed: {name}")
            return
        if stage not in _ALERT_OPEN:
            await context.bot.send_message(
                chat_id, f"Left as it is: {name} is in {_ALERT_LABEL.get(stage, stage or 'an unknown stage')}. "
                         f"Move it in HubSpot if you mean to.")
            return
        if applied:
            try:
                result = await asyncio.to_thread(
                    _alert_write, deal, f"Applied (via VJH alert button, {_panama_today()})", _ALERT_SENT)
            except Exception as e:
                logger.warning(f"🔘 ⏳ button: HubSpot write failed ({str(e)[:120]})")
                await context.bot.send_message(chat_id, _hs_problem(e))
                return
            if result == "ok":
                logger.info(f"🔘 ⏳ moved to Sent: {name}")
                await self._drop_alert_buttons(context, chat_id, alert_id)
                await context.bot.send_message(chat_id, f"⏳ moved to Sent: {name}")
            else:
                await context.bot.send_message(
                    chat_id, f"⚠️ Note saved, but HubSpot did not move the deal. Move {name} to ⏳ Sent by hand.")
            return
        from telegram import ForceReply
        goes = ("It goes on the HubSpot deal and VJH learns from it." if n_notes < _LEARNER_NOTE_CAP else
                f"It goes on the HubSpot deal. This deal already has {n_notes} notes and VJH's learner reads only "
                f"{_LEARNER_NOTE_CAP} per deal, so it may not see this one.")
        ask = await context.bot.send_message(
            chat_id,
            f"❌ Not a fit: {name}\nWhy? Reply with one line in your own words. {goes} "
            f"Type cancel to keep the job open.",
            reply_markup=ForceReply(input_field_placeholder="Why it is not a fit"))
        import time as _time
        self._nofit_waiting().setdefault(chat_id, {})[deal] = {
            "deal": deal, "name": name, "alert_id": alert_id,
            "ask_id": getattr(ask, "message_id", None), "at": _time.monotonic()}

    async def _take_nofit_reason(self, update, context, text: str) -> bool:
        """After ❌: her reply to the question (24 h) or to its alert (15 min), or her next plain text in this chat
        while exactly one ❌ question is under 15 min old, is the reason for THAT job.
        True = handled here. False = an ordinary message, which goes to the Q&A exactly as before."""
        chat = getattr(update, "effective_chat", None)
        chat_id = chat.id if chat else None
        waiting = (self.__dict__.get("_nofit_waiting_by_chat") or {}).get(chat_id)
        if not waiting:
            return False
        import time as _time
        now = _time.monotonic()
        for d in [d for d, q in waiting.items() if now - q["at"] > _NOFIT_REPLY_WINDOW_S]:
            waiting.pop(d, None)
        fresh = [q for q in waiting.values() if now - q["at"] <= _NOFIT_PLAIN_WINDOW_S]
        reply_to = getattr(update.message, "reply_to_message", None)
        if reply_to is not None:
            rid = getattr(reply_to, "message_id", None)
            p = next((q for q in waiting.values() if q.get("ask_id") is not None and q["ask_id"] == rid), None) \
                or next((q for q in fresh if q.get("alert_id") is not None and q["alert_id"] == rid), None)
        else:
            if len(fresh) > 1:          # which job? Ask, keep both open, save nothing
                await self._reply_plain(update, "Two or more jobs are waiting for a reason: "
                                                + "; ".join(q["name"] for q in fresh)
                                                + ". Swipe-reply to the question of the job you mean. Nothing saved.")
                return True
            p = fresh[0] if fresh else None
        if not p:
            return False
        waiting.pop(p["deal"], None)
        reason = (text or "").strip()[:_NOFIT_MAX]
        if reason.lower().strip(" .!") in _NOFIT_CANCEL:
            await self._reply_plain(update, f"Kept open. Nothing changed: {p['name']}")
            return True
        try:
            stage, name, _n = await asyncio.to_thread(_alert_deal, p["deal"])
        except Exception as e:
            logger.warning(f"🔘 ❌ reason: HubSpot read failed ({str(e)[:120]})")
            if getattr(e, "code", None) != 404:
                waiting[p["deal"]] = {**p, "at": _time.monotonic()}     # the same line can simply be sent again
            await self._reply_plain(update, _hs_problem(e, "Nothing saved.", "Send the reason again in a minute."))
            return True
        name = name or p["name"]
        if stage != _ALERT_LOST and stage not in _ALERT_OPEN:
            await self._reply_plain(update, f"Left as it is: {name} is in "
                                            f"{_ALERT_LABEL.get(stage, stage or 'an unknown stage')}. Nothing saved.")
            return True
        import html as _html
        note = _html.escape(reason, quote=False)
        try:
            result = await asyncio.to_thread(_alert_write, p["deal"], note,
                                             _ALERT_LOST if stage in _ALERT_OPEN else None)
        except Exception as e:
            logger.warning(f"🔘 ❌ reason: HubSpot write failed ({str(e)[:120]})")
            waiting[p["deal"]] = {**p, "at": _time.monotonic()}
            await self._reply_plain(update, _hs_problem(e, "Nothing saved.", "Send the reason again in a minute."))
            return True
        lines = [f"❌ saved: {reason}"]
        if result == "ok":
            logger.info(f"🔘 ❌ moved to No fit with her reason: {name}")
            lines.append(f"Moved to No fit: {name}" if stage in _ALERT_OPEN else f"Already in No fit: {name}")
            await self._drop_alert_buttons(context, chat_id, p.get("alert_id"))
        else:
            lines.append(f"⚠️ HubSpot did not move the deal. Move {name} to ❌ No fit by hand.")
        try:
            jfs = _judge_sync_module()
            if not jfs._rejection_reason([{"body": note}], jfs._clean_title(name)):
                lines.append("VJH's learner will skip this one (it needs three words or more, and not a pasted "
                             "posting). Add a sentence on the deal if it matters.")
        except Exception:
            pass
        await self._reply_plain(update, "\n".join(lines))
        return True

    # ═══════════════════════════════════════════════════════════════════════════
    # CALLBACK HELPER METHODS - Send messages directly via bot
    # ═══════════════════════════════════════════════════════════════════════════
    
    async def _send_menu(self, context, chat_id):
        """Send menu via callback — the same screen as /menu"""
        await context.bot.send_message(chat_id, MENU_TEXT, parse_mode='Markdown', reply_markup=_menu_keyboard())

    async def _send_workflow(self, context, chat_id):
        """Send workflow explanation via callback — the same screen as /workflow"""
        keyboard = [[InlineKeyboardButton("👤 What is mine", callback_data="manual")],
                    [InlineKeyboardButton("📋 Back to Menu", callback_data="menu")]]
        await context.bot.send_message(chat_id, WORKFLOW_TEXT, parse_mode='Markdown', reply_markup=InlineKeyboardMarkup(keyboard))

    async def _send_manual(self, context, chat_id):
        """Send manual tasks via callback — the same screen as /manual"""
        keyboard = [[InlineKeyboardButton("📋 Back to Menu", callback_data="menu")]]
        await context.bot.send_message(chat_id, MANUAL_TEXT, parse_mode='Markdown', reply_markup=InlineKeyboardMarkup(keyboard))

    async def _send_today(self, context, chat_id):
        """Send today's summary via callback"""
        jobs_count = self._count_todays_jobs()
        outreach = self._load_pending_outreach()
        now = datetime.utcnow()
        panama_hour = (now.hour - 5) % 24
        
        message = f"""📊 *TODAY'S SUMMARY*
📅 {now.strftime('%B %d, %Y')} | {panama_hour}:{now.strftime('%M')} Panama

*🔍 DISCOVERY*
• Jobs Found: {jobs_count}
• Companies: 200+

*📨 YOUR ACTION*
• LinkedIn Messages: {len(outreach)} pending

*⏰ NEXT CYCLE*
• In ~{60 - now.minute} minutes
"""
        keyboard = [
            [InlineKeyboardButton("💼 Jobs", callback_data="jobs"),
             InlineKeyboardButton("📨 Outreach", callback_data="outreach")],
            [InlineKeyboardButton("📋 Menu", callback_data="menu")]
        ]
        reply_markup = InlineKeyboardMarkup(keyboard)
        
        await context.bot.send_message(chat_id, message, parse_mode='Markdown', reply_markup=reply_markup)
    
    async def _send_outreach(self, context, chat_id):
        """Send pending outreach via callback"""
        outreach = self._load_pending_outreach()
        
        if not outreach:
            message = """📨 *PENDING OUTREACH*

✅ No pending messages right now!

The engine generates 1-2 new outreach per day.
Check back after the next job cycle.
"""
        else:
            message = f"""📨 *PENDING OUTREACH*

You have *{len(outreach)}* message(s) to send:

"""
            for i, m in enumerate(outreach[:5], 1):
                company = m.get('company', 'Unknown')
                contact = m.get('contact', 'Founder')
                platform = m.get('platform', 'LinkedIn')
                url = m.get('url', '')
                message += f"{i}. *{company}* → {contact}\n"
                if url:
                    message += f"   🔗 {url}\n"
                message += f"   via {platform}\n\n"

            if len(outreach) > 5:
                message += f"_...and {len(outreach) - 5} more_\n\n"
            message += "Use /outreach for full messages to copy-paste."
        
        keyboard = [[InlineKeyboardButton("📋 Menu", callback_data="menu")]]
        reply_markup = InlineKeyboardMarkup(keyboard)
        
        await context.bot.send_message(chat_id, message, parse_mode='Markdown', reply_markup=reply_markup)
    
    async def _send_jobs(self, context, chat_id):
        """Send jobs list via callback"""
        jobs = self._load_recent_jobs()
        
        if not jobs:
            message = "💼 No jobs found today yet. Check back in an hour!"
        else:
            top_jobs = sorted(jobs, key=lambda x: x.get('match_score', 0), reverse=True)[:5]
            message = f"💼 *Top {len(top_jobs)} Jobs Today*\n\n"
            
            for i, job in enumerate(top_jobs, 1):
                company = job.get('company', 'Unknown')
                title = job.get('title', 'Unknown')[:30]
                score = job.get('match_score', 0)
                if isinstance(score, float) and score < 1:
                    score = score * 100
                message += f"{i}. *{company}*\n   {title} ({score:.0f}%)\n\n"
        
        keyboard = [[InlineKeyboardButton("📋 Menu", callback_data="menu")]]
        reply_markup = InlineKeyboardMarkup(keyboard)
        
        await context.bot.send_message(chat_id, message, parse_mode='Markdown', reply_markup=reply_markup)
    
    async def _send_status(self, context, chat_id):
        """Send status via callback"""
        status = "RUNNING ✅" if not self.is_paused else "PAUSED ⏸️"
        now = datetime.utcnow()
        
        message = f"""⚙️ *SYSTEM STATUS*

• Mode: *{status}*
• Time: {now.strftime('%H:%M')} UTC
• Jobs Today: {self._count_todays_jobs()}
• Next Cycle: ~{60 - now.minute} min
"""
        keyboard = [[InlineKeyboardButton("📋 Menu", callback_data="menu")]]
        reply_markup = InlineKeyboardMarkup(keyboard)
        
        await context.bot.send_message(chat_id, message, parse_mode='Markdown', reply_markup=reply_markup)
    
    async def _send_stats(self, context, chat_id):
        """Send stats via callback — reads real data from local files"""
        jobs_today = self._count_todays_jobs()
        outreach = self._load_pending_outreach()

        # Count total outreach ever sent from log
        total_outreach = 0
        sent_outreach = 0
        try:
            log_file = "autonomous_data/outreach_log.jsonl"
            if os.path.exists(log_file):
                with open(log_file, 'r') as f:
                    for line in f:
                        line = line.strip()
                        if not line:
                            continue
                        try:
                            entry = json.loads(line)
                            total_outreach += 1
                            if entry.get('status') not in ('pending_manual_send', 'pending'):
                                sent_outreach += 1
                        except json.JSONDecodeError:
                            continue
        except Exception:
            pass

        # Count seen jobs (dedup cache)
        seen_jobs = 0
        try:
            seen_file = "autonomous_data/seen_jobs.json"
            if os.path.exists(seen_file):
                with open(seen_file, 'r') as f:
                    data = json.load(f)
                    seen_jobs = len(data) if isinstance(data, (list, dict)) else 0
        except Exception:
            pass

        message = f"""📈 *QUICK STATS*

*🔍 Jobs:*
   Today: {jobs_today}
   Total seen: {seen_jobs}

*📨 Outreach:*
   Pending to send: {len(outreach)}
   Total generated: {total_outreach}
   Completed: {sent_outreach}

_Use /stats for full DB stats (if connected)_
"""
        keyboard = [[InlineKeyboardButton("📊 Today", callback_data="today")],
                    [InlineKeyboardButton("📋 Menu", callback_data="menu")]]
        reply_markup = InlineKeyboardMarkup(keyboard)
        
        await context.bot.send_message(chat_id, message, parse_mode='Markdown', reply_markup=reply_markup)
    
    def _get_priority_export_path(self):
        """Resolve path to priority_companies_for_vibejob.json (same logic as sync)."""
        path = os.getenv("PRIORITY_YC_EXPORT_PATH")
        if not path:
            base = os.getenv("OPENCLAW_JOB_LIST_PATH", ".")
            path = os.path.join(base, "priority_companies_for_vibejob.json")
        if not os.path.isabs(path):
            path = os.path.join(os.getcwd(), path)
        if not os.path.exists(path) and os.path.exists("/home/ubuntu/job-list-filter"):
            path = "/home/ubuntu/job-list-filter/priority_companies_for_vibejob.json"
        return path

    async def _send_priority_list(self, context, chat_id):
        """Send priority companies list via menu callback. Syncs from JSON first for fresh data."""
        if not self.db_helper:
            await context.bot.send_message(chat_id, "Priority list not available (no database).")
            return
        try:
            # Sync from export file first (fresh data on manual click)
            path = self._get_priority_export_path()
            if os.path.exists(path):
                try:
                    self.db_helper.sync_priority_from_yc_file(path)
                except Exception:
                    pass
            rows = self.db_helper.list_priority_companies()
            total = len(rows)
            ts = rows[0]["created_at"].strftime("%Y-%m-%d %H:%M") if rows else "—"
            msg = f"🎯 *Priority Companies* ({total})\n\n"
            for r in rows[:30]:
                msg += f"• {r['company_name']} ({r['source']})\n"
            if total > 30:
                msg += f"\n... and {total - 30} more"
            msg += f"\n\n_Last updated: {ts}_"
            msg += "\n\n_Add/remove:_ /priority add or remove <name>"
            keyboard = [
                [
                    InlineKeyboardButton("🔄 Sync YC", callback_data="priority_sync"),
                    InlineKeyboardButton("🔄 Refresh", callback_data="priority_refresh"),
                ],
                [InlineKeyboardButton("📋 Back to Menu", callback_data="menu")],
            ]
            reply_markup = InlineKeyboardMarkup(keyboard)
            await context.bot.send_message(chat_id, msg, parse_mode='Markdown', reply_markup=reply_markup)
        except Exception as e:
            logger.exception("_send_priority_list error")
            await context.bot.send_message(chat_id, f"Error: {str(e)[:200]}")
    
    async def _do_priority_sync(self, context, chat_id):
        """Sync from YC export file (callback from button)."""
        if not self.db_helper:
            await context.bot.send_message(chat_id, "Priority list not available (no database).")
            return
        path = self._get_priority_export_path()
        if not os.path.exists(path):
            await context.bot.send_message(chat_id, f"📂 Export file not found. Ask OpenClaw for job shortlist first, then tap Refresh.")
            return
        try:
            added, skipped = self.db_helper.sync_priority_from_yc_file(path)
            await context.bot.send_message(chat_id, f"✅ Sync YC: {added} added, {skipped} skipped.")
            await self._send_priority_list(context, chat_id)
        except Exception as e:
            logger.exception("_do_priority_sync error")
            await context.bot.send_message(chat_id, f"Error: {str(e)[:150]}")
    
    async def _do_priority_refresh(self, context, chat_id):
        """Run pipeline + sync (callback from button)."""
        if not self.db_helper:
            await context.bot.send_message(chat_id, "Priority list not available (no database).")
            return
        await context.bot.send_message(chat_id, "🔄 Refreshing from YC (30–60 sec)...")
        import subprocess
        path = self._get_priority_export_path()
        base = os.path.dirname(path)
        script = os.path.join(base, "run_shortlist.sh")
        if not os.path.exists(script):
            await context.bot.send_message(chat_id, "❌ run_shortlist.sh not found.")
            return
        try:
            subprocess.run(["bash", script], cwd=base, capture_output=True, timeout=120, env={**os.environ})
        except subprocess.TimeoutExpired:
            await context.bot.send_message(chat_id, "⏱️ Pipeline timed out. Try again.")
            return
        except Exception as e:
            await context.bot.send_message(chat_id, f"❌ Pipeline error: {str(e)[:100]}")
            return
        if os.path.exists(path):
            added, skipped = self.db_helper.sync_priority_from_yc_file(path)
            await context.bot.send_message(chat_id, f"✅ Refreshed! {added} added, {skipped} skipped.")
            await self._send_priority_list(context, chat_id)
        else:
            await context.bot.send_message(chat_id, "❌ Export file not created.")
    
    # Helper methods
    
    def _load_recent_jobs(self) -> List[Dict]:
        """Load jobs from recent cache file"""
        cache_dir = "autonomous_data/ats_cache"
        if not os.path.exists(cache_dir):
            return []
        
        # Get most recent cache file
        files = sorted([f for f in os.listdir(cache_dir) if f.endswith('.json')])
        if not files:
            return []
        
        latest_file = os.path.join(cache_dir, files[-1])
        try:
            with open(latest_file, 'r') as f:
                data = json.load(f)
                return data.get('jobs', [])
        except:
            return []
    
    def _count_todays_jobs(self) -> int:
        """Count jobs found today"""
        jobs = self._load_recent_jobs()
        return len(jobs)
    
    def _load_keywords(self) -> List[str]:
        """Load current search keywords"""
        return ["AI Product Engineer", "Applied LLM Engineer", "Founding Engineer AI", "LLM Engineer", "AI Agent Engineer", "AI Developer Tools Engineer"]
    
    def _save_keywords(self, keywords: List[str]):
        """Save updated keywords"""
        pass
    
    def _load_target_companies(self) -> Dict:
        """Load target companies from JSON"""
        try:
            with open('src/data/target_companies.json', 'r') as f:
                return json.load(f)
        except:
            return {}
    
    async def send_notification(self, message: str, parse_mode: str = 'Markdown'):
        """Send a notification message"""
        try:
            await self.app.bot.send_message(
                chat_id=self.chat_id,
                text=message,
                parse_mode=parse_mode
            )
        except Exception as e:
            print(f" Failed to send Telegram notification: {e}")
    
    def run(self):
        """Start the bot (blocking)"""
        print(" Starting Enhanced Telegram Bot...")
        self.app.run_polling()
    
    async def start(self):
        """Start the bot asynchronously (non-blocking)"""
        print("🤖 Starting Enhanced Telegram Bot (async mode)...")
        await self.app.initialize()
        await self.app.start()
        # Explicitly include callback_query to ensure button clicks work
        await self.app.updater.start_polling(
            allowed_updates=["message", "callback_query"]
        )
        print("✅ Enhanced Bot polling started (messages + callbacks)")


def create_enhanced_bot(token: str = None, chat_id: str = None, db_helper=None):
    """Create and return bot instance"""
    if token is None:
        token = os.getenv('TELEGRAM_BOT_TOKEN')
    if chat_id is None:
        chat_id = os.getenv('TELEGRAM_CHAT_ID')
    
    if not token or not chat_id:
        raise ValueError("TELEGRAM_BOT_TOKEN and TELEGRAM_CHAT_ID required")
    
    return EnhancedTelegramBot(token, chat_id, db_helper)


if __name__ == '__main__':
    bot = create_enhanced_bot()
    bot.run()
